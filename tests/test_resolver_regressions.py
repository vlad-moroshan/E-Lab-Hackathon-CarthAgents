"""Source-only regressions for the current-code review; no model/network calls."""
import contextlib
import copy
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from carth import ingest, output
from carth.layout import layout
from carth.resolve import resolve, to_pred
from carth.spec import Section, Spec


def section(sheet='Sheet', **changes):
    data=dict(name=sheet or 'PDF', pages=[], sheet=sheet, year=2025, unit='EUR', granularity='asset',
              gross_or_net='gross', record_lines=1, record_start=r'^Asset ',
              fields=[dict(line=0,band=2,field='zugang')],roles=[],
              ids=[dict(field='asset_no',regex=r'Asset ([A-Z])',scope='record')],year_columns_note='')
    data.update(changes)
    return Section(**data)


def spec(*sections):
    return Spec(doc_type='register',is_asset_schedule=True,noise_pages=[],sections=list(sections))


class ResolverRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def workbook(self, name='source.xlsx', sheets=None):
        w=Workbook();w.remove(w.active)
        for title,cells in (sheets or {'Sheet':{'A1':'Asset A','B1':1}}).items():
            ws=w.create_sheet(title)
            for coord,value in cells.items():ws[coord]=value
        p=self.root/name;w.save(p)
        return p

    def materialize(self,path,s):
        d=ingest.ingest(path);lay=layout(d);refs=resolve(d,lay,s)
        return refs,output.extract(d,refs,path.stem)

    def test_01_mixed_units_survive_json_and_workbook(self):
        p=self.workbook(sheets={sh:{'A1':'Asset A','B1':1} for sh in ('EUR1','EUR2','TEUR')})
        refs,r=self.materialize(p,spec(section('EUR1'),section('EUR2'),section('TEUR',unit='TEUR')))
        self.assertIsNone(refs['unit'])
        self.assertEqual([v['value'] for v in r['values']],[100,100,100000])
        self.assertEqual([v['trace']['unit'] for v in r['values']],['EUR','EUR','TEUR'])
        jp,wp=output.write_outputs(r,self.root/'results')
        saved=json.loads(jp.read_text());w=load_workbook(wp)
        for value in saved['values']:
            dest=value['workbook']
            self.assertEqual(w[dest['sheet']][dest['cell']].value,value['value']/100)

    def test_02_sections_on_one_sheet_use_explicit_ranges(self):
        p=self.workbook(sheets={'Sheet':{'A1':'Asset A','B1':1,'A2':'Asset B','B2':2}})
        _,r=self.materialize(p,spec(section(year=2024,start_line='Sheet!1',end_line='Sheet!1'),
                                   section(year=2025,line_ids=['Sheet!2'])))
        self.assertEqual([(row['asset_no'],row['year']) for row in r['rows']],[('A',2024),('B',2025)])
        with self.assertRaisesRegex(ValueError,'ambiguous section'):
            self.materialize(p,spec(section(year=2024),section(year=2025)))

    def test_02_parallel_tables_require_disjoint_columns(self):
        p=self.workbook(sheets={'Sheet':{'A1':'Asset A','B1':1,'C1':2}})
        _,r=self.materialize(p,spec(section(year=2024,line_ids=['Sheet!1']),
             section(year=2025,line_ids=['Sheet!1'],fields=[dict(line=0,band=3,field='zugang')])))
        self.assertEqual([v['value'] for v in r['values']],[100,200])
        with self.assertRaisesRegex(ValueError,'ambiguous section'):
            self.materialize(p,spec(section(line_ids=['Sheet!1']),section(line_ids=['Sheet!1'])))

    def test_03_year_fallback_stays_in_its_section(self):
        p=self.workbook(sheets={'Cover':{'A1':'zum 31.12.2024','A2':'zum 31.12.2024'},
            'Data':{'A1':'zum 31.12.2025','A2':'Asset A','B2':1}})
        _,r=self.materialize(p,spec(section('Data',year=None)))
        self.assertEqual(r['rows'][0]['year'],2025)
        self.assertEqual(r['values'][0]['value'],100)

    def test_03_unknown_or_ambiguous_year_stays_null_and_flagged(self):
        for heading in ('no period printed','zum 31.12.2024 bis 31.12.2025'):
            p=self.workbook(sheets={'Sheet':{'A1':heading,'A2':'Asset A','B2':1}})
            _,r=self.materialize(p,spec(section(year=None)))
            self.assertIsNone(r['rows'][0]['year'])
            self.assertIsNone(r['values'][0]['value'])
            self.assertIn('reporting year unresolved',r['values'][0]['flags'])
            output.write_outputs(r,self.root/'unknown-year')

    def test_04_automatic_batch_uses_collision_safe_ids(self):
        from carth.resolve import main
        for name,amount in [('one',1),('two',2)]:
            folder=self.root/name;folder.mkdir()
            w=Workbook();w.active['A1']='Asset A';w.active['B1']=amount;w.save(folder/'client.xlsx')
        expected=ingest.document_ids(self.root)
        out=self.root/'out'
        with patch('carth.spec.infer',return_value=(spec(section()),{'seconds':0})), \
             patch.object(sys,'argv',['carth.resolve',str(self.root),'--out',str(out)]), \
             contextlib.redirect_stdout(io.StringIO()):
            main()
        self.assertEqual(len(expected),2)
        for path,doc_id in expected.items():
            saved=json.loads((out/f'{doc_id}.json').read_text())
            self.assertEqual(saved['document']['source'],str(path.resolve()))
            self.assertTrue((out/f'{doc_id}.xlsx').exists())
            self.assertTrue((out/f'{doc_id}.layout.json').exists())
            self.assertTrue((out/'pred'/f'{doc_id}.pred.json').exists())

    def test_05_failure_does_not_hide_success_and_returns_nonzero(self):
        from carth.resolve import main
        p=self.workbook();(self.root/'bad.xlsx').write_text('not a workbook')
        out=self.root/'results'
        with patch('carth.spec.infer',return_value=(spec(section()),{'seconds':0})), \
             contextlib.redirect_stdout(io.StringIO()):
            code=main([str(self.root),'--out',str(out),'--workers','2'])
        self.assertEqual(code,1)
        self.assertTrue((out/'source.xlsx').exists())
        failures=json.loads((out/'_failures.json').read_text())
        self.assertEqual(len(failures['failures']),1)
        self.assertEqual(failures['failures'][0]['doc_id'],'bad')
        self.assertFalse(json.loads((out/'_batch.json').read_text())['complete'])

    def test_05_incomplete_and_empty_selection_fail(self):
        from carth.resolve import main
        p=self.workbook();out=self.root/'out'
        missing=section(ids=[])
        with patch('carth.spec.infer',return_value=(spec(missing),{'seconds':0})), \
             contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main([str(p),'--out',str(out)]),1)
            self.assertEqual(main([str(p),'--out',str(self.root/'empty'),'--docs','unknown']),1)
        self.assertEqual(json.loads((out/'_failures.json').read_text())['incomplete'][0]['blanks'],1)

    def test_05_real_cli_failure_exit(self):
        p=self.root/'bad.xlsx';p.write_text('not a workbook')
        proc=subprocess.run([sys.executable,'-m','carth.resolve',str(p),'--out',str(self.root/'out')],
                            text=True,capture_output=True)
        self.assertEqual(proc.returncode,1,proc.stderr)
        self.assertIn('BadZipFile',proc.stdout)

    def test_06_prediction_conversion_preserves_scoring_identity(self):
        from carth.eval.score import score_doc
        for level in ('class','account','position'):
            p=self.workbook(name=f'{level}.xlsx',sheets={'Sheet':{'A1':f'{level} C','B1':1}})
            se=section(granularity=level,record_start=f'^{level} ',
                       ids=[dict(field=level,regex=f'{level} ([A-Z])',scope='record')])
            _,r=self.materialize(p,spec(se))
            key=dict(meta=dict(doc_id=level,family='synthetic',format='xlsx',variant='perfect',unit='EUR'),
                rows=[dict(year=2025,kind='asset',level=level,**{level:'C'},fields={'zugang':100},
                      printed={'zugang':'1,00'},pointer={'sheet':'Sheet','cell':'A1','cells':{'zugang':'B1'}})])
            direct=score_doc(p.with_suffix('.key.json'),key,r)
            converted=score_doc(p.with_suffix('.key.json'),key,to_pred(r))
            self.assertEqual((direct.score,converted.score),(1,1),level)

    def test_07_mixed_directions_keep_separate_lines_and_amounts(self):
        from reportlab.pdfgen import canvas
        p=self.root/'directions.pdf';c=canvas.Canvas(str(p),pagesize=(400,600))
        for i,code in enumerate('ABC'):
            c.drawString(100,500-i*20,f'Asset {code} {i+1},00')
        c.saveState();c.translate(100,100);c.rotate(90)
        for i,code in enumerate('DEF'):
            c.drawString(0,-i*20,f'Asset {code} {i+4},00')
        c.restoreState();c.save()
        doc=ingest.ingest(p);original=copy.deepcopy(doc);lay=layout(doc)
        self.assertEqual(doc,original)
        self.assertEqual(layout(doc),lay)
        self.assertEqual(len(lay['pages'][0]['lines']),6)
        sections=[]
        for frame in lay['pages'][0]['frames']:
            lines=[ln for ln in lay['pages'][0]['lines'] if ln['frame']==frame['id']]
            for line in lines:
                self.assertEqual(len({lay['tokens'][cell['tid']]['dir'] for cell in line['cells']}),1)
            band=next(b['id'] for b in lay['pages'][0]['bands'] if b['frame']==frame['id'])
            sections.append(section(sheet=None,pages=[1],line_ids=[ln['id'] for ln in lines],
                                    fields=[dict(line=0,band=band,field='zugang')]))
        result=output.extract(doc,resolve(doc,lay,spec(*sections)),'directions')
        self.assertEqual([v['value'] for v in result['values']],[100,200,300,400,500,600])
        for value in result['values']:
            self.assertEqual(value['pointer']['bbox'],next(t for t in original['tokens'] if t['id']==value['trace']['token'])['bbox'])

    def test_08_each_leaf_granularity_requires_its_own_identity(self):
        p=self.workbook()
        for granularity,primary in [('asset','asset_no'),('class','class'),('account','account'),('position','position')]:
            _,r=self.materialize(p,spec(section(granularity=granularity,ids=[])))
            self.assertIsNone(r['values'][0]['value'],granularity)
            self.assertIn(f'row identity unresolved: {primary}',r['values'][0]['flags'])
            self.assertIn(f'row identity unresolved: {primary}',r['rows'][0]['flags'])

    def test_08_heading_identity_is_valid_for_summary_leaves(self):
        p=self.workbook(sheets={'Sheet':{'A1':'Class C','A2':'Asset A','B2':1}})
        se=section(granularity='class',ids=[dict(field='class',regex=r'Class ([A-Z])',scope='heading')])
        _,r=self.materialize(p,spec(se))
        self.assertEqual(r['rows'][0]['class'],'C')
        self.assertEqual(r['values'][0]['value'],100)

    def test_09_invalid_offsets_and_patterns_are_rejected(self):
        for kwargs in [dict(record_lines=0),dict(record_start='['),
                       dict(fields=[dict(line=-1,band=2,field='zugang')]),
                       dict(fields=[dict(line=1,band=2,field='zugang')]),
                       dict(ids=[dict(field='asset_no',regex=r'Asset [A-Z]',scope='record')]),
                       dict(roles=[dict(regex='[',role='header')])]:
            with self.subTest(kwargs=kwargs),self.assertRaises(ValueError):section(**kwargs)

    def test_09_reference_bounds_are_checked_against_source(self):
        p=self.workbook()
        for band in (0,3):
            with self.assertRaisesRegex(ValueError,'workbook column'):
                self.materialize(p,spec(section(fields=[dict(line=0,band=band,field='zugang')])))
        with self.assertRaisesRegex(ValueError,'unknown sheet'):
            self.materialize(p,spec(section('Missing')))
        with self.assertRaisesRegex(ValueError,'unknown/out-of-section'):
            self.materialize(p,spec(section(line_ids=['Sheet!999'])))
        from reportlab.pdfgen import canvas
        pdf=self.root/'bands.pdf';c=canvas.Canvas(str(pdf))
        for i in range(3):c.drawString(20,700-i*20,f'Asset A {i+1},00')
        c.save()
        with self.assertRaisesRegex(ValueError,'unknown amount band'):
            self.materialize(pdf,spec(section(sheet=None,pages=[1],fields=[dict(line=0,band=999,field='zugang')])))
        with self.assertRaisesRegex(ValueError,'invalid PDF page'):
            self.materialize(pdf,spec(section(sheet=None,pages=[2])))

    def test_10_metric_counts_conflicting_votes(self):
        from collections import Counter
        from carth.eval.spec_metrics import score_spec
        metric=score_spec(spec(section()),{(0,2):Counter(zugang=5,ahk_close=5)},1)
        self.assertEqual(metric['n'],10)
        self.assertEqual(metric['right'],.5)
        self.assertEqual(metric['wrong'],.5)

    def test_10_metric_preserves_sheet_and_section_context(self):
        from carth.eval.spec_metrics import truth,score_spec
        p=self.workbook(sheets={'One':{'A1':'Asset A','B1':1},'Two':{'A1':'Asset B','B1':2}})
        lay=layout(ingest.ingest(p))
        key={'rows':[dict(pointer={'sheet':sheet,'cell':'A1','cells':{field:'B1'}})
                     for sheet,field in [('One','zugang'),('Two','ahk_close')]]}
        votes,lines=truth(key,lay)
        good=spec(section('One'),section('Two',fields=[dict(line=0,band=2,field='ahk_close')]))
        metric=score_spec(good,votes,lines,lay)
        self.assertEqual((metric['n'],metric['right'],metric['wrong']),(2,1,0))
        bad=spec(section('One'),section('Two'))
        metric=score_spec(bad,votes,lines,lay)
        self.assertEqual((metric['n'],metric['right'],metric['wrong']),(2,.5,.5))
        missing=score_spec(spec(section('One')),votes,lines,lay)
        self.assertEqual((missing['n'],missing['right'],missing['missing']),(2,.5,.5))

    def test_02_record_cannot_join_across_excluded_source_lines(self):
        p=self.workbook(sheets={'Sheet':{'A1':'Asset A','B1':1,'B2':2,'B3':3}})
        _,r=self.materialize(p,spec(section(record_lines=2,line_ids=['Sheet!1','Sheet!3'],
                                 fields=[dict(line=1,band=2,field='zugang')])))
        self.assertIsNone(r['values'][0]['value'])
        self.assertIn('record line not found',r['values'][0]['flags'])

    def test_02_two_tables_on_one_pdf_page_keep_their_years(self):
        from reportlab.pdfgen import canvas
        p=self.root/'tables.pdf';c=canvas.Canvas(str(p))
        for i,code in enumerate('ABC'):c.drawString(20,700-i*20,f'Asset {code} {i+1},00')
        c.save();doc=ingest.ingest(p);lay=layout(doc);lines=lay['pages'][0]['lines']
        first=section(sheet=None,pages=[1],year=2024,start_line=lines[0]['id'],end_line=lines[1]['id'],
                      fields=[dict(line=0,band=0,field='zugang')])
        last=section(sheet=None,pages=[1],year=2025,line_ids=[lines[2]['id']],
                     fields=[dict(line=0,band=0,field='zugang')])
        r=output.extract(doc,resolve(doc,lay,spec(first,last)),'tables')
        self.assertEqual([row['year'] for row in r['rows']],[2024,2024,2025])
        self.assertEqual([v['value'] for v in r['values']],[100,200,300])


if __name__=='__main__':unittest.main()
