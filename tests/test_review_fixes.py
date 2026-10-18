"""Regression cases for review findings; fixtures are synthetic and temporary."""
import contextlib
import copy
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from reportlab.pdfgen import canvas

from carth import ingest
from carth.cli import main
from carth.eval import score
from carth.eval.evidence import Ref, pdf_refs, xlsx_refs
from carth.eval.recall import doc_recall, recall_corpus
from carth.ingest.pdf_text import pdf_tokens
from carth.numbers import parse_amount, source_amount, split_glued
from carth.output import extract, validate, write_outputs


class ReviewFixTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def workbook(self):
        p = self.root / 'source.xlsx'
        w = Workbook()
        w.active['A1'] = '=1+2'
        w.active['B1'] = 0
        w.active['C1'] = ArrayFormula(ref='C1:C2', text='=ROW(C1:C2)')
        w.active['D1'] = DataTableFormula(ref='D1:D2', dt2D=False, r1='A1')
        w.active['F1'] = -12.34
        w.active['G1'] = -12.34
        w.active['A2'] = 'row boundary'
        w.save(p)
        return p

    def test_pointer_surrounds_characters(self):
        ref = Ref('token', 1, [20, 20, 50, 30])
        self.assertTrue(score._accepts(ref, {'page':1, 'bbox':[20,20,50,30]}, None))
        for b in ([35,25,35,25], [35,25,35.01,25.01], [20,24,50,26], [0,0,100,100],
                  [20,20,float('nan'),30], [50,30,20,20], [-1,20,50,30]):
            self.assertFalse(score._accepts(ref, {'page':1,'bbox':b}, None), b)
        self.assertFalse(score.valid_box([20,20,50,30], (40,40)))
        for kind in ('weak','multi','none'):
            self.assertFalse(score._accepts(Ref(kind,1,[20,20,50,30]), {'page':1,'bbox':[20,20,50,30]}, None))

    def test_identity_covers_source_granularity(self):
        row = dict(year=2025, kind='asset', level='class', **{'class':'C', 'account':'A', 'position':'P'})
        self.assertTrue(score._identity_ok(row,row))
        for k in row:
            self.assertFalse(score._identity_ok({**row,k:'WRONG'},row), k)

    def test_signed_substrings_and_unknown_columns_not_promoted(self):
        base = {'meta':{'unit':'EUR'}, 'rows':[{'year':2025,'fields':{'umbuchung':100}, 'printed':{'umbuchung':'1,00'},
                'pointer':{'page':1,'bbox':[0,0,200,20]}}]}
        for text in ('-1,00', 'ID1,00123', '1,00'):
            refs,_ = pdf_refs(base,[{'page':1,'bbox':[100,2,125,12],'text':text,'src':'pdf'}],{1:(200,200)},True)
            self.assertNotEqual(refs[0,'umbuchung'].kind, 'token')

    def test_unprinted_fixture_identity_is_not_required(self):
        p=self.workbook()
        key=dict(meta=dict(doc_id='source',family='x',format='xlsx',variant='perfect',unit='EUR'),rows=[
            dict(year=2025,kind='asset',asset_no='A',**{'class':'UNPRINTED', 'position':'UNPRINTED.POS'},
                 fields={'zugang':1234},pointer={'sheet':'Sheet','cell':'F1','cells':{'zugang':'F1'}})])
        pred=dict(rows=[dict(year=2025,kind='asset',asset_no='A',fields={'zugang':1234},
                            pointer={'sheet':'Sheet','cells':{'zugang':'F1'}})])
        self.assertEqual(score.score_doc(p.with_suffix('.key.json'),key,pred).score,1)
        pred['rows'][0]['class']='WRONG'
        self.assertEqual(score.score_doc(p.with_suffix('.key.json'),key,pred).score,-3)

    def test_field_learning_uses_other_rows_and_rejects_wrong_column(self):
        rows, toks = [], []
        for i in range(6):
            value = 100+i
            rows.append(dict(fields={'zugang':value}, printed={'zugang':f'1,{i:02d}'},
                             pointer={'page':1,'bbox':[0,i*20,200,i*20+20]}))
            toks.append(dict(page=1,text=f'1,{i:02d}',bbox=[100,i*20+2,125,i*20+12],src='pdf'))
        toks[-1]['bbox'] = [10,102,35,112]
        refs,_ = pdf_refs({'meta':{'unit':'EUR'},'rows':rows},toks,{1:(200,200)},True)
        self.assertEqual(refs[0,'zugang'].kind,'token')
        self.assertNotEqual(refs[5,'zugang'].kind,'token')
        rows[0]['flags']={'zugang':'cut_off'}
        refs,_ = pdf_refs({'meta':{'unit':'EUR'},'rows':rows},toks,{1:(200,200)},True)
        self.assertEqual(refs[0,'zugang'].kind,'none')

    def test_formulas_never_become_blank_zeros(self):
        p = self.workbook()
        doc = ingest.ingest(p)
        at = {t['cell']:t for t in doc['tokens']}
        for coord in ('A1','C1','D1'):
            self.assertEqual(at[coord]['kind'],'formula_uncached')
            self.assertEqual(at[coord]['cache_state'],'missing')
        key = dict(meta={'unit':'EUR'},rows=[dict(fields={'zugang':0},printed={'zugang':''},pointer={'sheet':'Sheet','cell':'A1','cells':{'zugang':'A1'}})])
        self.assertEqual(doc_recall(key,doc)['missing'],1)
        self.assertEqual(xlsx_refs(key,p,doc['tokens'])[0,'zugang'].kind,'none')
        key['rows'][0]['pointer']['cells']['zugang']='E1'
        self.assertEqual(xlsx_refs(key,p,doc['tokens'])[0,'zugang'].kind,'blank_cell')

    def test_source_key_conflict_scores_source_and_transfer_sign(self):
        p = self.workbook()
        k = dict(meta=dict(doc_id='source',family='synthetic',format='xlsx',variant='perfect',unit='EUR'),
                 rows=[dict(year=2025,kind='asset',asset_no='A',fields={'umbuchung':0},printed={'umbuchung':'0'},
                       pointer={'sheet':'Sheet','cell':'A1','cells':{'umbuchung':'F1'}})])
        for val, expected in ((0,-3),(-1234,1),(1234,-3)):
            q=copy.deepcopy(k);q['rows'][0]['fields']['umbuchung']=val
            self.assertEqual(score.score_doc(p.with_suffix('.key.json'),k,q).score,expected)
        oracle=score.oracle_pred(k,'evidence',p.with_suffix('.key.json'))
        self.assertEqual(oracle['rows'][0]['fields']['umbuchung'],-1234)
        self.assertEqual(score.score_doc(p.with_suffix('.key.json'),k,oracle).extra,0)

    def test_inferred_zero_cannot_read_a_populated_cell(self):
        p=self.workbook();w=load_workbook(p);w.active['F2']=99;w.save(p)
        key={'meta':{'unit':'EUR'},'rows':[
            dict(year=2025,kind='asset',fields={'zugang':1234},pointer={'sheet':'Sheet','cell':'A1','cells':{'zugang':'F1'}}),
            dict(year=2025,kind='asset',fields={'zugang':0},pointer={'sheet':'Sheet','cell':'A2'})]}
        refs=xlsx_refs(key,p,ingest.ingest(p)['tokens'])
        self.assertEqual(refs[1,'zugang'].kind,'weak')
        self.assertIsNone(refs[1,'zugang'].observed)

    def test_collision_safe_discovery_and_fingerprints(self):
        for part in ('one','two'):
            d=self.root/part;d.mkdir();w=Workbook();w.save(d/'client.xlsx')
        (self.root/'~$client.xlsx').write_text('not a workbook')
        w=Workbook();w.save(self.root/'_client.xlsx')
        paths=ingest.document_ids(self.root)
        self.assertEqual(len(paths),3)
        self.assertEqual(len(set(paths.values())),3)
        self.assertEqual(ingest.find_documents(self.root/'~$client.xlsx'),[])
        doc=ingest.ingest(next(iter(paths)))
        self.assertTrue(ingest.cache_matches(doc,next(iter(paths))))
        self.assertFalse(ingest.cache_matches(doc,list(paths)[1]))
        doc['frame_version']=0
        self.assertFalse(ingest.cache_matches(doc,next(iter(paths))))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main(['ingest',str(self.root),'--out',str(self.root/'tokens')]),0)
        self.assertEqual(len([p for p in (self.root/'tokens').glob('*.json') if not p.name.startswith('_failures')]),3)

    def test_ingest_isolates_bad_documents(self):
        p=self.workbook();(self.root/'broken.xlsx').write_text('broken')
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main(['ingest',str(self.root),'--out',str(self.root/'tokens')]),1)
        self.assertTrue((self.root/'tokens/source.json').exists())
        self.assertEqual(len(json.loads((self.root/'tokens/_failures.json').read_text())['failures']),1)

    def test_cardinal_and_small_rotations(self):
        for angle in (0,90,180,270,1,-1):
            p=self.root/f'{angle}.pdf';c=canvas.Canvas(str(p),pagesize=(400,600));c.translate(200,300);c.rotate(angle);c.drawString(0,0,'12.345,67');c.save()
            tok=pdf_tokens(p)[0]
            self.assertEqual(tok['text'],'12.345,67')
            self.assertEqual(tok['dir'],angle%360 if angle%90==0 else 0)
            if angle in (1,-1):
                self.assertGreater(tok['rbox'][2]-tok['rbox'][0],tok['rbox'][3]-tok['rbox'][1])
                self.assertAlmostEqual(tok['skew'],angle,places=2)

    def test_image_table_with_selectable_heading_needs_ocr(self):
        from PIL import Image
        from reportlab.lib.utils import ImageReader
        p=self.root/'mixed.pdf';c=canvas.Canvas(str(p),pagesize=(400,600));c.drawString(20,570,'Header')
        c.drawImage(ImageReader(Image.new('RGB',(300,400),'white')),20,20,width=300,height=400);c.save()
        d=ingest.ingest(p)
        self.assertTrue(d['pages'][0]['has_text'])
        self.assertTrue(d['pages'][0]['needs_ocr'])

    def test_ocr_orientation_quality_and_page_subset(self):
        import types
        try:
            import numpy
        except ImportError:
            sys.modules['numpy'] = types.ModuleType('numpy')
        from carth.ingest import ocr
        from PIL import Image
        class Page:
            def __init__(self, image): self.image=image
            def render(self, **kw): return self
            def to_pil(self): return self.image
        class Engine:
            name='stub'
            def orientation(self,img): return 90,1.0
            def words(self,img):
                return [('junk',[10,20,30,40],.4)]*10 if img.width>img.height else [('123,45',[10,20,30,40],.99)]
        pages=[Page(Image.new('L',(100,200))) for _ in range(3)]
        with patch.object(ocr.pdfium,'PdfDocument',return_value=pages):
            all_tokens=ocr.ocr_pages('stub',Engine(),[1,2,3],72)
            one=ocr.ocr_pages('stub',Engine(),[3],72)
        self.assertEqual(all_tokens[-1],one[0])
        self.assertEqual(one[0]['angle'],0)
        self.assertEqual(one[0]['text'],'123,45')
        self.assertFalse(one[0]['orientation_uncertain'])
        self.assertEqual(one[0]['classifier_angle'],90)

    def test_no_spill_or_merged_zero_inference(self):
        from carth.ingest.xlsx import empty_amount_cell
        p=self.workbook();w=load_workbook(p);w.active.merge_cells('E2:F2');w.save(p)
        doc=ingest.ingest(p)
        self.assertFalse(empty_amount_cell(doc,'Sheet','C2'))
        self.assertFalse(empty_amount_cell(doc,'Sheet','D2'))
        self.assertFalse(empty_amount_cell(doc,'Sheet','F2'))
        self.assertFalse(empty_amount_cell(doc,'Sheet','XFD1'))
        self.assertTrue(empty_amount_cell(doc,'Sheet','E1'))

    def test_runtime_does_not_import_eval(self):
        r=subprocess.run([sys.executable,'-c',"import carth.cli, carth.output, sys; assert not any(m.startswith('carth.eval') for m in sys.modules)"],capture_output=True,text=True)
        self.assertEqual(r.returncode,0,r.stderr)

    def test_parser_context_spacing_and_ambiguity(self):
        self.assertIsNone(parse_amount('1 2 3,45'))
        self.assertEqual(parse_amount('1\u00a0234,56').cents(),123456)
        for s in ('2025','123456789'):
            self.assertIsNone(source_amount(s,amount_column=False)[0])
        self.assertEqual(split_glued('1234'),[])
        self.assertIsNone(source_amount('1.234.56',amount_column=True,ocr=True)[0])
        self.assertEqual(source_amount('-1,00',amount_column=True,signed=True)[0],-100)

    def test_missing_caches_stay_in_recall_denominator(self):
        key={'meta':{'doc_id':'d','family':'x','format':'pdf'},'rows':[{'printed':{'zugang':'1,00'}}]}
        with patch('carth.eval.recall.load_keys',return_value={'d':(Path('d'),key)}):
            report=recall_corpus(self.root,self.root)
        self.assertIn('missing token documents: 1',report)
        self.assertIn('100.0',report)

    def test_json_and_real_template_round_trip(self):
        p=self.workbook();doc=ingest.ingest(p)
        spec={'sha256':doc['sha256'],'unit':'EUR','rows':[
            dict(id='a',year=2025,kind='asset',granularity='asset',asset_no='A', fields={
                'umbuchung':{'sheet':'Sheet','cell':'F1','amount_column':True},
                'afa_gj':{'sheet':'Sheet','cell':'G1','amount_column':True},
                'zugang':{'sheet':'Sheet','cell':'A1','amount_column':True}}),
            dict(id='p',year=2024,kind='position_subtotal',granularity='position',position='P',fields={
                'ahk_open':{'sheet':'Sheet','cell':'E1','amount_column':True}})]}
        result=extract(doc,spec,'synthetic')
        self.assertEqual([v['value'] for v in result['values']],[-1234,1234,None,0])
        jp,wp=write_outputs(result,self.root/'result')
        saved=json.loads(jp.read_text());validate(saved)
        wb=load_workbook(wp)
        for v in saved['values']:
            target=v['workbook'];self.assertIsNotNone(target)
            self.assertEqual(wb[target['sheet']][target['cell']].value,None if v['value'] is None else v['value']/100)
        self.assertIn('Quellwerte',wb.sheetnames)
        self.assertIsNone(wb[saved['values'][1]['workbook']['sheet']]['AC2'].value)
        self.assertIsNone(wb[saved['values'][1]['workbook']['sheet']]['AD2'].value)
        self.assertEqual(saved['values'][1]['workbook']['cell'],'AR2')
        self.assertEqual(len(wb.defined_names),len(load_workbook('challenge/Arbeitsversion-merged-layout-20.xlsx').defined_names))

    def test_writer_expands_grid_and_preserves_literal_totals(self):
        p=self.root/'many.xlsx';w=Workbook();ws=w.active
        for i in range(1,32):ws[f'A{i}']=i
        w.save(p);doc=ingest.ingest(p)
        spec={'sha256':doc['sha256'],'unit':'EUR','rows':[
            dict(id=f'r{i}',year=2025,kind='grand_total' if i==31 else 'asset',granularity='class',
                 **{'class':f'C{i}'},fields={'ahk_close':{'sheet':'Sheet','cell':f'A{i}','amount_column':True}})
            for i in range(1,32)]}
        result=extract(doc,spec,'many');jp,wp=write_outputs(result,self.root/'many-output')
        saved=json.loads(jp.read_text());wb=load_workbook(wp)
        for v in saved['values']:
            dest=v['workbook'];self.assertEqual(wb[dest['sheet']][dest['cell']].value,v['value']/100)
        self.assertEqual(wb['Gitter nach Anlagenklasse']['I38'].value,31)
        self.assertEqual(wb['Gitter nach Anlagenklasse']['V38'].value,'grand_total')

    def test_output_contract_rejects_invented_or_unsupported_values(self):
        d=ingest.ingest(self.workbook())
        r=extract(d,None,'blank');validate(r)
        self.assertTrue(r['flags'])
        r['seed']=123
        with self.assertRaises(Exception):validate(r)
        spec={'sha256':'wrong','unit':'EUR','rows':[]}
        with self.assertRaises(ValueError):extract(d,spec,'blank')


if __name__ == '__main__':
    unittest.main()
