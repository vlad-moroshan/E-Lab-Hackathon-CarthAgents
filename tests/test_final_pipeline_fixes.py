"""Final-pipeline regressions using synthetic sources; no model or OCR engine calls."""
import copy
import json
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from carth import ingest, output
from carth.layout import layout
from carth.resolve import _one, attach_ocr, resolve
from carth.spec import RowRole
from test_resolver_regressions import section, spec


class FinalPipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def source(self, rows):
        path = self.root / 'source.xlsx'
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        doc = ingest.ingest(path)
        return path, doc, layout(doc)

    def extract(self, doc, lay, s):
        return output.extract(doc, resolve(doc, lay, s), 'source')

    def test_overlapping_sections_withhold_conflicts_and_flag_saved_outputs(self):
        _, doc, lay = self.source([['Asset A', 1], ['Asset B', 2]])
        r = self.extract(doc, lay, spec(section(record_start='^Asset A'), section(record_start='^Asset B')))
        self.assertFalse(r['values'])
        self.assertTrue(any('ambiguous section ownership' in f for f in r['flags']))
        jp, wp = output.write_outputs(r, self.root/'out')
        self.assertEqual(json.loads(jp.read_text())['flags'], r['flags'])
        self.assertIn('ambiguous section ownership', load_workbook(wp)['Offene Punkte']['A4'].value)

    def test_conflicts_are_local_to_each_concurrent_document(self):
        _, doc, lay = self.source([['Asset A', 1]])
        def run(conflict):
            return self.extract(doc, lay, spec(section(), section()) if conflict else spec(section()))
        with ThreadPoolExecutor(2) as pool:
            bad, good = list(pool.map(run, [True, False]))
        self.assertTrue(bad['flags'])
        self.assertEqual(good['flags'], [])
        self.assertEqual(good['values'][0]['value'], 100)

    def test_uppercase_suffix_and_subnumber_are_preserved(self):
        _, doc, lay = self.source([['123A', 1, '123A forklift']])
        r = self.extract(doc, lay, spec(section(record_start='^123A',
            ids=[dict(field='asset_no', regex='^(123A)', scope='record')])))
        self.assertEqual(r['rows'][0]['asset_no'], '123A')
        self.assertIsNone(r['rows'][0]['sub_no'])

    def test_subtotal_routing_uses_its_level_and_preserves_leaf_granularity(self):
        _, doc, lay = self.source([['Summe P', 1]])
        for leaf in ('asset', 'class', 'account', 'position'):
            r = self.extract(doc, lay, spec(section(granularity=leaf, record_start='^Summe',
                roles=[dict(regex='^Summe', role='position_subtotal')],
                ids=[dict(field='position', regex='Summe ([A-Z])', scope='record')])))
            jp, wp = output.write_outputs(r, self.root/leaf)
            saved = json.loads(jp.read_text())
            self.assertEqual(saved['rows'][0]['granularity'], leaf)
            self.assertEqual(saved['values'][0]['workbook']['sheet'], 'Gitter nach Bilanzposition')
            ws = load_workbook(wp)['Gitter nach Bilanzposition']
            self.assertEqual((ws['B8'].value, ws['E8'].value), ('P', 1))

    def test_workbook_and_ledger_preserve_composite_asset_identity(self):
        _, doc, lay = self.source([['Asset A Sub 1', 1], ['Asset A Sub 2', 2]])
        r = self.extract(doc, lay, spec(section(ids=[
            dict(field='asset_no', regex='Asset ([A-Z])', scope='record'),
            dict(field='sub_no', regex='Sub ([12])', scope='record')])))
        _, wp = output.write_outputs(r, self.root/'out')
        wb = load_workbook(wp)
        ws = wb['Anlagenverzeichnis']
        self.assertEqual([(ws[f'A{i}'].value, ws[f'AT{i}'].value) for i in (2,3)], [('A','1'), ('A','2')])
        self.assertEqual([wb['Quellwerte'][f'K{i}'].value for i in (2,3)], ['1','2'])

    def test_missing_row_is_incomplete_in_manifest_json_and_workbook(self):
        path, _, _ = self.source([['Asset A', 1], ['Asset B', 2]])
        out = self.root/'out'
        out.mkdir()
        with patch('carth.spec.infer', return_value=(spec(section(record_start='^Asset A')), dict(seconds=0))):
            status = _one(path, out, None, 'source')
        self.assertEqual(status['status'], 'incomplete')
        self.assertIn('Sheet!2', '; '.join(status['flags']))
        self.assertIn('Sheet!2', '; '.join(json.loads((out/'source.json').read_text())['flags']))
        self.assertIn('Sheet!2', load_workbook(out/'source.xlsx')['Offene Punkte']['A4'].value)

    def test_coverage_includes_outside_explicit_range_but_not_explicit_noise(self):
        _, doc, lay = self.source([['Asset A', 1], ['Noise', 2]])
        s = section(start_line='Sheet!1', end_line='Sheet!1')
        r = self.extract(doc, lay, spec(s))
        self.assertTrue(any('Sheet!2' in f for f in r['flags']))
        s.roles = [RowRole(regex='^Noise', role='noise')]
        self.assertFalse(self.extract(doc, lay, spec(s))['flags'])

    def test_entire_omitted_sheet_is_flagged(self):
        path, _, _ = self.source([['Asset A', 1]])
        wb = load_workbook(path)
        wb.create_sheet('Omitted').append(['Asset B', 2])
        wb.save(path)
        doc = ingest.ingest(path)
        r = self.extract(doc, layout(doc), spec(section()))
        self.assertIn('Omitted!1', '; '.join(r['flags']))

    def ocr_doc(self, mixed=False):
        return dict(source=str(self.root/'scan.pdf'), sha256='a'*64, frame_version=ingest.FRAME_VERSION,
            pages=[dict(page=1, has_text=mixed, needs_ocr=True)],
            tokens=[dict(id='t0', page=1, text='Heading', src='pdf')] if mixed else [])

    def word(self, **changes):
        w = dict(id='old', page=1, text='999,00', src='ocr:doctr', orientation_uncertain=False,
                 bbox=[1,1,20,10])
        w.update(changes)
        return w

    def cache(self, doc, **changes):
        cached = copy.deepcopy(doc)
        cached['tokens'] = [self.word()]
        cached['ocr'] = dict(engine='doctr', model=ingest.DOCTR_MODEL, dpi=300, pages=[1],
                             pipeline_version=ingest.OCR_PIPELINE_VERSION)
        cached.update(changes)
        cache = self.root/'cache'
        cache.mkdir(exist_ok=True)
        (cache/'scan.json').write_text(json.dumps(cached))
        return cache

    def test_current_cache_reuses_only_ocr_and_preserves_native_tokens(self):
        doc = self.ocr_doc(mixed=True)
        cache = self.cache(doc)
        c = json.loads((cache/'scan.json').read_text())
        c['tokens'].append(dict(id='native', page=1, text='Heading', src='pdf'))
        (cache/'scan.json').write_text(json.dumps(c))
        with patch('carth.resolve._live_ocr') as live:
            attach_ocr(doc, Path(doc['source']), cache)
        live.assert_not_called()
        self.assertEqual([t['text'] for t in doc['tokens']], ['Heading', '999,00'])
        self.assertEqual(len({t['id'] for t in doc['tokens']}), 2)

    def test_legacy_stale_and_wrong_pipeline_caches_fall_back_to_live(self):
        for change in ({'sha256':None}, {'sha256':'b'*64}, {'frame_version':None},
                       {'ocr':dict(engine='doctr', dpi=300, pages=[1])},
                       {'ocr':dict(engine='paddle', dpi=300, pages=[1], pipeline_version=ingest.OCR_PIPELINE_VERSION)},
                       {'tokens':[dict(id='old', page=1, src='ocr:doctr', text='999,00')]}):
            with self.subTest(change=change):
                doc = self.ocr_doc()
                cache = self.cache(doc, **change)
                original = (cache/'scan.json').read_bytes()
                with patch('carth.resolve._live_ocr', return_value=[self.word(text='1,00')]) as live:
                    attach_ocr(doc, Path(doc['source']), cache)
                live.assert_called_once()
                self.assertEqual(doc['tokens'][0]['text'], '1,00')
                self.assertEqual((cache/'scan.json').read_bytes(), original)
                self.assertIn('rejected_cache', doc['ocr_cache'])

    def test_mixed_page_gets_live_ocr_without_losing_heading(self):
        doc = self.ocr_doc(mixed=True)
        with patch('carth.resolve._live_ocr', return_value=[self.word()]) as live:
            attach_ocr(doc, Path(doc['source']), None)
        live.assert_called_once_with(Path(doc['source']), [1])
        self.assertEqual([t['text'] for t in doc['tokens']], ['Heading','999,00'])

    def test_missing_ocr_evidence_flags_partial_native_result(self):
        for words in ([], RuntimeError('OCR offline')):
            doc = self.ocr_doc(mixed=True)
            kwargs = dict(side_effect=words) if isinstance(words, Exception) else dict(return_value=words)
            with patch('carth.resolve._live_ocr', **kwargs):
                attach_ocr(doc, Path(doc['source']), None)
            self.assertTrue(doc['flags'])

    def test_corrupt_cache_falls_back_and_empty_pages_stay_flagged(self):
        doc = self.ocr_doc()
        cache = self.cache(doc)
        (cache/'scan.json').write_text('[]')
        with patch('carth.resolve._live_ocr', return_value=[]) as live:
            attach_ocr(doc, Path(doc['source']), cache)
        live.assert_called_once()
        self.assertTrue(doc['flags'])

    def test_cache_retains_orientation_abstention(self):
        doc = self.ocr_doc()
        cache = self.cache(doc, tokens=[self.word(orientation_uncertain=True)])
        with patch('carth.resolve._live_ocr') as live:
            attach_ocr(doc, Path(doc['source']), cache)
        live.assert_not_called()
        refs = dict(sha256=doc['sha256'], unit='EUR', rows=[dict(id='r1', year=2025, kind='asset',
            granularity='asset', asset_no='A', fields=dict(zugang=dict(token='o0', amount_column=True)))])
        r = output.extract(doc, refs, 'scan')
        self.assertIsNone(r['values'][0]['value'])
        self.assertEqual(r['values'][0]['flags'], ['OCR orientation unresolved'])

    def test_failed_ocr_diagnostics_survive_without_layout(self):
        doc = self.ocr_doc()
        with patch('carth.resolve._live_ocr', side_effect=RuntimeError('offline')):
            attach_ocr(doc, Path(doc['source']), None)
        self.assertIn('OCR unavailable', '; '.join(output.extract(doc, None, 'scan')['flags']))


if __name__ == '__main__':
    unittest.main()
