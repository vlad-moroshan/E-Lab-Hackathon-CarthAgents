import base64
import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from carth.server import process_file


class UploadIntegrationTests(unittest.TestCase):
    def test_real_cells_to_json_and_populated_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / 'schedule.xlsx'
            wb = Workbook()
            wb.active.title = 'Assets'
            wb.active.append(['ID', 'Opening', 'Additions'])
            wb.active.append(['001', 123.45, 0])
            wb.save(source)
            specs = root / 'layouts'
            specs.mkdir()
            spec = {'sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'unit': 'EUR', 'rows': [
                {'id': 'r1', 'year': 2025, 'kind': 'asset', 'granularity': 'asset', 'asset_no': '001',
                 'fields': {'ahk_open': {'sheet': 'Assets', 'cell': 'B2', 'amount_column': True},
                            'zugang': {'sheet': 'Assets', 'cell': 'C2', 'amount_column': True}}}]}
            (specs / 'schedule.layout.json').write_text(json.dumps(spec))
            payload = process_file(source, specs, root / 'output')
            self.assertEqual([v['value'] for v in payload['result']['values']], [12345, 0])
            result = load_workbook(io.BytesIO(base64.b64decode(payload['workbook']['base64'])))
            self.assertEqual(result['Anlagenverzeichnis']['W2'].value, 123.45)
            self.assertEqual(result['Anlagenverzeichnis']['X2'].value, 0)
            self.assertEqual(result['Anlagenverzeichnis']['A2'].value, '001')
            self.assertEqual(payload['result']['values'][0]['pointer'], {'sheet': 'Assets', 'cell': 'B2'})
            with self.assertRaisesRegex(ValueError, 'reviewed layout'):
                process_file(source, None, root / 'missing', auto=False)
            spec['sha256'] = '0' * 64
            (specs / 'schedule.layout.json').write_text(json.dumps(spec))
            with self.assertRaisesRegex(ValueError, 'fingerprint'):
                process_file(source, specs, root / 'mismatch')
