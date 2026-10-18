#!/usr/bin/env python3
"""Shared core of the A3 hackathon corpus: invented fixed-asset data, a table layout spec, xlsx / text-PDF /
scan renderers, planted problems and the JSON answer key. Layout families (families/f*.py) build on it.

Every amount is an int in cents. Every company, person and asset name is invented (checked by check_corpus.py).
Pieces copied, not imported: `accumulated` and `de` from A3-synthetische-Arbeitsversionen/_generator/synth_av.py,
`post_scan` (now `scan`) from A3-Kontennachweis-Fakes/_generator/make_fakes.py.

Self-check:  python3 -I core.py
"""

from __future__ import annotations

import copy
import io
import json
import random
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageFilter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

PDFTOPPM = shutil.which("pdftoppm") or "/opt/homebrew/bin/pdftoppm"
SUPP = Path("/System/Library/Fonts/Supplemental")
for _name, _file in (("Narrow", "Arial Narrow.ttf"), ("Narrow-Bold", "Arial Narrow Bold.ttf"),
                     ("CourierNew", "Courier New.ttf"), ("CourierNew-Bold", "Courier New Bold.ttf")):
    if (SUPP / _file).exists():
        pdfmetrics.registerFont(TTFont(_name, str(SUPP / _file)))

AMOUNTS = ("ahk_open", "zugang", "abgang", "umbuchung", "zuschreibung", "ahk_close", "kum_afa_open", "afa_gj",
           "afa_abgang", "kum_afa_close", "bw_close", "bw_prev", "abgang_bw")
TOTAL_KINDS = ("class_subtotal", "account_subtotal", "position_subtotal", "grand_total", "carry_over")
KINDS = ("asset", "continuation", "heading", *TOTAL_KINDS, "blank", "noise", "page_break")

# ============================================================================ catalog (HGB § 266 Abs. 2 A., SKR03)

POSITIONS = {
    "A.I.2": "Entgeltlich erworbene Konzessionen, gewerbliche Schutzrechte und ähnliche Rechte und Werte sowie "
             "Lizenzen an solchen Rechten und Werten",
    "A.II.1": "Grundstücke, grundstücksgleiche Rechte und Bauten einschließlich der Bauten auf fremden Grundstücken",
    "A.II.2": "Technische Anlagen und Maschinen",
    "A.II.3": "Andere Anlagen, Betriebs- und Geschäftsausstattung",
    "A.II.4": "Geleistete Anzahlungen und Anlagen im Bau",
    "A.III.3": "Beteiligungen",
}
# (class code, account, class name, position, kind, ND months, EUR low, EUR high, item names). make_client takes the
# first n_classes: 3 brings GWG, 4 brings Anlagen im Bau (with a transfer), 8 brings Beteiligungen (with a Zuschreibung).
CATALOG = [
    ("3000", "0210", "Maschinen", "A.II.2", "afa", (96, 120, 156, 180), 8e3, 450e3,
     ["CNC-Drehmaschine", "CNC-Fräszentrum 5-Achs", "Bandsäge", "Abkantpresse", "Hydraulikpresse", "Schweißroboter",
      "Schraubenkompressor", "Laserschneidanlage", "Spritzgießmaschine", "Ständerbohrmaschine", "Flachschleifmaschine",
      "Verpackungsmaschine", "Etikettiermaschine", "Teilewaschanlage", "3D-Koordinatenmessmaschine"]),
    ("4100", "0420", "Büroeinrichtung", "A.II.3", "afa", (156,), 400, 25e3,
     ["Bürostuhl ergonomisch", "Schreibtisch höhenverstellbar", "Aktenschrank", "Besprechungstisch", "Empfangstheke",
      "Rollcontainer", "Regalsystem Archiv", "Konferenzstühle (8 Stück)", "Schrankwand Büro"]),
    ("4900", "0480", "Geringwertige Wirtschaftsgüter", "A.II.3", "gwg", (12,), 250, 800,
     ["Akkuschrauber", "Monitor 27 Zoll", "Bürostuhl", "Kaffeevollautomat", "Werkzeugkoffer", "Laserdrucker", "Tablet",
      "Messschieber digital", "Aktenvernichter", "Stehleuchte"]),
    ("5000", "0120", "Anlagen im Bau", "A.II.4", "aib", (0,), 20e3, 300e3, ["Anlagen im Bau"]),
    ("4200", "0320", "Pkw", "A.II.3", "afa", (72,), 18e3, 85e3,
     ["Pkw Kombi Diesel", "Pkw Kompaktklasse", "Pkw Elektro", "Pkw Mittelklasse Hybrid", "Kleintransporter"]),
    ("1000", "0027", "EDV-Software", "A.I.2", "afa", (36, 60), 1500, 120e3,
     ["ERP-Lizenz", "CAD-Software Lizenz", "Finanzbuchhaltung Software", "Office-Lizenzen",
      "Warenwirtschaft Modul Lager", "Zeiterfassung Software"]),
    ("2000", "0085", "Grund und Boden", "A.II.1", "land", (0,), 40e3, 600e3,
     ["Grund und Boden Betriebsgelände", "Grundstück Gewerbegebiet", "Erweiterungsfläche"]),
    ("6000", "0510", "Beteiligungen", "A.III.3", "fin", (0,), 2e3, 250e3,
     ["Geschäftsanteile Genossenschaftsbank", "Beteiligung Einkaufsgenossenschaft", "Anteile Servicegesellschaft"]),
    ("2100", "0100", "Fabrikbauten", "A.II.1", "afa", (396,), 150e3, 2.5e6,
     ["Produktionshalle", "Lagerhalle", "Anbau Sozialtrakt", "Werkstattgebäude"]),
    ("4300", "0350", "Lkw", "A.II.3", "afa", (108,), 30e3, 180e3,
     ["Lkw 7,5t Pritsche", "Sattelzugmaschine", "Kastenwagen", "Tandemanhänger"]),
    ("4400", "0490", "EDV-Hardware", "A.II.3", "afa", (36, 60), 600, 40e3,
     ["Serverschrank 42HE", "Server 2HE", "Notebook", "Netzwerk-Switch 48 Port", "Firewall-Appliance",
      "Workstation CAD", "NAS-Speicher"]),
    ("3100", "0240", "Technische Anlagen", "A.II.2", "afa", (120, 180, 240), 10e3, 400e3,
     ["Lüftungsanlage Halle", "Druckluftnetz", "Photovoltaikanlage Dach", "Absauganlage", "Kälteanlage",
      "Brückenkran 5t", "Notstromaggregat"]),
    ("4000", "0400", "Betriebsausstattung", "A.II.3", "afa", (60, 96, 120), 800, 60e3,
     ["Gabelstapler", "Elektro-Hubwagen", "Werkbank", "Palettenregal", "Schweißgerät", "Hochdruckreiniger"]),
]  # fmt: skip
METHOD = {"afa": "linear", "gwg": "GWG", "aib": "keine", "land": "keine", "fin": "keine"}
SYL_A = ("Arv", "Bel", "Dorn", "Elm", "Fal", "Gor", "Hav", "Isk", "Jor", "Kel", "Lor", "Mav", "Norv", "Orl", "Pell",
         "Rav", "Sor", "Tarv", "Ulv", "Vorn", "Wend", "Zerl")
SYL_B = ("enholt", "ingrod", "astedt", "umark", "elwede", "ensiek", "ofrath", "ibrunn", "anhoff", "ekamp", "oswerth",
         "unbrock")
TRADES = ("Maschinenbau", "Metallbau", "Kunststofftechnik", "Feinwerktechnik", "Holzverarbeitung", "Logistik",
          "Elektroanlagen", "Verpackungstechnik", "Druckerei", "Fahrzeugbau")
FORMS = ("GmbH", "GmbH & Co. KG", "AG", "KG")
GENERIC = {w.lower() for w in TRADES + ("GmbH", "mbH", "mbB", "Partner", "Steuerberatungsgesellschaft", "Steuerberater")}

# ============================================================================ data


def m(d: date) -> int:
    return d.year * 12 + d.month - 1


def accumulated(cost: int, nd: int, months: int, floor: int) -> int:
    """Linear AfA to date, monthly pro rata, never below the Erinnerungswert (round half up, integer only)."""
    if nd <= 0 or months <= 0 or cost <= 0:
        return 0
    return min(max(cost - floor, 0), (2 * cost * months + nd) // (2 * nd))


def fy_end_of(y: int, fe: tuple) -> date:
    return date(y, *fe)


def fy_of(d: date, fe: tuple) -> int:
    return d.year if (d.month, d.day) <= fe else d.year + 1


@dataclass
class AssetClass:
    code: str
    account: str
    name: str
    position: str
    kind: str
    nd: tuple = (0,)
    lo: float = 0
    hi: float = 0
    items: list = field(default_factory=list, repr=False)

    @property
    def bilanzposition(self) -> str:
        return f"{self.position} {POSITIONS[self.position]}"


@dataclass
class Asset:
    """One asset in one year. Laws: see check_corpus.LAWS."""

    year: int
    asset_no: str
    sub_no: str
    name: str
    cls: str
    account: str
    position: str
    acq_date: date
    nd_months: int
    method: str
    disposal_date: date | None
    ahk_open: int
    zugang: int
    abgang: int
    umbuchung: int
    zuschreibung: int
    ahk_close: int
    kum_afa_open: int
    afa_gj: int
    afa_abgang: int
    kum_afa_close: int
    bw_close: int
    bw_prev: int

    @property
    def abgang_bw(self) -> int:  # Restbuchwert of the disposal: the net-layout Abgang
        return self.abgang - self.afa_abgang

    @property
    def nd_years(self) -> str:
        return "" if not self.nd_months else f"{self.nd_months / 12:g}".replace(".", ",")


@dataclass
class _Life:  # one asset over its whole life; snapshots are cut from it
    k: AssetClass
    name: str
    acq: date
    nd: int
    floor: int
    ev: dict  # fiscal year -> [zugang, umbuchung]
    afa_from: date | None = None
    gone: date | None = None
    imp: dict = field(default_factory=dict)  # fiscal year -> außerplanmäßige AfA (+) or Zuschreibung (-)
    parent: _Life | None = None
    no: str = ""
    sub: str = ""

    def cost(self, y: int) -> int:
        return sum(z + u for fy, (z, u) in self.ev.items() if fy <= y)

    def planned(self, y: int, fe: tuple, until: date | None = None) -> int:
        c = self.cost(y)
        if c <= 0 or self.afa_from is None:
            return 0
        if self.k.kind == "gwg":
            return c  # written off in full in the year of acquisition
        return accumulated(c, self.nd, m(until or fy_end_of(y, fe)) - m(self.afa_from) + 1, self.floor)

    def snapshot(self, y: int, fe: tuple) -> Asset | None:
        gone_fy = fy_of(self.gone, fe) if self.gone else None
        if min(self.ev) > y or (gone_fy is not None and gone_fy < y):
            return None
        out = gone_fy == y
        c_open, (z, u) = self.cost(y - 1), self.ev.get(y, [0, 0])
        k_open = self.planned(y - 1, fe) + sum(v for fy, v in self.imp.items() if fy < y)
        imp = self.imp.get(y, 0)
        afa = self.planned(y, fe, self.gone if out else None) - self.planned(y - 1, fe) + max(imp, 0)
        zu = max(-imp, 0)
        abg = c_open + z + u if out else 0
        afa_abg = k_open + afa - zu if out else 0
        if not (c_open or z or u or k_open or afa):
            return None
        ahk, kum = c_open + z - abg + u, k_open + afa - afa_abg - zu
        k = self.k
        return Asset(y, self.no, self.sub, self.name, k.code, k.account, k.position, self.acq, self.nd,
                     METHOD[k.kind], self.gone if out else None, c_open, z, abg, u, zu, ahk, k_open, afa, afa_abg,
                     kum, ahk - kum, c_open - k_open)


@dataclass
class Client:
    seed: int
    company: str
    kanzlei: str
    signer: str
    year: int
    fy_end_md: tuple
    unit: str
    disposals_at: str
    classes: list
    assets: list = field(default_factory=list)
    lives: list = field(default_factory=list, repr=False)

    @property
    def fy_end(self) -> date:
        return fy_end_of(self.year, self.fy_end_md)

    @property
    def fy_start(self) -> date:
        return fy_end_of(self.year - 1, self.fy_end_md) + timedelta(days=1)

    @property
    def names(self) -> list[str]:  # every invented proper name; check_corpus searches them in the client files
        return [self.company, self.kanzlei, self.signer]

    @property
    def abgang_field(self) -> str:  # the one Abgang column a single-column layout prints
        return "abgang" if self.disposals_at == "cost" else "abgang_bw"

    def snapshot(self, year: int) -> Client:
        c = copy.copy(self)
        c.year = year
        order = {k.code: i for i, k in enumerate(self.classes)}
        snaps = (L.snapshot(year, self.fy_end_md) for L in self.lives)
        c.assets = sorted((a for a in snaps if a), key=lambda a: (order[a.cls], int(a.asset_no), a.sub_no))
        return c

    def years(self, n: int) -> list[Client]:
        """n consecutive years ending with this one, oldest first; closing of N-1 = opening of N."""
        return [self.snapshot(y) for y in range(self.year - n + 1, self.year + 1)]

    def by_class(self) -> list[tuple[AssetClass, list[Asset]]]:
        return [(k, g) for k in self.classes if (g := [a for a in self.assets if a.cls == k.code])]


def make_client(seed: int, *, year: int, fy_end: tuple = (12, 31), n_classes: int, assets_per_class: tuple = (3, 8),
                subnumbers: bool = False, unit: str = "EUR", disposals_at: str = "cost",
                numbering: str = "datev") -> Client:
    """Invented client with consistent asset lives; `.assets` is the roll-forward of `year`, `.years(n)` more."""
    assert 1 <= n_classes <= len(CATALOG) and disposals_at in ("cost", "book") and numbering in ("datev", "running")
    rng, fe = random.Random(f"a3hack-{seed}"), tuple(fy_end)
    classes = [AssetClass(*row) for row in CATALOG[:n_classes]]
    classes.sort(key=lambda k: (list(POSITIONS).index(k.position), k.account))
    lives: list[_Life] = []

    def day_in(y: int) -> date:
        start = fy_end_of(y - 1, fe) + timedelta(days=1)
        return start + timedelta(days=rng.randrange((fy_end_of(y, fe) - start).days + 1))

    def cost(k: AssetClass) -> int:
        c = int(k.lo * 100 * (k.hi / k.lo) ** rng.random())
        return c - c % 100 if rng.random() < 0.4 else c

    def item(k: AssetClass) -> str:
        base = rng.choice(k.items)
        if k.kind != "afa" or k.position == "A.II.1":
            return base
        return base + rng.choice(("", "", "", f" Nr. {rng.randrange(1, 40)}", f" Halle {rng.randrange(1, 6)}",
                                  f" Typ {rng.choice('ABDFHKLMRSTX')}{rng.randrange(10, 990)}", " (gebraucht)"))

    def life(k: AssetClass, fy: int, amount: int, *, name: str | None = None, ev: dict | None = None,
             nd: int | None = None) -> _Life:
        acq = day_in(fy)
        L = _Life(k, name or item(k), acq, rng.choice(k.nd) if nd is None else nd, 100 if k.kind == "afa" else 0,
                  ev or {fy: [amount, 0]}, acq if k.kind in ("afa", "gwg") else None)
        lives.append(L)
        return L

    keep = set()  # ids of the assets kept at Erinnerungswert 1,00
    for k in classes:
        if k.kind == "aib":
            continue
        n = rng.randint(1, 3) if k.kind in ("land", "fin") else rng.randint(*assets_per_class)
        for i in range(n):
            nd = rng.choice(k.nd)
            if k.kind in ("land", "fin"):
                fy = year - rng.randint(3, 25)
            elif k.kind == "gwg":
                fy = year - rng.choice((0, 0, 1, 1, 2, 3, 4))
            elif i == 0:  # old enough to be written down to the Erinnerungswert 1,00
                fy = year - -(-nd // 12) - rng.randint(1, 4)
            else:
                r = rng.random()
                fy = year if r < 0.22 else year - rng.randint(1, 4) if r < 0.45 else year - 4 - int(rng.triangular(0, 16, 1))
            L = life(k, fy, cost(k), nd=nd)
            if i == 0 and k.kind == "afa":
                keep.add(id(L))
            elif fy < year and k.kind in ("afa", "gwg") and rng.random() < 0.15:
                L.gone = day_in(rng.randint(max(fy + 1, year - 4), year))

    dep = [L for L in lives if L.k.kind == "afa"]
    if dep and not any(min(L.ev) == year for L in dep):  # at least one addition this year
        free = [L for L in dep if id(L) not in keep]
        if free:
            L = free[-1]
            L.acq = L.afa_from = day_in(year)
            L.ev, L.gone = {year: [sum(L.ev[min(L.ev)]), 0]}, None
        else:
            life(dep[0].k, year, cost(dep[0].k))
    if dep and not any(L.gone and fy_of(L.gone, fe) == year for L in lives):  # at least one disposal this year
        old = [L for L in lives if L.k.kind in ("afa", "gwg") and min(L.ev) < year and not L.gone and id(L) not in keep]
        (rng.choice(old) if old else life(dep[0].k, year - rng.randint(1, 5), cost(dep[0].k))).gone = day_in(year)
    dep = [L for L in lives if L.k.kind == "afa"]

    if subnumbers:  # later additions booked as sub-assets 1, 2 … of the main asset
        for P in [L for L in dep if min(L.ev) <= year - 2]:
            last = min(fy_of(P.gone, fe) - 1, year) if P.gone else year
            if rng.random() < 0.25 and last > min(P.ev):
                C = life(P.k, rng.randint(min(P.ev) + 1, last), int(P.cost(year) * rng.uniform(0.05, 0.3)),
                         name=P.name + rng.choice((" Erweiterung", " Nachrüstung", " Zusatzmodul")), nd=P.nd)
                C.parent, C.gone = P, P.gone

    aib = next((k for k in classes if k.kind == "aib"), None)
    if aib:  # last year's AiB, finished this year and transferred: Umbuchungen sum to zero
        to = next((k for k in classes if k.kind == "afa" and k.position == "A.II.2"), None) or dep[0].k
        z1, z2, name = cost(aib), cost(aib), item(to)
        life(aib, year - 1, 0, name=f"AiB {name}", ev={year - 1: [z1, 0], year: [z2, -(z1 + z2)]})
        life(to, year, 0, name=name, ev={year: [0, z1 + z2]})
        for _ in range(rng.randint(0, 1)):
            life(aib, year, cost(aib), name=f"Anlagen im Bau Halle {rng.randint(2, 6)}")

    fin = [L for L in lives if L.k.kind == "fin"]
    if fin:  # impaired two years ago, half written back this year
        x = fin[0].cost(year) * 3 // 10
        fin[0].imp = {year - 2: x, year: -(x // 2)}

    seq: dict[str, int] = {}
    mains = sorted((L for L in lives if L.parent is None), key=lambda L: L.acq)
    for i, L in enumerate(mains, 1):
        seq[L.k.code] = seq.get(L.k.code, 0) + 1
        L.no = str(1_000_000 + i) if numbering == "running" else str(int(L.k.account) * 1000 + seq[L.k.code])
        L.sub = "0" if subnumbers else ""
    for L in sorted((L for L in lives if L.parent), key=lambda L: L.acq):
        L.no, L.sub = L.parent.no, str(sum(1 for C in lives if C.parent is L.parent and C.acq <= L.acq))

    stems = rng.sample([a + b for a in SYL_A for b in SYL_B], 3)
    c = Client(seed, f"{stems[0]} {rng.choice(TRADES)} {rng.choice(FORMS)}",
               f"{stems[1]} & Partner mbB Steuerberatungsgesellschaft", f"{rng.choice('BDHKLMRSTW')}. {stems[2]}",
               year, fe, unit, disposals_at, classes, lives=lives)
    return c.snapshot(year)


# ============================================================================ formatting


def fmt(cents: int, style: str = "de", zero: str | None = None) -> str:
    """'de' 1.234,56 · 'de_nodots' 1234,56 · 'de_trailing_minus' 1.234,56- · 'paren' (1.234,56) · 'en' 1,234.56 ·
    'teur' 1.235 (rounded thousands). zero: None prints 0 normally, else the literal ('', '0,00', '-')."""
    if cents == 0 and zero is not None:
        return zero
    c = abs(cents)
    if style == "teur":
        body = f"{(c + 50_000) // 100_000:,}".replace(",", ".")
    else:
        e, ct = divmod(c, 100)
        body = {"en": f"{e:,}.{ct:02d}", "de_nodots": f"{e},{ct:02d}"}.get(style, f"{e:,}".replace(",", ".") + f",{ct:02d}")
    if cents >= 0 or body == "0":
        return body
    return {"de_trailing_minus": body + "-", "paren": f"({body})"}.get(style, "-" + body)


def fmt_date(d: date | None, style: str = "de") -> str:
    """'de' 31.12.2025 · 'de_short' 31.12.25 · 'iso' 2025-12-31 · 'month' 12/2025."""
    return "" if d is None else d.strftime({"de": "%d.%m.%Y", "de_short": "%d.%m.%y", "iso": "%Y-%m-%d", "month": "%m/%Y"}[style])


def round_teur(cents: int) -> int:
    return (1 if cents >= 0 else -1) * ((abs(cents) + 50_000) // 100_000) * 100_000


# ============================================================================ layout spec


@dataclass
class Cell:
    text: str
    colspan: int = 1


@dataclass(eq=False)  # identity: t.rows.index(row) must find this row, not an equal blank one
class Row:
    kind: str  # one of KINDS
    cells: list  # printed text per column; a 'noise' row holds the lines of its page instead
    truth: dict | None = None  # key data: year, asset_no, sub_no, class, account, position, name, …, fields, printed
    flags: dict = field(default_factory=dict)  # field -> planted problem ('*' = whole row)
    pointer: dict | None = None  # set by the last render


@dataclass
class Table:
    fields: list  # field per column: an Asset attribute ('asset_no', 'name', 'acq_date', AMOUNTS …) or None
    col_widths: list  # points
    align: list  # 'l' | 'r' | 'c' per column
    header_rows: list = field(default_factory=list)  # list of lists of Cell
    title_lines: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    num_style: str = "de"
    zero: str | None = None
    date_style: str = "de"
    unit: str = "EUR"
    font: str = "Helvetica"  # Helvetica, Courier, Times-Roman, Narrow, CourierNew
    font_size: float = 7.0
    landscape: bool = True
    repeat_header: bool = True  # title and column header on every page
    page_header: str = ""  # templates: {p} page, {n} pages, {date} print date
    page_footer: str = "Seite {p} von {n}"
    print_date: date | None = None
    rows_per_page: int | None = None
    rules: bool = True

    def text(self, f, v) -> str:
        if isinstance(v, date):
            return fmt_date(v, self.date_style)
        if f in AMOUNTS and isinstance(v, int):
            return fmt(v, self.num_style, self.zero)
        return "" if v is None else str(v)

    def make(self, kind: str, values: dict, truth: dict | None = None) -> Row:
        """values: {field name or column index: value}; ints in AMOUNTS columns are formatted as amounts."""
        cells = [self.text(f, values.get(i, values.get(f))) for i, f in enumerate(self.fields)]
        if truth is not None:
            truth["fields"] = {f: values[f] for f in self.fields if f in AMOUNTS and isinstance(values.get(f), int)}
            truth["printed"] = {f: cells[self.fields.index(f)] for f in truth["fields"]}
        return Row(kind, cells, truth)

    def add(self, kind: str, values: dict, truth: dict | None = None) -> Row:
        self.rows.append(r := self.make(kind, values, truth))
        return r

    def add_asset(self, a: Asset, kind: str = "asset") -> Row:
        values = {f: getattr(a, f) for f in self.fields if f and hasattr(a, f)}
        truth = {"year": a.year, "asset_no": a.asset_no, "sub_no": a.sub_no, "class": a.cls, "account": a.account,
                 "position": a.position, "name": a.name, "acq_date": a.acq_date, "nd_months": a.nd_months,
                 "method": a.method}
        return self.add(kind, values, truth)

    def total(self, kind: str, label: str, assets: list, *, cls=None, account=None, position=None, year=None) -> Row:
        """A sum row (not appended). The ids say which asset rows it sums: none = all rows of its year."""
        values = {f: sum(getattr(a, f) for a in assets) for f in self.fields if f in AMOUNTS}
        values["name" if "name" in self.fields else 0] = label
        ids = {"class": cls, "account": account, "position": position}
        return self.make(kind, values, {"year": year or assets[0].year, **{k: v for k, v in ids.items() if v}})

    def add_total(self, kind: str, label: str, assets: list, **ids) -> Row:
        self.rows.append(r := self.total(kind, label, assets, **ids))
        return r

    def set_amount(self, r: Row, f: str, cents: int) -> None:
        r.truth["fields"][f] = cents
        r.cells[self.fields.index(f)] = r.truth["printed"][f] = fmt(cents, self.num_style, self.zero)

    def copy(self) -> Table:
        return copy.deepcopy(self)


# ============================================================================ renderers


def _neg(text: str) -> bool:
    return text.startswith(("-", "(")) or text.endswith("-")


def render_xlsx(tables, path, sheet="Anlagenspiegel", numbers="number", number_format="#,##0.00",
                merged_headers=True, offset=(0, 0), extra_sheets=None) -> None:
    """tables: one Table or a list (stacked on one sheet, or one per sheet when `sheet` is a list).
    numbers='number' writes amounts as numbers (unless a problem changed the cell text), 'text' as printed strings."""
    tables = tables if isinstance(tables, list) else [tables]
    sheets = sheet if isinstance(sheet, list) else [sheet] * len(tables)
    wb = Workbook()
    wb.remove(wb.active)
    nxt: dict[str, int] = {}
    for t, name in zip(tables, sheets):
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        rr, c0 = nxt.get(name, offset[0] + 1), offset[1] + 1
        for i, s in enumerate(t.title_lines):
            ws.cell(rr, c0, s).font = Font(bold=i == 0)
            rr += 1
        rr += bool(t.title_lines)
        for hr in t.header_rows:
            col = c0
            for cell in hr:
                x = ws.cell(rr, col, cell.text or None)
                x.font, x.alignment = Font(bold=True), Alignment(horizontal="center" if cell.colspan > 1 else None)
                if merged_headers and cell.colspan > 1:
                    ws.merge_cells(start_row=rr, start_column=col, end_row=rr, end_column=col + cell.colspan - 1)
                col += cell.colspan
            rr += 1
        for row in t.rows:
            if row.kind in ("page_break", "noise"):
                continue
            printed, cells = (row.truth or {}).get("printed", {}), {}
            for j, (f, text) in enumerate(zip(t.fields, row.cells)):
                v, nf = text or None, None
                if numbers == "number" and text and f in printed and text == printed[f]:
                    cents = abs(row.truth["fields"][f]) * (-1 if _neg(text) else 1)
                    v, nf = (cents // 100_000, "#,##0") if t.unit == "TEUR" else (cents / 100, number_format)
                elif numbers == "number" and f == "acq_date" and text and row.truth:
                    v, nf = datetime.combine(row.truth["acq_date"], datetime.min.time()), "DD.MM.YYYY"
                x = ws.cell(rr, c0 + j, v)
                if nf:
                    x.number_format = nf
                if row.kind in TOTAL_KINDS or row.kind == "heading":
                    x.font = Font(bold=True)
                if f in printed:
                    cells[f] = x.coordinate
            row.pointer = {"sheet": name, "cell": f"{get_column_letter(c0)}{rr}", "cells": cells}
            rr += 1
        nxt[name] = rr + 2
        for j, w in enumerate(t.col_widths):
            ws.column_dimensions[get_column_letter(c0 + j)].width = round(w / 5.2, 1)
    for name, data in (extra_sheets or {}).items():
        ws = wb.create_sheet(name)
        for line in data:
            ws.append(list(line))
    wb.save(path)
    _freeze_zip(path)


def _freeze_zip(path) -> None:
    """Same input, same bytes: fixed zip dates and document dates."""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w") as z:
        for name, data in items:
            if name == "docProps/core.xml":
                data = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-01-01T00:00:00Z", data)
            z.writestr(zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0)), data, zipfile.ZIP_DEFLATED)


def _bold(font: str) -> str:
    b = {"Helvetica": "Helvetica-Bold", "Courier": "Courier-Bold", "Times-Roman": "Times-Bold"}.get(font, font + "-Bold")
    return b if b in pdfmetrics.getRegisteredFontNames() or b in pdfmetrics.standardFonts else font


def _plan(t: Table) -> list:
    """Pages of one table: a list of rows per table page, a lone Row per noise page."""
    H = (landscape(A4) if t.landscape else A4)[1]
    head = len(t.title_lines) + bool(t.title_lines) + len(t.header_rows) + bool(t.header_rows)
    plan, cur = [], []
    for r in t.rows:
        if r.kind in ("page_break", "noise"):
            plan += ([cur] if cur else []) + ([r] if r.kind == "noise" else [])
            cur = []
            continue
        n = int((H - 82) / (t.font_size * 1.35)) - (head if t.repeat_header or not any(isinstance(p, list) for p in plan) else 0)
        if len(cur) >= min(n, t.rows_per_page or n):
            plan, cur = plan + [cur], []
        cur.append(r)
    return plan + ([cur] if cur else [])


def render_pdf(tables, path) -> list[dict]:
    """Text PDF via reportlab. tables: one Table or a list (each starts a new page, own page numbers).
    Sets row.pointer = {page, line, bbox: [x0, top, x1, bottom]} (points, top-left origin, like pdfplumber)."""
    tables = tables if isinstance(tables, list) else [tables]
    c = canvas.Canvas(str(path), invariant=1, pageCompression=1)
    page, out = 0, []
    for t in tables:
        W, H = landscape(A4) if t.landscape else A4
        lh, size = t.font_size * 1.35, t.font_size
        x0 = max(24.0, (W - sum(t.col_widths)) / 2)
        xs = [x0 + sum(t.col_widths[:i]) for i in range(len(t.col_widths))]
        x1, top = x0 + sum(t.col_widths), H - 40
        plan = _plan(t)
        n, tp = sum(isinstance(p, list) for p in plan), 0
        for p in plan:
            page += 1
            c.setPageSize((W, H))
            if isinstance(p, Row):  # noise page: a letter, a blank sheet, an unrelated print
                c.setFont("Helvetica", 10)
                for i, s in enumerate(p.cells):  # text after a tab is right-aligned at 480 pt
                    left, _, right = s.partition("\t")
                    c.drawString(70, H - 90 - i * 14, left)
                    c.drawRightString(480, H - 90 - i * 14, right)
                p.pointer = {"page": page, "line": 1, "bbox": [0, 0, round(W, 1), round(H, 1)]}
                c.showPage()
                continue
            tp += 1
            fill = {"p": tp, "n": n, "date": fmt_date(t.print_date, t.date_style)}
            c.setFont(t.font, size)
            if t.page_header:
                c.drawRightString(x1, H - 24, t.page_header.format(**fill))
            if t.page_footer:
                c.drawCentredString(W / 2, 22, t.page_footer.format(**fill))
            y = top
            if tp == 1 or t.repeat_header:
                for i, s in enumerate(t.title_lines):
                    c.setFont(_bold(t.font) if i == 0 else t.font, size + (2 if i == 0 else 1))
                    c.drawString(x0, y, s)
                    y -= lh
                y -= lh * bool(t.title_lines)
                c.setFont(_bold(t.font), size)
                for hr in t.header_rows:
                    col = 0
                    for cell in hr:
                        w = sum(t.col_widths[col:col + cell.colspan])
                        _put(c, cell.text, xs[col], w, "c" if cell.colspan > 1 else t.align[col], y)
                        col += cell.colspan
                    y -= lh
                if t.header_rows and t.rules:
                    c.setLineWidth(0.5)
                    c.line(x0, y + lh * 0.55, x1, y + lh * 0.55)
                y -= lh * 0.4 * bool(t.header_rows)
            for r in p:
                total = r.kind in TOTAL_KINDS
                c.setFont(_bold(t.font) if total or r.kind == "heading" else t.font, size)
                if total and t.rules:
                    first = next((i for i, f in enumerate(t.fields) if f in AMOUNTS), 0)
                    c.setLineWidth(0.4)
                    c.line(xs[first], y + size + 0.6, x1, y + size + 0.6)
                    if r.kind == "grand_total":
                        c.line(xs[first], y - 2.6, x1, y - 2.6)
                        c.line(xs[first], y - 3.8, x1, y - 3.8)
                for i, s in enumerate(r.cells):
                    if s:
                        _put(c, s, xs[i], t.col_widths[i], t.align[i], y)
                r.pointer = {"page": page, "line": round((top - y) / lh) + 1,
                             "bbox": [round(x0, 1), round(H - y - size * 0.8, 1), round(x1, 1), round(H - y + size * 0.25, 1)]}
                out.append(r.pointer)
                y -= lh
            c.showPage()
    c.save()
    return out


def _put(c, text: str, x: float, w: float, align: str, y: float) -> None:
    if align == "r":
        c.drawRightString(x + w - 2, y, text)
    elif align == "c":
        c.drawCentredString(x + w / 2, y, text)
    else:
        c.drawString(x + 2, y, text)


def scan(pdf_in, pdf_out, dpi=150, skew_deg=0.5, noise=0.02, fax=False, rotate=None, seed=0) -> None:
    """Image-only PDF (no text layer) that looks scanned: grey paper, skew, noise, blur, JPEG or fax.
    rotate: {page index (0-based): degrees}. Adapted from make_fakes.post_scan; noise is seeded, so output is stable."""
    rng = random.Random(f"scan-{seed}")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([PDFTOPPM, "-r", str(dpi), "-gray", "-png", str(pdf_in), f"{tmp}/pg"], check=True)
        out = []
        for i, f in enumerate(sorted(Path(tmp).glob("pg*.png"))):
            im = Image.open(f).convert("L").point(lambda v: int(28 + v * 0.86))  # grey paper, toner never fully black
            if rotate and i in rotate:
                im = im.rotate(rotate[i], expand=True)
            im = im.rotate(rng.uniform(-skew_deg, skew_deg), resample=Image.BICUBIC, fillcolor=238)
            if noise:
                im = Image.blend(im, Image.frombytes("L", im.size, rng.randbytes(im.width * im.height)), noise)
            im = im.filter(ImageFilter.GaussianBlur(0.55))
            if fax:  # fax 'Normal' mode: half the vertical lines, hard black/white
                w, h = im.size
                im = im.resize((w, h // 2)).resize((w, h), Image.NEAREST).point(lambda v: 255 if v > 165 else 0).convert("1")
            else:
                buf = io.BytesIO()
                im.save(buf, "JPEG", quality=42)
                im = Image.open(io.BytesIO(buf.getvalue()))
            out.append(im)
        out[0].save(pdf_out, "PDF", resolution=dpi, save_all=True, append_images=out[1:],
                     title=None, creationDate=None, modDate=None)


def outline(pdf_in, pdf_out) -> None:
    """Clean PDF with NO text layer: Ghostscript turns every glyph into vector shapes (as some ERP print drivers do).
    Looks sharp, needs OCR. Use key format 'outline'. The Omit flags drop the date and ID, so rebuilds are byte-identical."""
    subprocess.run(["/opt/homebrew/bin/gs", "-q", "-o", str(pdf_out), "-sDEVICE=pdfwrite", "-dNoOutputFonts", "-dOmitInfoDate",
                    "-dOmitID", "-dOmitXMP", str(pdf_in)], check=True)


# ============================================================================ planted problems
# Each mutates the table in place (t.copy() first to keep the clean one) and returns (table, record).
# Order when combining: p_teur, p_glued / p_non_tying / p_wrap, p_no_repeat_header, p_carry_over, p_noise.
# A record's "where" items may hold Row objects; write_key turns them into their pointers.


def _amount_cols(t: Table) -> list[int]:
    return [i for i, f in enumerate(t.fields) if f in AMOUNTS]


def p_glued(t: Table, rng: random.Random, n: int = 2):
    """Two adjacent amounts printed with no space between them (the right column gets both)."""
    cols = _amount_cols(t)
    pairs = lambda r: [(i, j) for i, j in zip(cols, cols[1:]) if j == i + 1 and r.cells[i] and r.cells[j]]  # noqa: E731
    where = []
    for r in rng.sample(rows := [r for r in t.rows if r.kind == "asset" and pairs(r)], min(n, len(rows))):
        i, j = rng.choice(pairs(r))
        r.cells[j], r.cells[i] = r.cells[i] + r.cells[j], ""
        r.flags[t.fields[i]] = r.flags[t.fields[j]] = "glued"
        where.append({"row": r, "fields": [t.fields[i], t.fields[j]], "printed": r.cells[j]})
    return t, {"type": "glued_numbers", "where": where,
               "expected_behaviour": "split the glued string into the two amounts of the two adjacent columns"}


def p_non_tying(t: Table, rng: random.Random):
    """A printed subtotal that does not equal the sum of its rows. The key keeps the printed value and flags it."""
    subs = [r for r in t.rows if r.kind in ("class_subtotal", "account_subtotal", "position_subtotal")]
    r = rng.choice(subs or [r for r in t.rows if r.kind == "grand_total"])
    f = rng.choice([f for f, v in r.truth["fields"].items() if v] or list(r.truth["fields"]))
    steps = (100_000, 200_000, 900_000, 1_000_000) if t.unit == "TEUR" else (900, 9_000, 10_000, 100_000)
    true, delta = r.truth["fields"][f], rng.choice(steps)
    wrong = true + delta if true - delta < 0 or rng.random() < 0.5 else true - delta
    t.set_amount(r, f, wrong)
    r.flags["*"] = "non_tying"
    return t, {"type": "non_tying_subtotal", "where": [{"row": r, "field": f, "printed": wrong, "sum_of_rows": true}],
               "expected_behaviour": "report the printed value; flag that it differs from the sum of its rows "
                                     f"by {fmt(wrong - true)}"}


def p_wrap(t: Table, rng: random.Random, n: int = 2):
    """Asset description wrapped onto a continuation line below the row."""
    ni, where = t.fields.index("name"), []
    for r in rng.sample(rows := [r for r in t.rows if r.kind == "asset" and " " in r.cells[ni]], min(n, len(rows))):
        words = r.cells[ni].split()
        r.cells[ni], tail = " ".join(words[: max(1, len(words) // 2)]), " ".join(words[max(1, len(words) // 2):])
        cont = Row("continuation", [tail if i == ni else "" for i in range(len(t.fields))])
        t.rows.insert(t.rows.index(r) + 1, cont)
        where.append({"row": r, "continuation": cont, "name": r.truth["name"]})
    return t, {"type": "wrapped_description", "where": where,
               "expected_behaviour": "join the continuation line to the asset above; it is not a row of its own"}


def p_carry_over(t: Table, label: str = "Übertrag"):
    """Page break inside a class: 'Übertrag' line at the foot of one page and the head of the next, placed at the
    latest point on a page where two asset rows of one class meet. Run after p_teur / p_wrap / p_no_repeat_header."""
    same = lambda a, b: a.kind == b.kind == "asset" and a.truth["class"] == b.truth["class"]  # noqa: E731
    pages = [p for p in _plan(t) if isinstance(p, list)]
    cands = [(k == len(pages) - 1, len(p) - j, p[j])
             for k, p in enumerate(pages) for j in range(len(p) - 1) if same(p[j], p[j + 1])]
    if not cands:
        raise ValueError("p_carry_over needs two asset rows of one class next to each other")
    last = min(cands, key=lambda x: x[:2])[2]
    i = t.rows.index(last) + 1
    group = (last.truth["class"], last.truth["year"])
    run = [r for r in t.rows[:i] if r.kind == "asset" and (r.truth["class"], r.truth["year"]) == group]
    vals = {f: sum(r.truth["fields"][f] for r in run) for f in last.truth["fields"]}
    vals["name" if "name" in t.fields else 0] = label
    ids = {"year": last.truth["year"], "class": last.truth["class"]}
    a, b = t.make("carry_over", vals, dict(ids)), t.make("carry_over", vals, dict(ids))
    t.rows[i:i] = [a, Row("page_break", []), b]
    return t, {"type": "carry_over", "where": [{"row": a}, {"row": b}],
               "expected_behaviour": "carry-over lines repeat a running sum; they are not assets and not totals"}


def p_no_repeat_header(t: Table):
    """Title and column header only on page 1; ensures at least two pages."""
    t.repeat_header = False
    half = max(5, -(-sum(r.kind not in ("noise", "page_break") for r in t.rows) // 2))
    t.rows_per_page = min(t.rows_per_page or half, half)
    return t, {"type": "header_only_on_first_page", "where": [{"pages": "2+"}],
               "expected_behaviour": "keep the column meaning of page 1 for the following pages"}


def p_teur(t: Table):
    """All amounts in TEUR, each rounded on its own, so sums can miss by rounding. Run before other problems."""
    t.unit, t.num_style = "TEUR", "teur"
    for r in t.rows:
        for f, v in list((r.truth or {}).get("fields", {}).items()):
            t.set_amount(r, f, round_teur(v))
    t.title_lines = [re.sub(r"\bEUR\b", "TEUR", s) for s in t.title_lines]
    for hr in t.header_rows:
        for cell in hr:
            cell.text = re.sub(r"\bEUR\b", "TEUR", cell.text)
    return t, {"type": "teur_rounding", "where": [{"unit": "TEUR"}],
               "expected_behaviour": "amounts are thousand EUR, rounded one by one; report as printed, "
                                     "accept rounding differences"}


def noise_lines(c: Client, kind: str) -> list[str]:
    rng = random.Random(f"noise-{c.seed}-{kind}")
    if kind == "blank":
        return []
    if kind == "cover":
        return [c.kanzlei, "", c.company, "z. Hd. Geschäftsführung", "", fmt_date(c.fy_end + timedelta(days=45)), "",
                f"Jahresabschluss zum {fmt_date(c.fy_end)}", "", "Sehr geehrte Damen und Herren,", "",
                "anbei erhalten Sie die Unterlagen zum Anlagevermögen für den Jahresabschluss.",
                "Bitte prüfen Sie die Aufstellung und teilen Sie uns Abweichungen innerhalb von",
                "zwei Wochen mit.", "", "Mit freundlichen Grüßen", "", c.signer, "Steuerberater"]
    lines = ["Kontennachweis zur Bilanz - Umlaufvermögen (Auszug)", c.company, ""]
    for no, name in (("1200", "Bank"), ("1400", "Forderungen aus Lieferungen und Leistungen"),
                     ("1548", "Vorsteuer laufendes Jahr"), ("3980", "Bestand Waren"), ("1000", "Kasse")):
        lines.append(f"{no}   {name}\t{fmt(rng.randrange(50_000, 90_000_000))}")
    return lines


def p_noise(t: Table, c: Client, kind: str = "cover", at: str = "start"):
    """A page that is not part of the schedule: 'cover' letter, 'blank' sheet or 'unrelated' print (PDF only).
    'middle' goes between two table pages, so run it last (after p_carry_over)."""
    r, pages = Row("noise", noise_lines(c, kind)), [p for p in _plan(t) if isinstance(p, list)]
    mid = t.rows.index(pages[len(pages) // 2 - 1][-1]) + 1 if len(pages) > 1 else len(t.rows) // 2  # between pages
    t.rows.insert({"start": 0, "end": len(t.rows), "middle": mid}[at], r)
    return t, {"type": f"noise_page_{kind}", "where": [{"row": r}],
               "expected_behaviour": "ignore this page; it holds no schedule rows"}


def p_multi_year(tables: list[Table]):
    """Several years in one file: render the list with render_pdf / render_xlsx. Keys carry each row's year."""
    years = sorted({r.truth["year"] for t in tables for r in t.rows if r.truth})
    return tables, {"type": "multi_year", "where": [{"years": years}],
                    "expected_behaviour": "extract every year on its own and tag each row with its year"}


# ============================================================================ answer key


FORMATS = {"xlsx", "pdf", "scan", "outline"}
VARIANTS = {"perfect", "problem", "layout-mutation"}


def key_meta(c: Client, **kw) -> dict:
    """meta for write_key; pass doc_id, family, doc_type, format ('xlsx'|'pdf'|'scan'|'outline'), variant
    ('perfect'|'problem'|'layout-mutation': the content state only, a scan of a perfect print is 'perfect'); add more names to invented_names if the family invents any."""
    return {"unit": c.unit, "years": [c.year], "period_end": c.fy_end.isoformat(), "company": c.company,
            "invented_names": c.names, "seed": c.seed, **kw}


def _plain(o):
    if isinstance(o, date):
        return o.isoformat()
    if isinstance(o, Row):
        return o.pointer
    raise TypeError(type(o))


def write_key(path, meta: dict, rows, planted=()) -> dict:
    """rows: a Table, a list of Tables (unit taken from the first) or a list of Rows. Call right after rendering."""
    if isinstance(rows, Table):
        rows = [rows]
    if rows and isinstance(rows[0], Table):
        meta = {**meta, "unit": rows[0].unit}
        rows = [r for t in rows for r in t.rows]
    missing = {"doc_id", "family", "doc_type", "format", "unit", "years", "period_end", "variant"} - meta.keys()
    assert not missing, f"key meta lacks {sorted(missing)}"
    assert meta["format"] in FORMATS and meta["variant"] in VARIANTS, f"bad format/variant {meta['format']}/{meta['variant']}"
    keyed = [{"kind": r.kind, **r.truth, "pointer": r.pointer, **({"flags": r.flags} if r.flags else {})}
             for r in rows if r.truth]
    key = {"meta": meta, "rows": keyed, "planted": [p for p in planted if p]}
    Path(path).write_text(json.dumps(key, ensure_ascii=False, indent=1, default=_plain) + "\n", encoding="utf-8")
    return key


# ============================================================================ self-check


if __name__ == "__main__":
    assert fmt(123456) == "1.234,56" and fmt(-123456, "de_trailing_minus") == "1.234,56-"
    assert fmt(-123456, "paren") == "(1.234,56)" and fmt(123456, "en") == "1,234.56" and fmt(123456, "de_nodots") == "1234,56"
    assert fmt(123_456_789, "teur") == "1.235" and fmt(0, zero="-") == "-" and fmt(-40_000, "teur") == "0"
    assert accumulated(120_000, 12, 6, 100) == 60_000 and accumulated(120_000, 12, 99, 100) == 119_900
    for s in range(40):
        c = make_client(s, year=2025, n_classes=1 + s % 13, assets_per_class=(2, 9), subnumbers=s % 2 == 0,
                        fy_end=(6, 30) if s % 3 == 0 else (12, 31), numbering="running" if s % 5 == 0 else "datev")
        ys = c.years(5)
        for y in ys:
            for a in y.assets:
                assert a.ahk_close == a.ahk_open + a.zugang - a.abgang + a.umbuchung, a
                assert a.kum_afa_close == a.kum_afa_open + a.afa_gj - a.afa_abgang - a.zuschreibung, a
                assert a.bw_close == a.ahk_close - a.kum_afa_close and a.bw_prev == a.ahk_open - a.kum_afa_open, a
                assert min(a.ahk_open, a.zugang, a.abgang, a.zuschreibung, a.ahk_close, a.kum_afa_open, a.afa_gj,
                           a.afa_abgang, a.kum_afa_close, a.bw_close, a.bw_prev) >= 0, a
            assert sum(a.umbuchung for a in y.assets) == 0
            assert len({(a.asset_no, a.sub_no) for a in y.assets}) == len(y.assets)
        for y0, y1 in zip(ys, ys[1:]):
            close = {(a.asset_no, a.sub_no): (a.ahk_close, a.kum_afa_close) for a in y0.assets if a.ahk_close or a.kum_afa_close}
            opens = {(a.asset_no, a.sub_no): (a.ahk_open, a.kum_afa_open) for a in y1.assets if a.ahk_open or a.kum_afa_open}
            assert close == opens, (s, y1.year)
        a = c.assets
        assert any(x.zugang and not x.ahk_open for x in a) and any(x.abgang for x in a), s
        assert any(x.bw_close == 100 and x.ahk_close > 100 for x in a), s
        if any(k.kind == "gwg" for k in c.classes):
            assert any(x.method == "GWG" for x in a), s
        if any(k.kind == "aib" for k in c.classes):
            assert any(x.umbuchung for x in a), s
    assert make_client(3, year=2025, n_classes=6).assets == make_client(3, year=2025, n_classes=6).assets
    print("core self-check OK")
