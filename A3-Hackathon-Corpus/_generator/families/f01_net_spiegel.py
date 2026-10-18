"""f01_net_spiegel: net Anlagenspiegel ("Netto-Anlagespiegel ... steuerrechtlich") at account level, as printed by a
German mid-market fixed-asset module. Starts from the prior-year book value, no AHK columns.

Layout: landscape, Arial 7 pt, nearly all bold, airy (a blank line after every line). Three levels: numbered
Bilanzposition "0." to "9." with client labels > Gruppe (indent 1) > Konto (indent 2, one line of figures).
"Summe <Gruppe>" after every group (even a one-account group), "Summe <n>. <Position>" after every position, no grand
total. Heading block of 3 lines + a unit row with sign hints ("(+)EUR", "(-)EUR", "(+/-)EUR") on every page. Every zero
printed as "0,00". Footer "Datum: DD.MM.YY" left, "Seite: n" right. Group "Pkw" holds two accounts (Pkw + Lkw).
Sister family f02_asset_spiegel prints the same client tree one line per asset (same software run): it imports
client(), tree() and the helpers below.

Key: a Konto line is keyed as kind "asset" with level "account" and asset_no None (it is the leaf row of this
document); class = printed group label, position = printed position label, hgb_position = HGB § 266 code.
"""

import random
import string
import sys
import tempfile
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from reportlab.pdfbase.pdfmetrics import registerFont, stringWidth
from reportlab.pdfbase.ttfonts import TTFont

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import core  # noqa: E402
from core import Cell  # noqa: E402

FAMILY = Path(__file__).stem
for _name, _file in (("Arial", "Arial.ttf"), ("Arial-Bold", "Arial Bold.ttf")):  # embedded, as in the program's print
    registerFont(TTFont(_name, str(core.SUPP / _file)))
PAIR_SEED, PAIR_DOCS = 150, ("f01_net_spiegel_03", "f02_asset_spiegel_02")
KW = dict(year=2025, n_classes=13, assets_per_class=(5, 12), subnumbers=True)
IND1, IND2 = " " * 5, " " * 12
# Client tree of this software run: (position, [(group, [core class codes])]). Not § 266 order, no Finanzanlagen.
TREE = [
    ("0. Immaterielle Vermögensgegenstände", [("EDV-Software", ["1000"])]),
    ("1. Grund und Boden", [("Grund und Boden", ["2000"])]),
    ("2. Betriebsgebäude", [("Fabrikbauten", ["2100"])]),
    ("3. Maschinenpark", [("Maschinen", ["3000"])]),
    ("4. Betriebstechnik", [("Technische Anlagen Kto. 0240", ["3100"])]),
    ("5. Fuhrpark", [("Pkw", ["4200", "4300"])]),
    ("6. Werkstatt und Lager", [("Betriebsausstattung", ["4000"])]),
    ("7. Büro und Verwaltung", [("Büroeinrichtung", ["4100"]), ("EDV-Hardware", ["4400"])]),
    ("8. Geringwertige Wirtschaftsgüter", [("GWG", ["4900"])]),
    ("9. Anlagen im Bau", [("Anlagen im Bau", ["5000"])]),
]
FIELDS = ["name", "bw_prev", "zugang", None, "abgang_bw", "umbuchung", "zuschreibung", "afa_gj", None,
          "kum_afa_close", "bw_close"]
HEADS = [("Bilanzposition",), ("Buchwert", "Vorjahr"), ("Zugänge,", "Geschäftsjahr"),
         ("davon aktivierte", "Zinsen", "Fremdkapital"), ("Abgänge /", "Minderungen", "Geschäftsjahr"),
         ("Umbuchungen", "Geschäftsjahr"), ("Zuschreibungen", "Geschäftsjahr"), ("Normal-AfA", "Geschäftsjahr"),
         ("Sonder-AfA", "Geschäftsjahr"), ("Abschreibungen", "(kumuliert)"), ("Buchwert", "31.12.")]
ZERO = (3, 8)  # columns printed as 0,00 with no key field
UNITS = ["", "EUR", "(+)EUR", "(+)EUR", "(-)EUR", "(+/-)EUR", "(+)EUR", "EUR", "EUR", "(-)EUR", "EUR"]
SUMS = ("class_subtotal", "account_subtotal", "position_subtotal")


# ---------------------------------------------------------------- shared with f02_asset_spiegel


def client(seed: int, years: int = 1):
    """core.make_client with this software's tree: the run holds no Finanzanlagen, so that class is dropped."""
    out = []
    for c in core.make_client(seed, **KW).years(years):
        fin = {k.code for k in c.classes if k.kind == "fin"}
        c.classes = [k for k in c.classes if k.code not in fin]
        c.assets = [a for a in c.assets if a.cls not in fin]
        out.append(c)
    return out if years > 1 else out[0]


def tree(c) -> list:
    """[(position, [(group, [(AssetClass, assets)])])] in print order, empty parts dropped."""
    by = {k.code: (k, g) for k, g in c.by_class()}
    out = []
    for pos, groups in TREE:
        gs = [(g, [by[x] for x in codes if x in by]) for g, codes in groups]
        if gs := [(g, accts) for g, accts in gs if accts]:
            out.append((pos, gs))
    return out


def print_date(c, seed: int) -> date:
    rng = random.Random(f"printdate-{seed}")
    return date(c.year + 1, rng.randint(2, 9), rng.randint(1, 28))


def centred(text: str, w: float, align: str, font: str, size: float = 7.0) -> str:
    """Pad with spaces so core's left/right-aligned header cell prints centred over its column."""
    pad = " " * max(0, round((w - 4 - stringWidth(text, font, size)) / 2 / stringWidth(" ", font, size)))
    return text + pad if align == "r" else pad + text if align == "l" else text


def header(heads, units, widths, align, font, pdf=True) -> list:
    """PDF: heading lines top-aligned and centred, a blank line, the unit row. xlsx: one heading row + unit row."""
    if not pdf:
        return [[Cell(" ".join(h)) for h in heads], [Cell(u) for u in units]]
    rows = [[Cell(centred(h[i], w, a, font) if i < len(h) else "") for h, w, a in zip(heads, widths, align)]
            for i in range(max(map(len, heads)))]
    return rows + [[Cell("") for _ in heads], [Cell(centred(u, w, a, font)) for u, w, a in zip(units, widths, align)]]


def footer(left: str, right: str, width: float, font: str, size: float = 7.0) -> str:
    """'left ... right' as one line spanning the table width (core centres the footer line)."""
    gap = width - stringWidth((left + right).format(p="9", n="9", date="00.00.00"), font, size)
    return left + " " * max(1, round(gap / stringWidth(" ", font, size))) + right


def zero_cols(r, cols) -> None:
    """Amount columns with no key field (aktivierte Zinsen, Sonder-AfA) are always zero: print '0,00'."""
    for i in cols:
        r.cells[i] = "0,00"


def cent_drift(t, rng: random.Random, share: float):
    """The program's rounding drift: about `share` of the sum lines printed 1 or 2 cents off the sum of their rows."""
    sums = [r for r in t.rows if r.kind in SUMS and not r.flags]
    where = []
    for r in rng.sample(sums, max(1, round(len(sums) * share))):
        f = rng.choice([f for f, v in r.truth["fields"].items() if v > 2] or list(r.truth["fields"]))
        true = r.truth["fields"][f]
        wrong = true + rng.choice((-2, -1, 1, 2) if true > 2 else (1, 2))
        t.set_amount(r, f, wrong)
        r.flags["*"] = "non_tying"
        where.append({"row": r, "field": f, "printed": wrong, "sum_of_rows": true})
    return t, {"type": "non_tying_subtotal", "where": where,
               "expected_behaviour": "report the printed value; flag that it differs from the sum of its rows by "
                                     "0,01 or 0,02 EUR (rounding drift of the program)"}


def xlsx_fix(path, t, int_cols=()) -> None:
    """Excel export: the '0,00' of unkeyed amount columns as number 0, the ND column as integers."""
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    for r in t.rows:
        if not r.pointer:
            continue
        row = ws[r.pointer["cell"]].row
        for j, f in enumerate(t.fields):
            x = ws.cell(row, j + 1)
            if f is None and x.value == "0,00":
                x.value, x.number_format = 0, "#,##0.00"
            elif j in int_cols and isinstance(x.value, str) and x.value.isdigit():
                x.value = int(x.value)
    wb.save(path)
    core._freeze_zip(path)


def short_ids(c, seed: int):
    """Inventory numbers of this program: 1-4 digits in order of acquisition, not tied to the account; split items
    (core sub-assets) get a letter suffix a, b, … after the main number. Returns a -> (asset_no, suffix)."""
    rng, ids = random.Random(f"invnr-{seed}"), {}
    n = rng.randint(3, 60)
    for _, no in sorted({(a.acq_date, a.asset_no) for a in c.assets if a.sub_no in ("", "0")}):
        n += rng.choice((1, 1, 1, 2, 3, 7))
        ids[no] = str(n)
    return lambda a: (ids[a.asset_no], "" if a.sub_no in ("", "0") else string.ascii_lowercase[int(a.sub_no) - 1])


# ---------------------------------------------------------------- this family


def table(c, seed: int, pdf: bool = True) -> core.Table:
    widths, align, font = [196] + [56] * 10, ["l"] + ["r"] * 10, "Arial-Bold"
    title = f"Netto-Anlagespiegel {core.fmt_date(c.fy_end)} steuerrechtlich"
    heads = HEADS[:-1] + [("Buchwert", f"{c.fy_end:%d.%m.}")]
    t = core.Table(fields=FIELDS, col_widths=widths, align=align, zero=None, date_style="de_short", font=font,
                   rules=False, header_rows=header(heads, UNITS, widths, align, font, pdf),
                   title_lines=[title] + [""] * 5 if pdf else [title, c.company],
                   page_header=c.company if pdf else "", print_date=print_date(c, seed),
                   page_footer=footer("Datum: {date}", "Seite: {p}", sum(widths), font))
    gap = (lambda: t.add("blank", {})) if pdf else (lambda: None)
    for pos, groups in tree(c):
        t.add("heading", {0: pos})
        gap()
        p_assets = []
        for g, accts in groups:
            t.add("heading", {0: IND1 + g})
            gap()
            g_assets = []
            for k, assets in accts:
                r = t.total("asset", f"{IND2}{k.account}    {k.name}", assets, cls=g, account=k.account, position=pos)
                r.truth.update(asset_no=None, name=k.name, level="account", hgb_position=k.position)
                zero_cols(r, ZERO)
                t.rows.append(r)
                gap()
                g_assets += assets
            zero_cols(t.add_total("class_subtotal", f"{IND1}Summe {g}", g_assets, cls=g), ZERO)
            gap()
            p_assets += g_assets
        zero_cols(t.add_total("position_subtotal", f"Summe {pos}", p_assets, position=pos), ZERO)
        gap()
        gap()
    return t


MUT_FIELDS = ["name", "bw_prev", "zugang", "umbuchung", "abgang_bw", "zuschreibung", "afa_gj", "bw_close",
              "kum_afa_close"]


def table_mutation(c, seed: int) -> core.Table:
    """Same document, other print set-up: portrait, Arial Narrow, renamed and reordered columns, no group level,
    Abgang and AfA printed with a trailing minus, zero blank, rules, a grand total, 'Seite x von y'."""
    d0, d1 = core.fmt_date(c.fy_start), core.fmt_date(c.fy_end)
    widths, align, font = [172] + [46] * 8, ["l"] + ["r"] * 8, "Narrow"
    heads = [("Konto / Bezeichnung",), ("Restbuchwert", d0), ("Zugang",), ("Umbuchung",),
             ("Abgang zum", "Restbuchwert"), ("Zuschreibung",), ("Abschreibung", "lfd. Jahr"), ("Restbuchwert", d1),
             ("kumulierte", "Abschreibung")]
    t = core.Table(fields=MUT_FIELDS, col_widths=widths, align=align, zero="", num_style="de_trailing_minus",
                   font=font, landscape=False, header_rows=header(heads, [""] * 9, widths, align, "Narrow-Bold")[:-2],
                   title_lines=[c.company, f"Anlagenspiegel (Nettodarstellung) Steuerbilanz zum {d1}",
                                "Beträge in EUR"],
                   page_header="Druckdatum {date}", page_footer="Seite {p} von {n}", print_date=print_date(c, seed),
                   rows_per_page=30)

    def minus(r) -> None:  # outflows printed negative; the key keeps the positive amount
        for f in ("abgang_bw", "afa_gj"):
            if v := r.truth["fields"][f]:
                r.cells[MUT_FIELDS.index(f)] = r.truth["printed"][f] = core.fmt(-v, t.num_style)

    for pos, groups in tree(c):
        t.add("heading", {0: pos})
        p_assets = []
        for _g, accts in groups:
            for k, assets in accts:
                r = t.total("asset", f"   {k.account} {k.name}", assets, account=k.account, position=pos)
                r.truth.update(asset_no=None, name=k.name, level="account", hgb_position=k.position)
                minus(r)
                t.rows.append(r)
                p_assets += assets
        minus(t.add_total("position_subtotal", f"Zwischensumme {pos.split('. ', 1)[1]}", p_assets, position=pos))
        t.add("blank", {})
    minus(t.add_total("grand_total", "Anlagevermögen gesamt", c.assets))
    return t


def build(out: Path) -> None:
    d = out / FAMILY
    d.mkdir(parents=True, exist_ok=True)

    def key(n, c, fmt, variant, t, planted=(), **extra):
        meta = core.key_meta(c, doc_id=f"{FAMILY}_{n:02d}", family=FAMILY, doc_type="Anlagenspiegel (netto)",
                             format=fmt, variant=variant, **extra)
        core.write_key(d / f"{FAMILY}_{n:02d}.key.json", meta, t, planted)

    def pdf(n):
        return d / f"{FAMILY}_{n:02d}.pdf"

    pair = {"paired_with": PAIR_DOCS[1],
            "pairing": "same client and software run as the asset-level print; group (class) and position sums agree"}

    c = client(PAIR_SEED)  # 03 perfect text PDF, paired with f02_asset_spiegel_02
    t = table(c, PAIR_SEED)
    core.render_pdf(t, pdf(3))
    key(3, c, "pdf", "perfect", t, **pair)

    c = client(102)  # 07 perfect Excel export
    t = table(c, 102, pdf=False)
    core.render_xlsx(t, d / f"{FAMILY}_07.xlsx", sheet="Netto-Anlagespiegel")
    xlsx_fix(d / f"{FAMILY}_07.xlsx", t)
    key(7, c, "xlsx", "perfect", t)

    with tempfile.TemporaryDirectory() as tmp:
        c = client(103)  # 01 scan, heavy skew
        t = table(c, 103)
        core.render_pdf(t, Path(tmp) / "print.pdf")
        core.scan(Path(tmp) / "print.pdf", pdf(1), skew_deg=2.5, noise=0.03, seed=103)
        key(1, c, "scan", "perfect", t)

        c = client(104)  # 05 clean print without text layer
        t = table(c, 104)
        core.render_pdf(t, Path(tmp) / "print.pdf")
        core.outline(Path(tmp) / "print.pdf", pdf(5))
        key(5, c, "outline", "perfect", t)

    c, rng = client(105), random.Random(105)  # 02 problems: rounding drift in sum lines + wrapped account labels
    t = table(c, 105)
    planted = [cent_drift(t, rng, 1 / 7)[1]]
    t, rec = core.p_wrap(t, rng, n=3)
    for w in rec["where"]:  # core joins the words again: put the indent back
        w["row"].cells[0], w["continuation"].cells[0] = IND2 + w["row"].cells[0], IND2 + w["continuation"].cells[0]
    planted.append(rec)
    core.render_pdf(t, pdf(2))
    key(2, c, "pdf", "problem", t, planted)

    c, rng = client(106), random.Random(106)  # 06 problems: glued amounts + a cover letter in front
    t = table(c, 106)
    planted = [core.p_glued(t, rng, n=3)[1], core.p_noise(t, c, "cover", "start")[1]]
    core.render_pdf(t, pdf(6))
    key(6, c, "pdf", "problem", t, planted)

    cs = client(107, years=2)  # 08 problem: two fiscal years in one file
    ts, rec = core.p_multi_year([table(y, 107) for y in cs])
    core.render_pdf(ts, pdf(8))
    key(8, cs[-1], "pdf", "problem", ts, [rec], years=[y.year for y in cs])

    c = client(108)  # 04 layout mutation
    t = table_mutation(c, 108)
    core.render_pdf(t, pdf(4))
    key(4, c, "pdf", "layout-mutation", t)
