"""f05_sap_class_gitter: SAP-like HGB Anlagengitter at class level, 13-column wide version, an ALV list exported to
Excel and printed. Each record is a boxed 4-line block (key line, AHK line, AfA line, Buchwert line); subtotal levels
only by stars: * Bestandskonto, ** Bilanzposition, *** Geschäftsbereich, **** Buchungskreis; selection page first.

Class rows are the finest rows here: the key gives them kind 'asset', asset_no null, level 'class' (check_corpus ties
totals to 'asset' rows). Docs 01/02 are the trio with f06/f07 01/02 (_sap_common.TRIO).
"""

import random
import sys
import tempfile
from dataclasses import replace
from itertools import groupby
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

sys.path[:0] = [str(Path(__file__).resolve().parents[1]), str(Path(__file__).resolve().parent)]
import _sap_common as S  # noqa: E402
import core  # noqa: E402
from core import Cell  # noqa: E402

FAMILY = Path(__file__).stem
LINES4 = [["star", "bukr", "bilpos", "konto", "klasse", "klname", None, "curr"],  # slot 0: level stars
          [None, "ahk_open", "zugang", "abgang", "umbuchung", "nil", "nil", "ahk_close"],
          [None, "kum_afa_open", "afa_gj", "afa_abgang", "nil", "nil", "zuschreibung", "kum_afa_close"],
          [None, "bw_prev", None, None, None, None, None, "bw_close"]]
HEADS4 = [["", "BuKr GsBe", "BilanzPos", "BestKtoAHK", "Klasse", "", "", "Währg"],
          ["", "AHK GJ-Beg", "Zugang", "Abgang", "Umbuchung", "Nachaktivg", "Invest.Förderg", "aktuelle AHK"],
          ["", "AfA GJ-Beg", "AfA des Jahres", "AfA Abgang", "AfA Umbuchung", "AfA Nachaktivg", "Zuschreibungen",
           "kumulierte AfA"],
          ["", "Buchwrt GJ-Beg", "", "", "", "", "", "lfd Buchwert"]]
FIELDS4 = [f for line in LINES4 for f in line]
WIDTHS4 = [16] + [66] * 7
FLAT = ["bukr", "bilpos", "konto", "klasse", "name", "ahk_open", "zugang", "abgang", "umbuchung", "ahk_close",
        "kum_afa_open", "afa_gj", "afa_abgang", "zuschreibung", "kum_afa_close", "bw_prev", "bw_close", "curr"]
FLAT_HEADS = ["BuKr", "BilPos", "Konto", "Klasse", "Bezeichnung", "AHK GJ-Beg", "Zugang", "Abgang", "Umbuchung",
              "aktuelle AHK", "AfA GJ-Beg", "AfA des Jahres", "AfA Abgang", "Zuschreibung", "kumulierte AfA",
              "Buchwrt GJ-Beg", "lfd Buchwert", "Währg"]
FLAT_WIDTHS = [19, 29, 19, 19, 78] + [51] * 12 + [19]
WORDS = {"*": "Summe Bestandskonto", "**": "Summe Bilanzposition", "***": "Summe Geschäftsbereich",
         "****": "Summe Buchungskreis"}


def table(c: core.Client, flat: bool = False) -> core.Table:
    fields = FLAT if flat else FIELDS4
    t = core.Table(fields=fields, col_widths=FLAT_WIDTHS if flat else WIDTHS4 * 4,
                   align=["r" if f in core.AMOUNTS or f == "nil" else "l" for f in fields],
                   header_rows=[[Cell(h) for h in line] for line in ([FLAT_HEADS] if flat else HEADS4)],
                   font="CourierNew", font_size=5.6 if flat else 6.8, rows_per_page=26 if flat else 10)
    S.titles(t, c, "Anlagengitter  Bereich 01 Handelsbilanz", "Gitter HGB §268 Abs.2 - 13 Spalten (breit, unvollständig)")
    bk, y = S.bukr(c), c.year

    def add(kind, level, A, **v):
        name = v.pop("name", "")
        if flat:
            v["name"] = name or WORDS[level]
        elif level == "****":
            v.update(star="****" + bk, bukr="")  # glued, as the list prints it
        else:
            v.update(star="" if level == "class" else level, klname=name)
        truth = {"year": y, **v.pop("ids"), "level": level}
        return t.add(kind, {"bukr": bk, **v, **S.sums(A, fields)}, truth)

    ks = S.ledger_classes(c)
    for pos, group in groupby(ks, key=lambda k: k.position):
        for i, k in enumerate(group):
            A = [a for a in c.assets if a.cls == k.code]
            name = k.name[:22] if flat else k.name  # the one-line list cuts at the field length
            ids = {"asset_no": None, "sub_no": None, "class": k.code, "account": k.account, "position": pos, "name": name}
            r = add("asset", "class", A, ids=ids, bilpos=S.POS7[pos], konto=k.account, klasse=k.code, name=name)
            r.band = i == 0  # the list prints some class blocks bold on a grey band
            add("account_subtotal", "*", A, ids={"account": k.account}, bilpos=S.POS7[pos], konto=k.account)
        add("position_subtotal", "**", [a for a in c.assets if a.position == pos], ids={"position": pos},
            bilpos=S.POS7[pos])
    add("grand_total", "***", c.assets, ids={})
    add("grand_total", "****", c.assets, ids={})
    return S.finish(t)


def style(r, name_slots=(5, 6)):
    if not r.truth:
        return None, (), False
    if r.truth["level"] != "class":
        return S.YELLOW, (), False
    band = getattr(r, "band", False)
    return (S.GREY if band else None), [(0, *name_slots, S.LILAC)], band


def cover(t: core.Table) -> list:
    n = sum(r.kind == "asset" for r in t.rows)
    return [[["Sortierkriterien", "Aufst.", "Abst.", "Zwischensumme"], ["Buchungskreis", "X", "", "X"],
             ["Geschäftsbereich", "X", "", "X"], ["Bilanzposition", "X", "", "X"], ["Bestandskonto AHK", "X", "", "X"],
             ["Anlagenklasse", "X", "", ""]],
            [["Datenstatistik", "Anzahl"], ["Übergebene Sätze", str(n)],
             ["Ermittelte Summenzeilen", str(sum(r.kind in core.TOTAL_KINDS for r in t.rows))]]]


LOOK = S.Look(font="CourierNew", size=6.8, x0=14, widths=WIDTHS4, lines=[list(range(i, i + 8)) for i in (0, 8, 16, 24)],
              heads=HEADS4, head_fill=S.BLUE, label_color=S.LABEL, box="record", style=style, cover=cover,
              cover_fill=S.BLUE, gap=3)
LOOK_FLAT = S.Look(font="CourierNew", size=5.6, head_fill=S.BLUE, label_color=S.LABEL, box="grid",
                   style=lambda r: style(r, (4, 4)),
                   cover=cover, cover_fill=S.BLUE)


def render_xlsx4(t: core.Table, path: Path) -> None:
    """The Excel export itself: cover block, report header, 4-row blocks with borders and fills. Sets row.pointer."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Anlagengitter"
    reg, bld, thin = Font(name="Courier New", size=7), Font(name="Courier New", size=7, bold=True), Side(style="thin")
    paint = {k: PatternFill("solid", fgColor="%02X%02X%02X" % tuple(round(v * 255) for v in rgb))
             for k, rgb in (("head", S.BLUE), ("sum", S.YELLOW), ("band", S.GREY), ("name", S.LILAC))}
    nf = "#,##0;#,##0-" if t.unit == "TEUR" else "#,##0.00;#,##0.00-"
    rr = 1
    for tab in cover(t):
        for k, line in enumerate(tab):
            for j, s in enumerate(line):
                x = ws.cell(rr, j + 1, int(s) if s.isdigit() else s or None)
                x.font, x.fill = (bld, paint["head"]) if k == 0 else (reg, PatternFill())
            rr += 1
        rr += 1
    for s in t.title_lines:
        ws.cell(rr, 1, s.strip()).font = bld
        rr += 1
    rr += 1
    for line in HEADS4:
        for j, s in enumerate(line):
            x = ws.cell(rr, j + 1, s or None)
            x.font, x.fill = bld, paint["head"]
        rr += 1
    for r in t.rows:
        if not r.truth:
            continue
        printed, cells = r.truth["printed"], {}
        band = getattr(r, "band", False)
        tone = paint["sum"] if r.truth["level"] != "class" else paint["band"] if band else None
        for li, line in enumerate(LINES4):
            for j, f in enumerate(line):
                s = r.cells[8 * li + j]
                if f in printed and s == printed[f]:
                    cents = abs(r.truth["fields"][f]) * (-1 if core._neg(s) else 1)
                    v = cents // 100_000 if t.unit == "TEUR" else cents / 100
                else:
                    v = 0 if f == "nil" else s or None
                x = ws.cell(rr + li, j + 1, v)
                x.font = bld if band else reg
                x.border = Border(left=thin if j == 0 else None, right=thin if j == 7 else None,
                                  top=thin if li == 0 else None, bottom=thin if li == 3 else None)
                if isinstance(v, (int, float)):
                    x.number_format = nf
                lilac = li == 0 and j in (5, 6) and r.truth["level"] == "class"
                if lilac or tone:
                    x.fill = paint["name"] if lilac else tone
                if f in printed:
                    cells[f] = x.coordinate
        r.pointer = {"sheet": ws.title, "cell": f"A{rr}", "cells": cells}
        rr += 5
    ws.column_dimensions["A"].width = 9
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 15
    wb.save(path)
    core._freeze_zip(path)


def build(out: Path) -> None:
    d = out / FAMILY
    d.mkdir(parents=True, exist_ok=True)

    def key(n, c, fmt, variant, t, planted=(), **kw):
        S.write_key(d, FAMILY, n, c, fmt, variant, t, planted, row_level="class", **kw)

    def via_print(t, look, then):  # render to a temp print, then scan / outline it
        with tempfile.TemporaryDirectory() as tmp:
            S.render(t, Path(tmp) / "print.pdf", look)
            then(Path(tmp) / "print.pdf")

    c = S.client(S.TRIO_SEED, year=2025, n_classes=13, assets_per_class=(4, 9))  # 01 + 02: the trio ledger
    t = table(c)
    S.render(t, d / f"{FAMILY}_01.pdf", LOOK)
    key(1, c, "pdf", "perfect", t, trio=True)
    t = table(c)
    render_xlsx4(t, d / f"{FAMILY}_02.xlsx")
    key(2, c, "xlsx", "perfect", t, trio=True)

    c = S.client(512, year=2025, fy_end=(9, 30), n_classes=10, fin=True)  # 03 fax scan
    t = table(c)
    via_print(t, LOOK, lambda p: core.scan(p, d / f"{FAMILY}_03.pdf", dpi=200, fax=True, seed=503))
    key(3, c, "scan", "perfect", t)

    c = S.client(513, year=2024, n_classes=13, fin=True)  # 04 clean, no text layer
    t = table(c)
    via_print(t, LOOK, lambda p: core.outline(p, d / f"{FAMILY}_04.pdf"))
    key(4, c, "outline", "perfect", t)

    c = S.client(514, year=2025, n_classes=12, fin=True)  # 05 right-most column cut at the page edge
    t, cut = table(c), ["ahk_close", "kum_afa_close", "bw_close"]
    for r in t.rows:
        r.flags.update(dict.fromkeys(cut, "cut_off"))
    S.render(t, d / f"{FAMILY}_05.pdf", replace(LOOK, cut=4))
    key(5, c, "pdf", "problem", t, [{
        "type": "column_cut_off", "where": [{"fields": cut, "pages": "all", "chars_lost": 4}],
        "expected_behaviour": "the right-most column runs past the page edge and loses its last characters "
                              "(cents, sign); rebuild it from the row laws: aktuelle AHK = AHK GJ-Beg + Zugang + Abgang "
                              "+ Umbuchung, kumulierte AfA = AfA GJ-Beg + AfA des Jahres + AfA Abgang + Zuschreibungen, "
                              "lfd Buchwert = aktuelle AHK + kumulierte AfA (as printed, with signs)"}])

    c, rng = S.client(515, year=2025, n_classes=11, fin=True), random.Random(515)  # 06 TEUR + a sum that does not tie
    t = table(c)
    P = [core.p_teur(t)[1]]
    S.finish(t)
    P.append(core.p_non_tying(t, rng)[1])
    S.finish(t)
    S.render(t, d / f"{FAMILY}_06.pdf", LOOK)
    key(6, c, "pdf", "problem", t, P)

    c, rng = S.client(516, year=2025, fy_end=(6, 30), n_classes=13), random.Random(516)  # 07 glued, header once, blank
    t = table(c)
    P = [core.p_glued(t, rng, n=3)[1]]
    S.finish(t)
    P += [core.p_no_repeat_header(t)[1], core.p_noise(t, c, "blank", "middle")[1]]
    S.render(t, d / f"{FAMILY}_07.pdf", LOOK)
    key(7, c, "pdf", "problem", t, P)

    c = S.client(517, year=2025, n_classes=13, fin=True)  # 08 one line per record, levels as words
    t = table(c, flat=True)
    S.render(t, d / f"{FAMILY}_08.pdf", LOOK_FLAT)
    key(8, c, "pdf", "layout-mutation", t)
