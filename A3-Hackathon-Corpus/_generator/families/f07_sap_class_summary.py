"""f07_sap_class_summary: SAP-like depreciation report, one row per Anlagenklasse prefixed with its 7-digit
Bilanzposition, movements as net columns (Bew AnschWert, Bew Wertber), boxed groups per Bilanzposition closed by a
'*<BilPos>' subtotal, no grand total; Courier text layer; selection page first; the last subtotal alone on the last page.

Class rows are the finest rows: kind 'asset', asset_no null, level 'class'. The two net columns have no core field:
their values are in each row's "net_fields" (bew_ahk = zugang - abgang + umbuchung, bew_wb = afa_abgang +
zuschreibung). Docs 01/02 are the trio with f05/f06 01/02 (_sap_common.TRIO).
"""

import random
import sys
import tempfile
from itertools import groupby
from pathlib import Path

import openpyxl

sys.path[:0] = [str(Path(__file__).resolve().parents[1]), str(Path(__file__).resolve().parent)]
import _sap_common as S  # noqa: E402
import core  # noqa: E402
from core import Cell  # noqa: E402

FAMILY = Path(__file__).stem
FIELDS = ["bilpos", "name", "klasse", "ahk_open", "bew_ahk", "kum_afa_open", "afa_gj", "bew_wb", "bw_close", "bw_prev",
          "curr"]
HEADS = ["BilPos", "Bezeichnung", "Klasse", "kum AnschWert", "Bew AnschWert", "kum AfA", "gep AfA", "Bew Wertber",
         "Ende-Buchwert", "Beg-Buchwert", "Währg"]
WIDTHS = [42, 107, 32] + [76] * 7 + [30]
MUT_FIELDS = ["klasse", "name", "bilpos", "ahk_open", "bew_ahk", "kum_afa_open", "afa_gj", "bew_wb", "bw_prev",
              "bw_close", "curr"]
MUT_HEADS = ["Klasse", "Klassenbezeichnung", "Bil.-Pos.", "AHK Beginn", "Bewegung AHK", "AfA kum. Beg.",
             "AfA geplant", "Bewegung WB", "Buchwert Beg.", "Buchwert Ende", "Whrg"]
MUT_WIDTHS = [34, 107, 48] + [76] * 7 + [28]
NET = ("bew_ahk", "bew_wb")
ROOM = 36  # class rows a page holds


def net(A: list) -> dict:
    return {"bew_ahk": sum(a.zugang - a.abgang + a.umbuchung for a in A),
            "bew_wb": sum(a.afa_abgang + a.zuschreibung for a in A)}


def table(c: core.Client, mut: bool = False) -> core.Table:
    fields = MUT_FIELDS if mut else FIELDS
    t = core.Table(fields=fields, col_widths=MUT_WIDTHS if mut else WIDTHS,
                   align=["r" if f in core.AMOUNTS or f in NET else "l" for f in fields],
                   header_rows=[[Cell(h) for h in (MUT_HEADS if mut else HEADS)]], font="Courier", font_size=8.3)
    S.titles(t, c, "Abschreibungen - Bereich 01 Handelsbilanz")
    t.neg = "paren" if mut else "trailing"
    for pos, group in groupby(S.ledger_classes(c), key=lambda k: k.position):
        A_pos = []
        for k in group:
            A = [a for a in c.assets if a.cls == k.code]
            A_pos += A
            name = k.name[:20]  # the report cuts class names at 20 characters
            t.add("asset", {"bilpos": S.POS7[pos], "name": name, "klasse": k.code, **S.sums(A, fields)},
                  {"year": c.year, "asset_no": None, "sub_no": None, "class": k.code, "account": k.account,
                   "position": pos, "name": name, "level": "class", "net_fields": net(A)})
        label = {"bilpos": S.POS7[pos], "name": f"Summe {S.POS7[pos]}"} if mut else {"bilpos": "*" + S.POS7[pos]}
        t.add("position_subtotal", {**label, **S.sums(A_pos, fields)},
              {"year": c.year, "position": pos, "net_fields": net(A_pos)})
    n = len(t.rows)
    t.rows_per_page = n - 1 if n - 1 <= ROOM and not mut else ROOM  # last subtotal alone on the last page
    return S.finish(t)


def cover(c: core.Client):
    def tabs(t: core.Table) -> list:
        return [[["Selektionsfeld", "Wert"], ["Buchungskreis", S.bukr(c)], ["Berichtsdatum", core.fmt_date(t.report_date)],
                 ["Bewertungsbereich 1", "01"], ["Sortiervariante", "ZB7"], ["", "X"], ["Anzeigevariante", "/AFA-KLASSE"],
                 ["ALV Grid Control benutzen", "X"]],
                [["Sortierkriterien", "Aufst.", "Abst.", "Zwischensumme"], ["Bilanzposition", "X", "", "X"],
                 ["Bestandskonto AHK", "X", "", ""], ["Anlagenklasse", "X", "", ""]],
                [["Datenstatistik", "Anzahl"], ["Übergebene Sätze", str(sum(r.kind == "asset" for r in t.rows))],
                 ["Ermittelte Summenzeilen", str(sum(r.kind == "position_subtotal" for r in t.rows))]]]
    return tabs


def look(c: core.Client, mut: bool = False) -> S.Look:
    if mut:  # every cell boxed, no selection page
        return S.Look(font="Courier", size=8.3, W=791, H=595, box="grid")
    return S.Look(font="Courier", size=8.3, W=791, H=595, box="group", gap=5, cover=cover(c),
                  group=lambda r: r.truth["position"] if r.kind == "asset" else id(r))


def build(out: Path) -> None:
    d = out / FAMILY
    d.mkdir(parents=True, exist_ok=True)

    def key(n, c, fmt, variant, t, planted=(), **kw):
        S.write_key(d, FAMILY, n, c, fmt, variant, t, planted, row_level="class",
                    net_fields="bew_ahk = Bew AnschWert (zugang - abgang + umbuchung), "
                               "bew_wb = Bew Wertber (afa_abgang + zuschreibung); signed as printed", **kw)

    def via_print(t, lk, then):
        with tempfile.TemporaryDirectory() as tmp:
            S.render(t, Path(tmp) / "print.pdf", lk)
            then(Path(tmp) / "print.pdf")

    c = S.client(S.TRIO_SEED, year=2025, n_classes=13, assets_per_class=(4, 9))  # 01 + 02: the trio ledger
    t = table(c)
    S.render(t, d / f"{FAMILY}_01.pdf", look(c))
    key(1, c, "pdf", "perfect", t, trio=True)
    t = table(c)
    path = d / f"{FAMILY}_02.xlsx"
    core.render_xlsx(t, path, sheet="Abschreibungen", number_format="#,##0.00;#,##0.00-",
                     extra_sheets={"Selektion": [row for tab in cover(c)(t) for row in [*tab, []]]})
    wb = openpyxl.load_workbook(path)  # the export holds the net columns as numbers too
    ws = wb["Abschreibungen"]
    for r in t.rows:
        for f in NET:
            x = ws.cell(ws[r.pointer["cell"]].row, FIELDS.index(f) + 1, r.truth["net_fields"][f] / 100)
            x.number_format = "#,##0.00;#,##0.00-"
    wb.save(path)
    core._freeze_zip(path)
    key(2, c, "xlsx", "perfect", t, trio=True)

    c = S.client(532, year=2025, n_classes=13, fin=True)  # 03 scan, heavily skewed
    t = table(c)
    via_print(t, look(c), lambda p: core.scan(p, d / f"{FAMILY}_03.pdf", skew_deg=2.5, noise=0.03, seed=703))
    key(3, c, "scan", "perfect", t)

    c = S.client(533, year=2025, fy_end=(9, 30), n_classes=12, fin=True)  # 04 clean, no text layer
    t = table(c)
    via_print(t, look(c), lambda p: core.outline(p, d / f"{FAMILY}_04.pdf"))
    key(4, c, "outline", "perfect", t)

    c, rng = S.client(534, year=2025, n_classes=13, fin=True), random.Random(534)  # 05 two report runs, one sum off
    ts = [table(y) for y in c.years(2)]
    P = [core.p_non_tying(ts[1], rng)[1]]
    S.finish(ts[1])
    ts, rec = core.p_multi_year(ts)
    S.render(ts, d / f"{FAMILY}_05.pdf", look(c))
    key(5, c, "pdf", "problem", ts, [rec, *P], years=[c.year - 1, c.year])

    c, rng = S.client(535, year=2025, n_classes=11, fin=True), random.Random(535)  # 06 TEUR, glued, cover letter
    t = table(c)
    P = [core.p_teur(t)[1]]
    S.finish(t)
    P.append(core.p_glued(t, rng)[1])
    S.finish(t)
    P.append(core.p_noise(t, c, "cover", "start")[1])
    S.render(t, d / f"{FAMILY}_06.pdf", look(c))
    key(6, c, "pdf", "problem", t, P)

    c = S.client(536, year=2024, n_classes=13, fin=True)  # 07 renamed, reordered, (negatives), sums in words
    t = table(c, mut=True)
    S.render(t, d / f"{FAMILY}_07.pdf", look(c, mut=True))
    key(7, c, "pdf", "layout-mutation", t)
