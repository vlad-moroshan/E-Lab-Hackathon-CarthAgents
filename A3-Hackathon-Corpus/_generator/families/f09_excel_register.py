"""f09_excel_register: SAP asset history sheet exported to Excel (ALV), edited by an auditor. Look-alike of
_intern/profiles/Book2.md. Sheet1: one row per asset with the full Gitter columns (21), asset rows inside a collapsed
outline (level 2, hidden), the class only in the 2-digit prefix of the 9-digit text asset id, SAP signs (Abgang and
all AfA columns negative, AfA Abgang positive), zero cells "#,##0", others "#,##0.00", an auditor's SUBTOTAL row
above the table, two yellow value-total rows below it, check formulas and a reconciliation with hard-coded
constants. Sheet2: Gitter by asset class with two auditor columns (SUM, difference).

Key: fields are positive cents as everywhere (signs are layout); `printed` holds the text as shown (-1.234,56).
The yellow rows and the SUBTOTAL row are all 'grand_total' rows; Sheet2 class rows are 'class_subtotal' rows
that sum the Sheet1 assets of their class. A cell whose shown value differs from the truth is flagged per field.
xlsx: formulas are calculated by LibreOffice (soffice --convert-to xlsx), which saves the cached values.
"""

import subprocess
import sys
import tempfile
from dataclasses import replace
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.pdfbase.pdfmetrics import stringWidth

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import core  # noqa: E402
from core import Cell  # noqa: E402

FAMILY = Path(__file__).stem
SOFFICE = "/opt/homebrew/bin/soffice"
COLS1 = [("asset_no", "Anlage"), ("sub_no", "Unter-nummer"), ("acq_date", "Aktivierung am"),
         ("name", "Anlagenbezeichnung"), ("cur", "Währ-ung"), ("ahk_open", "AHK GJ-Beg"), ("zugang", "Zugang"),
         ("abgang", "Abgang"), ("umbuchung", "Umbuchung"), ("z_nachakt", "Nachaktivg"), ("z_invest", "Invest.Förderg"),
         ("ahk_close", "aktuelle AHK"), ("kum_afa_open", "AfA GJ-Beg"), ("afa_gj", "AfA des Jahres"),
         ("afa_abgang", "AfA Abgang"), ("z_afa_umb", "AfA Umbuchung"), ("z_afa_nachakt", "AfA Nachaktivg"),
         ("zuschreibung", "Zuschreibungen"), ("kum_afa_close", "kumulierte AfA"), ("bw_prev", "Buchwrt GJ-Beg"),
         ("bw_close", "lfd Buchwert")]
COLS2 = [("bukrs", "Buchungskreis"), ("cls", "Anlagenklasse"), ("name", "Bezeichnung"), *COLS1[4:11],
         ("x_sum", "SUMME AHK / Zugang / Abgang / Umbuchung"), ("ahk_close", "aktuelle AHK"),
         ("x_diff", "Differenz zu aktuelle AHK"), *COLS1[12:]]
COLS8 = [("asset_no", "Anlagennummer"), ("sub_no", "UNr."), ("name", "Bezeichnung"), ("acq_date", "Aktiviert"),
         ("bw_close", "Restbuchwert Ende"), ("bw_prev", "Restbuchwert Beginn"), ("ahk_open", "Anschaffungswert Beginn"),
         ("zugang", "Zugänge"), ("abgang", "Abgänge"), ("umbuchung", "Umbuchungen"), ("ahk_close", "Anschaffungswert Ende"),
         ("kum_afa_open", "Abschreibung kum. Beginn"), ("afa_gj", "Abschreibung GJ"), ("afa_abgang", "Abschreibung Abgänge"),
         ("zuschreibung", "Zuschreibungen"), ("kum_afa_close", "Abschreibung kum. Ende")]
NEG = {"abgang", "kum_afa_open", "afa_gj", "kum_afa_close"}  # SAP: shown negative
ACC = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
NF8 = '#,##0.00;-#,##0.00;""'
GREY, YELLOW1, YELLOW2, AUDIT = "DDDDDD", "FFFB8B", "FFF844", "E2EFDA"


def font(**kw) -> Font:
    return Font(name="Aptos Narrow", size=11, **kw)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def client(seed: int, year: int = 2025, fy_end: tuple = (6, 30)) -> core.Client:
    return core.make_client(seed, year=year, fy_end=fy_end, n_classes=13, assets_per_class=(2, 5), subnumbers=True)


def bukrs(c: core.Client) -> str:
    return str(1000 + c.seed * 7 % 9000)


def text(v: int, style: str) -> str:
    """Shown text of a signed amount: 'sap' number cells, 'acc' SUBTOTAL row, 'list' SAP text export, 'plain'."""
    if v == 0:
        return {"sap": "0", "acc": "-", "list": "0,00", "plain": ""}[style]
    return core.fmt(v, "de_trailing_minus" if style == "list" else "de")


def heads(cols: list) -> list:
    """PDF heading: a long SAP column text wraps onto a second line at the space or hyphen nearest its middle."""
    rows = [[], []]
    for _, h in cols:
        cut = [i for i, ch in enumerate(h) if ch in " -"] if len(h) > 10 else []
        i = min(cut, key=lambda i: abs(i - len(h) / 2)) if cut else len(h)
        rows[0].append(Cell(h[:i + 1].strip()))
        rows[1].append(Cell(h[i + 1:].strip()))
    return rows


def show(t: core.Table, r, style: str, sap: bool = True) -> None:
    """Signed shown values r.val per value column (truth, or r.shown where a problem changed a cell) and r.cells."""
    shown, r.val = getattr(r, "shown", {}), {}
    for j, f in enumerate(t.fields):
        if f in core.AMOUNTS:
            v = shown.get(f, r.truth["fields"][f]) * (-1 if sap and f in NEG else 1)
        elif f.startswith("z_"):
            v = 0
        elif f == "x_sum":
            v = sum(r.val[k] for k in ("ahk_open", "zugang", "abgang", "umbuchung"))
        elif f == "x_diff":
            v = r.val["ahk_close"] - r.val["x_sum"]
        else:
            continue
        r.val[f], r.cells[j] = v, text(v, getattr(r, "style", style))
        if f in r.truth["printed"]:
            r.truth["printed"][f] = r.cells[j]


def total(t: core.Table, kind: str, label: str, kids: list, **ids):
    """A sum row over asset rows; the sum of what the kids show, flagged where that is not the truth."""
    r = t.add_total(kind, label, [k.asset for k in kids], **ids)
    r.cells[t.fields.index("name")], r.cells[0] = "", label
    for f in r.truth["fields"]:
        s = sum(getattr(k, "shown", {}).get(f, k.truth["fields"][f]) for k in kids)
        if s != r.truth["fields"][f]:
            r.shown = {**getattr(r, "shown", {}), f: s}
            r.flags[f] = next(k.flags[f] for k in kids if f in k.flags)
    return r


def tables(c: core.Client, *, summen: bool = True, style: str = "sap", transfer_gap: bool = False):
    """Sheet1 (assets) and Sheet2 (classes). transfer_gap: one receiving asset shows Umbuchung 0 (column incomplete)."""
    bk = bukrs(c)
    t1 = core.Table(fields=[f for f, _ in COLS1], col_widths=[10] * len(COLS1), align=["l"] * 5 + ["r"] * 16,
                    header_rows=heads(COLS1), zero=None, font="Narrow", font_size=6,
                    title_lines=[f"Jahresabschluss per {core.fmt_date(c.fy_end)}", f"BK: {bk}"])
    assets = sorted((replace(a, asset_no=f"{a.cls[:2]}{int(a.asset_no) % 1000 + 1000:07d}") for a in c.assets),
                    key=lambda a: (a.asset_no, int(a.sub_no)))
    rows = []
    for a in assets:
        r = t1.add_asset(a)
        r.asset, r.cells[t1.fields.index("cur")] = a, "EUR"
        rows.append(r)
    planted = None
    if transfer_gap:
        r = max((r for r in rows if r.truth["fields"]["umbuchung"] > 0), key=lambda r: r.truth["fields"]["umbuchung"])
        r.shown, r.flags["umbuchung"] = {"umbuchung": 0}, "transfer_incomplete"
        planted = {"type": "transfer_breaks_cost_law", "where": [{"row": r, "field": "umbuchung", "printed": 0,
                                                                  "implied": r.truth["fields"]["umbuchung"]}],
                   "expected_behaviour": "the Umbuchung cell is empty in the export although 'aktuelle AHK' includes "
                                         "the transfer, so the row does not roll forward; the key holds the implied "
                                         "transfer (aktuelle AHK - AHK GJ-Beg - Zugang + Abgang); the column totals "
                                         "and the class tab miss it too (flagged)"}
    if summen:
        s = total(t1, "grand_total", "SUMMEN", rows)
        t1.rows.insert(0, t1.rows.pop())  # above the table
        s.style = "acc"
    total(t1, "grand_total", f"Buchungskreis {bk}", rows)
    total(t1, "grand_total", "", rows)
    for r in t1.rows:
        show(t1, r, style)
    t2 = core.Table(fields=[f for f, _ in COLS2], col_widths=[10] * len(COLS2), align=["l"] * 4 + ["r"] * 18,
                    header_rows=heads(COLS2), zero=None, font="Narrow", font_size=6)
    classes = []
    for k in sorted(c.classes, key=lambda k: k.code):
        if kids := [r for r in rows if r.truth["class"] == k.code]:
            r = total(t2, "class_subtotal", "", kids, cls=k.code)
            r.cells[:4] = [bk, k.code, k.name if len(k.name) <= 20 else "GWG", "EUR"]
            classes.append(r)
    for label in (f"Buchungskreis {bk}", ""):
        r = total(t2, "grand_total", label, rows)
        r.cells[3] = "EUR"
    for r in t2.rows:
        show(t2, r, style)
    return t1, t2, planted


def fit(t: core.Table, room: float = 792) -> core.Table:
    """PDF column widths from the widest text per column; smaller font until the table fits the landscape page."""
    while True:
        bold = core._bold(t.font)
        w = [max(stringWidth(s, bold, t.font_size) for s in col)
             for col in zip(*[r.cells for r in t.rows], *[[c.text for c in hr] for hr in t.header_rows])]
        t.col_widths = [x + 4 for x in w]
        if sum(t.col_widths) <= room or t.font_size <= 4.5:
            return t
        t.font_size -= 0.25


def put(ws, t: core.Table, r, x: int, *, as_text=False, formula=None, nf=None, bold=False, color=None) -> None:
    """One table row at Excel row x; sets the key pointer."""
    cells = {}
    for j, f in enumerate(t.fields):
        y = ws.cell(x, j + 1)
        fx = formula(f, get_column_letter(j + 1)) if formula and f in r.val else None
        if fx:
            y.value, y.number_format = fx, ACC if r.cells[0] == "SUMMEN" else (nf or "#,##0.00")
        elif f in r.val and not as_text:
            y.value, y.number_format = r.val[f] / 100, (nf or ("#,##0.00" if r.val[f] else "#,##0"))
        elif f == "acq_date" and r.kind == "asset" and not as_text:
            d = r.truth["acq_date"]
            y.value, y.number_format = datetime(d.year, d.month, d.day), "DD.MM.YYYY"
        else:
            y.value = r.cells[j] or None
        y.font = font(bold=bold or f in ("x_sum", "x_diff"))
        if color or f in ("x_sum", "x_diff"):
            y.fill = fill(color or AUDIT)
        if f in r.truth["printed"]:
            cells[f] = y.coordinate
    r.pointer = {"sheet": ws.title, "cell": f"A{x}", "cells": cells}


def header(ws, x: int, cols: list, widths: dict) -> None:
    for j, (f, h) in enumerate(cols, 1):
        y = ws.cell(x, j, h)
        y.fill, y.font = fill(AUDIT if f.startswith("x_") else GREY), font(bold=f.startswith("x_"))
        y.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(j)].width = widths.get(f, 15)
    ws.row_dimensions[x].height = 30


def sheet1(ws, t: core.Table, c: core.Client, *, as_text=False, outline=True, recon=True) -> int:
    """Rows 1-12 header block (SUBTOTAL row 9), table heading row 13, assets, two yellow rows, checks, notes."""
    ws.cell(4, 1, t.title_lines[0]).font = font(bold=True)
    ws.cell(6, 1, t.title_lines[1]).font = font()
    header(ws, 13, COLS1, {"asset_no": 12, "sub_no": 8, "acq_date": 11, "name": 40, "cur": 6})
    summen = t.rows[0] if t.rows[0].cells[0] == "SUMMEN" else None
    assets, (bk, grand) = [r for r in t.rows if r.kind == "asset"], t.rows[-2:]
    first, last = 14, 13 + len(assets)
    if summen:
        put(ws, t, summen, 9, bold=True, formula=lambda f, L: f"=SUBTOTAL(9,{L}{first}:{L}{last})")
    for x, r in enumerate(assets, first):
        put(ws, t, r, x, as_text=as_text)
        if outline:
            ws.row_dimensions[x].outlineLevel, ws.row_dimensions[x].hidden = 2, True
    put(ws, t, bk, last + 1, as_text=as_text, color=YELLOW1)
    put(ws, t, grand, last + 2, as_text=as_text, color=YELLOW2)
    if not recon:
        return last + 3
    for j in range(6, len(COLS1) + 1):  # auditor's check: bottom total minus SUBTOTAL row
        L = get_column_letter(j)
        ws.cell(last + 4, j, f"=+{L}{last + 2}-{L}9").number_format = "#,##0.00"
    x, U = last + 6, get_column_letter(len(COLS1))
    ws.cell(x, 1, "Abstimmung Bilanz").font = font(bold=True)
    lines = [("lfd Buchwert lt. Anlagengitter", f"=+{U}{last + 2}")]
    for g, name in (("A.I", "Immaterielle VG"), ("A.II", "Sachanlagen"), ("A.III", "Finanzanlagen")):
        parts = [v for p in core.POSITIONS if p.startswith(g + ".")
                 if (v := sum(a.bw_close for a in c.assets if a.position == p))]
        if parts:
            lines.append((f"{name} lt. Bilanz", "=" + "+".join(f"{v / 100:.2f}" for v in parts)))
    lines.append(("Differenz", f"={U}{x}-SUM({U}{x + 1}:{U}{x + len(lines) - 1})"))
    for i, (label, fx) in enumerate(lines):
        ws.cell(x + i, 4, label).font = font()
        ws.cell(x + i, len(COLS1), fx).number_format = "#,##0.00"
    return x + len(lines)


def sheet2(ws, t: core.Table) -> None:
    header(ws, 1, COLS2, {"bukrs": 14, "cls": 12, "name": 21, "cur": 8})
    classes, (bk, grand) = t.rows[:-2], t.rows[-2:]
    n = 1 + len(classes)
    for x, r in enumerate(classes, 2):
        put(ws, t, r, x, formula=lambda f, L, x=x: {"x_sum": f"=SUM(E{x}:H{x})", "x_diff": f"=L{x}-K{x}"}.get(f))
        ws.row_dimensions[x].outlineLevel = 1
    for x, r, color in ((n + 1, bk, YELLOW1), (n + 2, grand, YELLOW2)):
        put(ws, t, r, x, color=color, formula=lambda f, L: f"=SUM({L}2:{L}{n})" if f.startswith("x_") else None)


def forecast(ws, c: core.Client, x: int) -> dict:
    """A second, unrelated table under the reconciliation: next year's planned AfA per class (ALV look)."""
    ws.cell(x, 1, f"AfA-Vorschau GJ {c.year + 1}").font = font(bold=True)
    for j, h in ((1, "Anlagenklasse"), (4, "Bezeichnung"), (6, "Buchwert"), (7, "Plan-AfA"), (8, "Restbuchwert")):
        y = ws.cell(x + 1, j, h)
        y.fill, y.font = fill(GREY), font()
    y = x + 2
    for k, assets in sorted(c.by_class(), key=lambda p: p[0].code):
        bw = sum(a.bw_close for a in assets)
        plan = sum(min(a.bw_close, a.afa_gj) for a in assets if a.disposal_date is None)
        ws.cell(y, 1, k.code), ws.cell(y, 4, k.name)
        for j, v in ((6, bw), (7, -plan), (8, bw - plan)):
            ws.cell(y, j, v / 100).number_format = "#,##0.00" if v else "#,##0"
        y += 1
    return {"sheet": "Sheet1", "range": f"A{x}:H{y - 1}"}


def mutation(c: core.Client) -> core.Table:
    """Same register, other system settings: one tab, columns reordered and renamed, positive signs, rows visible
    with an ALV subtotal per class, zero cells blank by number format."""
    t = core.Table(fields=[f for f, _ in COLS8], col_widths=[10] * len(COLS8), align=["l"] * 4 + ["r"] * 12,
                   header_rows=[[Cell(h) for _, h in COLS8]], zero="",
                   title_lines=[c.company, f"Anlagenliste nach Anlagenklassen zum {core.fmt_date(c.fy_end)}"])
    every = []
    for k, assets in sorted(c.by_class(), key=lambda p: p[0].code):
        kids = []
        for a in assets:
            r = t.add_asset(replace(a, asset_no=f"{a.cls[:2]}{int(a.asset_no) % 1000 + 1000:07d}"))
            r.asset = a
            kids.append(r)
        total(t, "class_subtotal", f"Summe Anlagenklasse {k.code} {k.name}", kids, cls=k.code)
        every += kids
    total(t, "grand_total", "Gesamtsumme", every)
    for r in t.rows:
        show(t, r, "plain", sap=False)
    return t


def recalc(src: Path, dst: Path) -> None:
    """LibreOffice calculates every formula and saves the workbook with cached values; fixed zip dates after."""
    files = sorted(src.glob("*.xlsx"))
    subprocess.run([SOFFICE, "--headless", f"-env:UserInstallation=file://{src}/lo", "--convert-to", "xlsx",
                    "--outdir", str(dst), *map(str, files)], check=True, capture_output=True, timeout=600)
    for f in files:
        assert (dst / f.name).exists(), f"soffice did not convert {f.name}"
        core._freeze_zip(dst / f.name)


def build(out: Path) -> None:
    d = out / FAMILY
    d.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        def key(n, c, fmt, variant, ts, planted=()):
            meta = core.key_meta(c, doc_id=f"{FAMILY}_{n:02d}", family=FAMILY, doc_type="Anlagenverzeichnis", format=fmt,
                                 variant=variant, company_code=bukrs(c))
            core.write_key(d / f"{FAMILY}_{n:02d}.key.json", meta, ts, planted)

        def book(c, n, variant, *, extra=False, **kw):
            t1, t2, planted = tables(c, **kw)
            wb = Workbook()
            ws1, ws2 = wb.active, wb.create_sheet("Sheet2")
            ws1.title = "Sheet1"
            end = sheet1(ws1, t1, c)
            sheet2(ws2, t2)
            planted = [planted]
            if extra:
                planted.append({"type": "unrelated_table", "where": [forecast(ws1, c, end + 3)],
                                "expected_behaviour": "the table under the reconciliation is a forecast of next "
                                                      "year's AfA per class, not part of the register; its columns "
                                                      "F-H sit under 'AHK GJ-Beg / Zugang / Abgang' by accident"})
            wb.save(tmp / f"{FAMILY}_{n:02d}.xlsx")
            key(n, c, "xlsx", variant, [t1, t2], planted)

        def pdf(c, path):
            t1, t2, _ = tables(c)
            ts = [fit(t1), fit(t2)]
            core.render_pdf(ts, path)
            return ts

        c = client(901)  # 01 perfect workbook
        book(c, 1, "perfect")

        c = client(902, fy_end=(9, 30))  # 02 printed with the outline expanded
        key(2, c, "pdf", "perfect", pdf(c, d / f"{FAMILY}_02.pdf"))

        c = client(903, fy_end=(12, 31))  # 03 scan of such a print
        ts = pdf(c, tmp / "print03.pdf")
        core.scan(tmp / "print03.pdf", d / f"{FAMILY}_03.pdf", dpi=200, seed=903)
        key(3, c, "scan", "perfect", ts)

        c = client(904, year=2024)  # 04 clean print, glyphs as outlines (no text layer)
        ts = pdf(c, tmp / "print04.pdf")
        core.outline(tmp / "print04.pdf", d / f"{FAMILY}_04.pdf")
        key(4, c, "outline", "perfect", ts)

        c = client(905)  # 05 transfer column incomplete: one row does not roll forward
        book(c, 5, "problem", transfer_gap=True)

        c = client(906, fy_end=(12, 31))  # 06 raw SAP list export: every value as German text, trailing minus
        t1, _, _ = tables(c, summen=False, style="list")
        wb = Workbook()
        wb.active.title = "Sheet1"
        sheet1(wb.active, t1, c, as_text=True, outline=False, recon=False)
        wb.save(tmp / f"{FAMILY}_06.xlsx")
        key(6, c, "xlsx", "problem", [t1], [{
            "type": "numbers_as_text", "where": [{"sheet": "Sheet1", "cells": "every amount and date"}],
            "expected_behaviour": "amounts are German-format text with SAP trailing minus (1.234,56-), dates are "
                                  "text; parse them; the sign is layout, the key holds magnitudes"}])

        c = client(907)  # 07 a second, unrelated table on the same sheet
        book(c, 7, "problem", extra=True)

        c = client(908, fy_end=(3, 31))  # 08 layout mutation
        t = mutation(c)
        wb = Workbook()
        ws = wb.active
        ws.title = "Anlagen"
        for i, s in enumerate(t.title_lines, 1):
            ws.cell(i, 1, s).font = font(bold=i == 1)
        header(ws, 4, COLS8, {"asset_no": 14, "sub_no": 6, "name": 40, "acq_date": 11})
        for x, r in enumerate(t.rows, 5):
            put(ws, t, r, x, nf=NF8, bold=r.kind != "asset")
            if r.kind == "asset":
                ws.row_dimensions[x].outlineLevel = 1
        wb.save(tmp / f"{FAMILY}_08.xlsx")
        key(8, c, "xlsx", "layout-mutation", t)

        recalc(tmp, d)
