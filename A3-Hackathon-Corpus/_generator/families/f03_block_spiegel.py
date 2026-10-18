"""f03_block_spiegel: DATEV-like print "Entwicklung des Anlagevermögens" (look-alike of profile dsfdvfdvd).

Part 1 (summary): one 3-line block per account (Ansch-/Herst-K / Abschreibung / Buchwerte), unlabeled grand total.
Part 2 (detail): one 3-line block per asset (AHK / Absch / BW) with date, AfA-Art and "ND/MM rate" in their own
column, an account sum block (name only, thin rules) and the grand total again. Each account starts a new page; its
heading line repeats on every page. Negative movements (Abgang, Zuschreibung, outgoing Umbuchung) go on a second
line under their column with a trailing minus. "Blatt n" top right, "Handelsrecht" bottom right. The BW line
repeats the Zugang, Umbuchung and AfA figures. Contract: see f00_demo.
"""

import random
import sys
import tempfile
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import core  # noqa: E402

FAMILY = Path(__file__).stem
# One Table field per amount. The None gaps leave p_glued only the two pairs that sit side by side on one line.
FIELDS = ["name", "ahk_open", "zugang", None, "abgang", None, "umbuchung", None, "ahk_close", None,
          "kum_afa_open", "afa_gj", None, "afa_abgang", None, "zuschreibung", None, "kum_afa_close", None,
          "bw_prev", None, "abgang_bw", None, "bw_close"]
# (field, block line 0 AHK / 1 Abschreibung / 2 Buchwert, amount column 0 opening .. 4 closing, printed again)
SLOTS = [("ahk_open", 0, 0, 0), ("zugang", 0, 1, 0), ("abgang", 0, 1, 0), ("umbuchung", 0, 2, 0),
         ("ahk_close", 0, 4, 0), ("kum_afa_open", 1, 0, 0), ("afa_gj", 1, 1, 0), ("afa_abgang", 1, 1, 0),
         ("zuschreibung", 1, 3, 0), ("kum_afa_close", 1, 4, 0), ("bw_prev", 2, 0, 0), ("zugang", 2, 1, 1),
         ("abgang_bw", 2, 1, 0), ("umbuchung", 2, 2, 1), ("afa_gj", 2, 3, 1), ("zuschreibung", 2, 3, 1),
         ("bw_close", 2, 4, 0)]
NEG = {"abgang", "afa_abgang", "zuschreibung", "abgang_bw"}  # positive in the key, printed as negatives
BALANCES = {"ahk_open", "ahk_close", "kum_afa_open", "kum_afa_close", "bw_prev", "bw_close"}  # zero prints 0,00
LAND = dict(size=landscape(A4), font=7.4, pitch=10.2, left=34, right=805, id_x=40, name_x=72, info_x=290,
            label_r=379, amt_r=(463, 548, 630, 714, 800), amt_w=72, bottom=58, title="Entwicklung des Anlagevermögens",
            labels={"summary": ("Ansch-/Herst-K", "Abschreibung", "Buchwerte"), "detail": ("AHK", "Absch", "BW")})
PORT = dict(LAND, size=A4, font=7.0, pitch=9.6, left=24, right=571, id_x=26, name_x=56, info_x=178, label_r=262,
            amt_r=(323, 385, 447, 509, 571), amt_w=58, title="Anlagenspiegel",
            labels={"summary": ("AK/HK", "kum. AfA", "Restbuchwert"), "detail": ("AK/HK", "kum. AfA", "RBW")})


def isneg(s: str) -> bool:
    return s.startswith(("-", "(")) or s.endswith("-")


def heads(part: str, d0: str, d1: str, mut: bool) -> list:
    det = part == "detail"
    if mut:
        rows = [["Anlagegut", "Zugangsdatum" if det else "", "Wertart", "Anfangsstand", "Zugänge", "Umbuchungen",
                 "AfA lfd. Jahr", "Endstand"],
                ["", "AfA-Methode" if det else "", "", d0, "Abgänge (-)", "", "Zuschreib. (-)", d1],
                ["", "ND / Satz %" if det else "", "", *["EUR"] * 5]]
    else:
        rows = [["Bezeichnung", "Datum" if det else "", "Entw." if det else "Entwicklung", "Stand zum", "Zugang", "",
                 "Abschreibung", "Stand zum"],
                ["", "AfA-Art" if det else "", "der", d0, "Abgang-", "Umbuchung", "Zuschreibung-", d1],
                ["", "ND  %" if det else "", "", *["EUR"] * 5]]
    return [[core.Cell(s) for s in row] for row in rows]


def tidy(t: core.Table, r: core.Row) -> core.Row:
    """Zero movements print blank, zero balances 0,00."""
    for f, v in r.truth["fields"].items():
        if v == 0 and f not in BALANCES:
            r.cells[t.fields.index(f)] = r.truth["printed"][f] = ""
    return r


def tables(c: core.Client, town: str, L: dict) -> list:
    """[summary, detail] for client c in layout L (PORT = the mutation: synonyms, leading minus, portrait)."""
    mut = L is PORT
    d0, d1, out = core.fmt_date(c.fy_start), core.fmt_date(c.fy_end), []
    for part in ("summary", "detail"):
        t = core.Table(fields=FIELDS, col_widths=[0] * len(FIELDS), align=["r"] * len(FIELDS),
                       header_rows=heads(part, d0, d1, mut), num_style="de" if mut else "de_trailing_minus",
                       title_lines=[L["title"], f"vom {d0} bis {d1}", c.company, town], font_size=L["font"],
                       landscape=not mut, page_header="Blatt {p}", page_footer="Handelsrecht")
        t.part = part
        for k, assets in c.by_class():
            acc = assets[0].account
            if part == "detail":
                t.add("heading", {"name": f"{acc}  {k.name}"})
                for a in assets:
                    tidy(t, t.add_asset(a))
            tidy(t, t.add_total("account_subtotal", f"{acc}  {k.name}" if part == "summary" else k.name, assets,
                                cls=k.code, account=acc))
        tidy(t, t.add_total("grand_total", "", c.assets))
        out.append(t)
    return out


def afa_word(tr: dict) -> str:
    if tr["method"] in ("GWG", "keine"):
        return {"GWG": "GWG-Sofort", "keine": "Keine AfA"}[tr["method"]]
    return {"A.I.2": "Immat.WG", "A.II.1": f"Lin.Geb.{round(1200 / tr['nd_months'])}"}.get(tr["position"], "Linear")


def wrap(text: str, width: float, size: float, n: int = 3) -> list:
    lines, cur = [], ""
    for w in text.split(" "):
        if cur and len(lines) < n - 1 and stringWidth(f"{cur} {w}", "Helvetica", size) > width:
            lines, cur = lines + [cur], w
        else:
            cur = f"{cur} {w}" if cur else w
    return lines + [cur]


def _sign(t: core.Table) -> None:
    """NEG amounts print as negatives (trailing minus, or leading in the 'de' style); idempotent."""
    neg = (lambda s: s + "-") if t.num_style == "de_trailing_minus" else (lambda s: "-" + s)
    for r in t.rows:
        p = (r.truth or {}).get("printed", {})
        for f in NEG & p.keys():
            i = t.fields.index(f)
            if r.cells[i] and not isneg(r.cells[i]) and f not in r.flags:
                r.cells[i] = neg(r.cells[i])
            if p[f] and not isneg(p[f]):
                p[f] = neg(p[f])


def lines_of(t: core.Table, r: core.Row, L: dict, wrap_names: bool = True) -> list:
    """Printed lines of one row: dicts of id, bez, info, label, amt {col: (text, field, primary)}, rule, under."""
    if r.kind in ("heading", "continuation"):
        return [{"id" if r.kind == "heading" else "bez": r.cells[0], "amt": {}}]
    lines, mains = [], []
    for ln in range(3):
        main, neg = {"label": L["labels"][t.part][ln], "amt": {}}, {"amt": {}}
        for f, line, col, again in SLOTS:
            if line == ln and f in r.truth["fields"]:
                s = r.truth["printed"][f] if again else r.cells[t.fields.index(f)]
                if s:
                    (neg if f in NEG or isneg(s) else main)["amt"][col] = (s, f, not again)
        mains.append(len(lines))
        lines += [main] + ([neg] if neg["amt"] else [])
    tr = r.truth
    if r.kind == "asset":
        nd = tr["nd_months"]
        bez = wrap(r.cells[0], L["info_x"] - L["name_x"] - 6, L["font"]) if wrap_names else [r.cells[0]]
        info = [core.fmt_date(tr["acq_date"]), afa_word(tr),
                f"{nd // 12}/{nd % 12:02d}  {core.fmt(round(120_000 / nd))}" if nd and tr["method"] != "keine" else ""]
        for k, i in enumerate(mains):
            lines[i].update(bez=bez[k] if k < len(bez) else "", info=info[k])
        lines[0]["id"] = tr["asset_no"]
    else:
        lines[0]["id"] = r.cells[0]
    if r.kind == "account_subtotal" and t.part == "detail":
        for i in mains:
            lines[i]["rule"] = 0.4
    if r.kind == "grand_total":
        lines[0]["rule"], lines[-1]["under"] = 1.1, True
    return lines


def _line(cv, L: dict, ln: dict, y: float) -> None:
    cv.setFont("Helvetica", L["font"])
    for key, x in (("id", L["id_x"]), ("bez", L["name_x"]), ("info", L["info_x"])):
        if ln.get(key):
            cv.drawString(x, y, ln[key])
    if ln.get("label"):
        cv.drawRightString(L["label_r"], y, ln["label"])
    for col, (s, *_) in ln["amt"].items():
        cv.drawRightString(L["amt_r"][col], y, s)
    x0, x1 = L["amt_r"][0] - L["amt_w"], L["amt_r"][-1]
    if ln.get("rule"):
        cv.setLineWidth(ln["rule"])
        cv.line(x0, y + L["font"] + 0.8, x1, y + L["font"] + 0.8)
    if ln.get("under"):
        cv.setLineWidth(0.5)
        cv.line(x0, y - 2.6, x1, y - 2.6)
        cv.line(x0, y - 4.4, x1, y - 4.4)


def _head(cv, t: core.Table, L: dict) -> float:
    """Title, company, location, 3 heading rows and the full-width rule; returns the first content baseline."""
    W, H = L["size"]
    t1, t2 = t.title_lines[0], " " + t.title_lines[1]
    w1 = stringWidth(t1, "Helvetica", 11)
    x = (W - w1 - stringWidth(t2, "Helvetica", 8.5)) / 2
    cv.setFont("Helvetica", 11)
    cv.drawString(x, H - 58, t1)
    cv.setFont("Helvetica", 8.5)
    cv.drawString(x + w1, H - 58, t2)
    cv.drawCentredString(W / 2, H - 85, t.title_lines[2])
    cv.setFont("Helvetica-Bold", 8.5)
    cv.drawCentredString(W / 2, H - 116, t.title_lines[3])
    cv.setFont("Helvetica", L["font"])
    for i, hr in enumerate(t.header_rows):
        y = H - 134 - i * 9
        for col, cell in enumerate(hr):
            if cell.text and col < 2:
                cv.drawString((L["left"], L["info_x"])[col], y, cell.text)
            elif cell.text:
                cv.drawRightString(L["label_r"] if col == 2 else L["amt_r"][col - 3], y, cell.text)
    cv.setLineWidth(0.6)
    cv.line(L["left"], H - 168, L["right"], H - 168)
    return H - 188


def _noise(cv, r: core.Row, st: dict) -> None:
    """A page that is not part of the print (cover letter, blank, unrelated list), portrait like a letter."""
    W, H = A4
    st["page"] += 1
    cv.setPageSize((W, H))
    cv.setFont("Helvetica", 10)
    for i, s in enumerate(r.cells):
        left, _, right = s.partition("\t")
        cv.drawString(70, H - 90 - i * 14, left)
        cv.drawRightString(480, H - 90 - i * 14, right)
    r.pointer = {"page": st["page"], "line": 1, "bbox": [0, 0, round(W, 1), round(H, 1)]}
    cv.showPage()


def render_pdf(tabs: list, path: Path, L: dict) -> None:
    """Text PDF of [summary, detail]; blocks never split; sets each row's pointer {page, line, bbox}."""
    W, H = L["size"]
    lh, fs = L["pitch"], L["font"]
    cv = canvas.Canvas(str(path), invariant=1, pageCompression=1)
    st = {"page": 0, "blatt": 0, "y": None}

    def close():
        if st["y"] is not None:
            cv.showPage()
            st["y"] = None

    def open_page(t, head, heading):
        close()
        st["page"] += 1
        st["blatt"] += 1
        cv.setPageSize((W, H))
        cv.setFont("Helvetica", fs)
        cv.drawRightString(L["right"], H - 36, t.page_header.format(p=st["blatt"]))
        cv.setFont("Helvetica", 8)
        cv.drawRightString(L["right"], 37, t.page_footer)
        y = _head(cv, t, L) if head else H - 62
        if heading:  # the account heading repeats on every page of its account
            _line(cv, L, heading, y)
            y -= 2 * lh
        st["y"] = y

    for t in tabs:
        _sign(t)
        first, heading = True, None
        for r in t.rows:
            if r.kind in ("noise", "page_break"):
                close()
                if r.kind == "noise":
                    _noise(cv, r, st)
                continue
            lines = lines_of(t, r, L)
            if r.kind in ("heading", "grand_total"):
                heading = None
                if r.kind == "heading":  # each account starts a new page
                    close()
            if st["y"] is None or st["y"] - (len(lines) - 1) * lh < L["bottom"]:
                open_page(t, first or t.repeat_header, heading)
                first = False
            heading = lines[0] if r.kind == "heading" else heading
            y0 = st["y"]
            for ln in lines:
                _line(cv, L, ln, st["y"])
                st["y"] -= lh
            r.pointer = {"page": st["page"], "line": round((H - 40 - y0) / lh) + 1,
                         "bbox": [L["left"], round(H - y0 - fs * 0.8, 1), L["right"], round(H - st["y"] - lh + fs * 0.25, 1)]}
            st["y"] -= lh  # one blank line after every block
    close()
    cv.save()


def render_xlsx(tabs: list, path: Path, L: dict, sheets: list) -> None:
    """Excel export of the same print: one sheet per part, one sheet row per printed line, amounts as numbers
    (negatives negative, shown with a trailing minus). Sets each row's pointer {sheet, cell, cells}."""
    wb = Workbook()
    wb.remove(wb.active)
    for t, name in zip(tabs, sheets):
        _sign(t)
        ws = wb.create_sheet(name)
        for i, s in enumerate([f"{t.title_lines[0]} {t.title_lines[1]}", *t.title_lines[2:]]):
            ws.cell(i + 1, 1, s).font = Font(bold=i == 0)
        rr = 5
        for hr in t.header_rows:
            for j, cell in enumerate(hr):
                if cell.text:
                    ws.cell(rr, (1, 3, 4, 5, 6, 7, 8, 9)[j], cell.text).font = Font(bold=True)
            rr += 1
        rr += 1
        for r in t.rows:
            if r.kind in ("noise", "page_break"):
                continue
            top, cells, bold = rr, {}, r.kind != "asset"
            for ln in lines_of(t, r, L, wrap_names=False):
                for key, col in (("id", 1), ("bez", 2), ("info", 3), ("label", 4)):
                    if ln.get(key):
                        ws.cell(rr, col, ln[key]).font = Font(bold=bold and col == 1)
                for col, (s, f, primary) in ln["amt"].items():
                    clean = s == r.truth["printed"][f] or not primary
                    x = ws.cell(rr, 5 + col, r.truth["fields"][f] / 100 * (-1 if isneg(s) else 1) if clean else s)
                    x.number_format, x.font = "#,##0.00;#,##0.00-", Font(bold=bold)
                    if primary:
                        cells[f] = x.coordinate
                rr += 1
            r.pointer = {"sheet": name, "cell": f"A{top}", "cells": cells}
            rr += 1
        for col, w in zip("ABCDEFGHI", (9, 44, 12, 14, 15, 15, 15, 15, 15)):
            ws.column_dimensions[col].width = w
    wb.save(path)
    core._freeze_zip(path)  # same input, same bytes


def client(seed: int, **kw) -> tuple:
    """Client with 3-digit accounts, 6-digit ids (account + counter), supplier text after a slash on some
    descriptions, and an invented town. Returns (client, town, extra invented names)."""
    c = core.make_client(seed, assets_per_class=(4, 12), **kw)
    rng, names = random.Random(f"{FAMILY}-{seed}"), []
    *stems, town = rng.sample([a + b for a in core.SYL_A for b in core.SYL_B], 11)
    out = []
    for a in c.assets:
        name = a.name
        if a.method != "keine" and len(names) < len(stems) and rng.random() < 0.3:  # at most 10 suppliers
            names.append(firm := f"{stems[len(names)]} {rng.choice(('GmbH', 'KG', 'GmbH & Co. KG'))}")
            name += f" / {rng.choice(('Fa.', 'Lieferant', 'Händler'))} {firm}"
        out.append(replace(a, asset_no=f"{int(a.asset_no):06d}", account=f"{int(a.account):03d}", name=name))
    c.assets = out
    town = town[:-2] + rng.choice(("hausen", "stedt", "ingen", "feld"))
    return c, town, names + [town]


def build(out: Path) -> None:
    d = out / FAMILY
    d.mkdir(parents=True, exist_ok=True)

    def key(n, c, names, fmt, variant, tabs, planted=()):
        meta = core.key_meta(c, doc_id=f"{FAMILY}_{n:02d}", family=FAMILY, doc_type="Anlagenspiegel", format=fmt,
                             variant=variant)
        meta["invented_names"] = c.names + names
        core.write_key(d / f"{FAMILY}_{n:02d}.key.json", meta, tabs, planted)

    def pdf(n, seed, variant="perfect", L=LAND, problems=None, **kw):
        c, town, names = client(seed, **kw)
        tabs = tables(c, town, L)
        planted = problems(tabs, c, random.Random(seed)) if problems else []
        render_pdf(tabs, d / f"{FAMILY}_{n:02d}.pdf", L)
        key(n, c, names, "pdf", variant, tabs, planted)

    def image(n, seed, fmt, **kw):  # perfect print, then scanned or outlined: no text layer
        c, town, names = client(seed, **kw)
        tabs = tables(c, town, LAND)
        with tempfile.TemporaryDirectory() as tmp:
            render_pdf(tabs, Path(tmp) / "print.pdf", LAND)
            if fmt == "scan":
                core.scan(Path(tmp) / "print.pdf", d / f"{FAMILY}_{n:02d}.pdf", fax=True, seed=seed)
            else:
                core.outline(Path(tmp) / "print.pdf", d / f"{FAMILY}_{n:02d}.pdf")
        key(n, c, names, fmt, "perfect", tabs)

    pdf(1, 301, year=2025, n_classes=8)
    c, town, names = client(302, year=2024, n_classes=7)  # 02 Excel export of the same print
    tabs = tables(c, town, LAND)
    render_xlsx(tabs, d / f"{FAMILY}_02.xlsx", LAND, ["Kontenübersicht", "Anlagen"])
    key(2, c, names, "xlsx", "perfect", tabs)
    image(3, 303, "scan", year=2025, fy_end=(6, 30), n_classes=6)
    image(4, 304, "outline", year=2025, n_classes=9)
    # 05: glued amounts in the dense detail lines + one account sum block that does not tie
    pdf(5, 305, "problem", problems=lambda ts, c, rng: [core.p_glued(ts[1], rng, 3)[1], core.p_non_tying(ts[1], rng)[1]],
        year=2025, n_classes=8)
    # 06: column heading only on the first detail page + a cover letter in front
    pdf(6, 306, "problem", problems=lambda ts, c, rng: [core.p_no_repeat_header(ts[1])[1], core.p_noise(ts[0], c)[1]],
        year=2024, n_classes=7)
    pdf(7, 307, "layout-mutation", L=PORT, year=2025, n_classes=8)  # portrait, synonyms, leading minus
