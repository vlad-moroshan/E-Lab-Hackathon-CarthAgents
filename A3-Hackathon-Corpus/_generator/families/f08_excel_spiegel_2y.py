"""f08_excel_spiegel_2y: hand-made Excel "Entwicklung des Anlagevermögens" at HGB § 266 line-item level, current year
on the first tab and prior year on the second (shifted one column left). Look-alike of _intern/profiles/Sthg.md:
3-row merged heading with real date cells, spacer columns, roll-forward by formula, hand-typed `=a+b+c` movement
cells, unlabelled SUM subtotals per roman-numeral group, a "lt. Bilanz" / "Differenz" check pair and an italic
reconciliation block with cell comments below the table.

Key: the leaf rows are balance-sheet line items, not assets (meta row_level = "bilanzposition"). They are keyed as
kind 'asset' (= leaf row) with position = the line item (A.II.2) and class = its § 266 group (A.II); the group sums
are 'position_subtotal' rows keyed by that group. The layout has no Zuschreibungen column, so the client's
write-back of the Beteiligung is dropped from the data.

xlsx: openpyxl writes the formulas, LibreOffice (soffice --convert-to xlsx) calculates them and saves the cached
values, as Excel does on save; check_corpus then reads the calculated values against the key.
"""

import random
import re
import subprocess
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from reportlab.pdfbase.pdfmetrics import stringWidth

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import core  # noqa: E402
from core import Cell  # noqa: E402

FAMILY = Path(__file__).stem
SOFFICE = "/opt/homebrew/bin/soffice"
GROUPS = {"A.I": ("I.", "Immaterielle Vermögensgegenstände"), "A.II": ("II.", "Sachanlagen"),
          "A.III": ("III.", "Finanzanlagen")}
CAPTION = {"A.I.2": "Entgeltlich erworbene Konzessionen und Lizenzen",
           "A.II.1": "Grundstücke, grundstücksgleiche Rechte und Bauten", "A.II.2": "Technische Anlagen und Maschinen",
           "A.II.3": "Andere Anlagen, Betriebs- und Geschäftsausstattung",
           "A.II.4": "Geleistete Anzahlungen und Anlagen im Bau", "A.III.3": "Beteiligungen"}
AMT = ["ahk_open", "zugang", "abgang", "umbuchung", "ahk_close", "kum_afa_open", "afa_gj", "afa_abgang",
       "kum_afa_close", "bw_close", "bw_prev"]
INPUT = ("ahk_open", "zugang", "abgang", "umbuchung", "kum_afa_open", "afa_gj", "afa_abgang")
MOVES = ("zugang", "abgang", "umbuchung", "afa_gj", "afa_abgang")
LAW = {"ahk_close": "{ahk_open}+{zugang}-{abgang}+{umbuchung}", "kum_afa_close": "{kum_afa_open}+{afa_gj}-{afa_abgang}",
       "bw_close": "{ahk_close}-{kum_afa_close}", "bw_prev": "{ahk_open}-{kum_afa_open}", "diff": "{bilanz}-{bw_close}"}
VALUE = set(AMT) | {"bilanz", "diff"}
THIN, DOUBLE = Side(style="thin"), Side(style="double")


def arial(**kw) -> Font:
    return Font(name="Arial", size=11, **kw)


def client(seed: int, year: int = 2025, fy_end: tuple = (12, 31)) -> core.Client:
    c = core.make_client(seed, year=year, fy_end=fy_end, n_classes=8, assets_per_class=(3, 6))
    for L in c.lives:  # no Zuschreibungen column in this layout: keep the impairment, drop the write-back
        L.imp = {y: v for y, v in L.imp.items() if v > 0}
    return c.snapshot(year)


def table(c: core.Client, *, num: bool = True, mutation: bool = False) -> core.Table:
    d0, d1, dp = c.fy_start, c.fy_end, c.fy_end.replace(year=c.year - 1)
    D0, D1, DP = map(core.fmt_date, (d0, d1, dp))
    if mutation:  # one header row, synonyms, no spacers, no check columns, sums above their rows
        num, fields = False, ["name", *AMT]
        header = [[Cell(s) for s in ("Bilanzposten", f"AK/HK {D0}", "Zugang", "Abgang", "Umbuchung", f"AK/HK {D1}",
                                     f"AfA kumuliert {D0}", "AfA lfd. Jahr", "AfA auf Abgänge", f"AfA kumuliert {D1}",
                                     f"Restbuchwert {D1}", f"Restbuchwert {DP}")]]
        title = [f"Anlagenspiegel {c.company}", f"Geschäftsjahr {D0} - {D1}, Beträge in EUR"]
    else:
        lead = ["num", "name"] if num else ["name"]
        fields = lead + ["ahk_open", "zugang", "abgang", "umbuchung", None, "ahk_close", None, "kum_afa_open", "afa_gj",
                         "afa_abgang", "kum_afa_close", None, "bw_close", "bw_prev", None, None, "bilanz", "diff"]
        header = [[Cell("", len(lead)), Cell("Anschaffungs- und Herstellungskosten", 6), Cell(""),
                   Cell("kumulierte Abschreibungen", 4), Cell(""), Cell("Buchwerte", 2), Cell("", 2),
                   Cell("lt. Bilanz"), Cell("Differenz")],
                  [Cell(s) for s in [""] * len(lead) + [D0, "Zugänge", "Abgänge", "Umbuchungen", "", D1, "", D0, "Zugänge",
                                                        "Abgänge", D1, "", D1, DP, "", "", "(aus Bilanz)", ""]],
                  [Cell("EUR" if f in VALUE else "") for f in fields]]
        title = ["Entwicklung des Anlagevermögens", f"Wirtschaftsjahr vom {D0} bis {D1}"]
    t = core.Table(fields=fields, col_widths=[10] * len(fields), align=["r" if f in VALUE else "l" for f in fields],
                   header_rows=header, title_lines=title, zero=None, font_size=6.5, page_header=c.company)
    t.dates = {D0: d0, D1: d1, DP: dp}
    label = (lambda n, s: {"num": n, "name": s}) if num else (lambda n, s: {"name": f"{n} {s}"})
    subs = []
    for g, (roman, gname) in GROUPS.items():
        poss = [p for p in core.POSITIONS if p.startswith(g + ".") and any(k.position == p for k in c.classes)]
        if not poss:
            continue
        assets = [a for a in c.assets if a.position in poss]
        if mutation:
            sub = t.add_total("position_subtotal", f"{roman} {gname}", assets, cls=g, year=c.year)
        else:
            t.add("heading", label(roman, gname))
            t.add("blank", {})
        sub_items = []
        for p in poss:
            pa = [a for a in assets if a.position == p]
            r = t.add("asset", {**{f: sum(getattr(a, f) for a in pa) for f in AMT}, **label(f"{p[-1]}.", CAPTION[p])},
                      {"year": c.year, "asset_no": None, "sub_no": None, "class": g, "account": None, "position": p,
                       "name": CAPTION[p]})
            r.parts = {f: [v for k in c.classes if (v := sum(getattr(a, f) for a in pa if a.cls == k.code))] for f in MOVES}
            for f in INPUT:  # typed cells stay empty when there is nothing to type
                if not r.truth["fields"][f]:
                    r.cells[fields.index(f)] = r.truth["printed"][f] = ""
            sub_items.append(r)
        if not mutation:
            sub = t.add_total("position_subtotal", "", assets, cls=g, year=c.year)
            t.add("blank", {})
        sub.kids, sub.from_blank = sub_items, len(subs) == 1 and not mutation  # one SUM range starts a row higher
        subs.append(sub)
    subs[0].paste = {"kum_afa_open"}  # one AfA-start subtotal pasted as a value
    grand = t.total("grand_total", "Anlagevermögen gesamt" if mutation else "", c.assets, year=c.year)
    t.rows.insert(0, grand) if mutation else t.rows.append(grand)
    grand.kids = subs
    for r in t.rows:
        if r.truth and "bilanz" in fields:
            r.cells[fields.index("bilanz")], r.cells[fields.index("diff")] = core.fmt(r.truth["fields"]["bw_close"]), "0,00"
    return t


def fit(t: core.Table, room: float = 790) -> core.Table:
    """PDF column widths from the widest text per column; smaller font until the table fits the landscape page."""
    while True:
        bold = core._bold(t.font)
        w = [0.0] * len(t.fields)
        for cells in [r.cells for r in t.rows] + [[c.text if c.colspan == 1 else "" for c in hr for _ in range(c.colspan)]
                                                  for hr in t.header_rows]:
            w = [max(a, stringWidth(s, bold, t.font_size)) for a, s in zip(w, cells)]
        t.col_widths = [x + 5 if x else 5 for x in w]
        if sum(t.col_widths) <= room or t.font_size <= 5:
            return t
        t.font_size -= 0.25


def content(t: core.Table, r, f: str, col: dict, values: bool, text: tuple):
    """What one value cell holds: a typed number, a hand-typed sum `=a+b`, a formula, or German text."""
    if f in text:
        return r.cells[t.fields.index(f)] or None
    fields, x = r.truth["fields"], r.x
    cents = fields["bw_close"] if f == "bilanz" else fields.get(f, 0)
    if r.kind == "asset":
        if f in LAW and not values and all(k in col for k in re.findall(r"\{(\w+)\}", LAW[f])):
            return "=" + LAW[f].format(**{k: f"{v}{x}" for k, v in col.items()})
        if f == "bilanz":  # copied from the balance sheet: binary float noise on some lines
            return cents / 100 + (1e-10 if x % 2 else 0)
        if not cents and f in INPUT:
            return None
        parts = r.parts.get(f, [])
        if len(parts) > 1 and not values:
            return ("=" + "+".join(f"{p / 100:.2f}" for p in parts)).replace("+-", "-")
        return cents / 100
    if values or f in getattr(r, "paste", ()):
        return cents / 100
    if f == "diff":
        return f"={col['bilanz']}{x}-{col['bw_close']}{x}"
    if r.kind == "grand_total":
        return "=" + "+".join(f"{col[f]}{k.x}" for k in r.kids)
    return f"=SUM({col[f]}{r.kids[0].x - r.from_blank}:{col[f]}{r.kids[-1].x})"


def write_tab(ws, t: core.Table, top: int = 1, *, values=False, text=(), draft=None) -> int:
    """Write one year block; sets each keyed row's pointer. draft: {row: cents} for a hidden stale Zugänge column."""
    col = {f: get_column_letter(j + 1) for j, f in enumerate(t.fields) if f}
    for i, s in enumerate(t.title_lines):
        ws.cell(top + i, t.fields.index("name") + 1, s).font = arial(bold=True)
    x, seen = top + len(t.title_lines) + 1, {}
    for i, hr in enumerate(t.header_rows):
        j = 1
        for cell in hr:
            y, d = ws.cell(x, j), t.dates.get(cell.text)
            if d:  # real date cells; a repeated date links to the first one (=C5)
                y.value = f"={seen[d]}" if d in seen else datetime(d.year, d.month, d.day)
                y.number_format = "DD.MM.YYYY"
                seen.setdefault(d, y.coordinate)
            elif cell.text:
                y.value = cell.text
            y.font = arial(bold=cell.text in ("lt. Bilanz", "Differenz"))
            if cell.text and cell.colspan > 1:
                ws.merge_cells(start_row=x, start_column=j, end_row=x, end_column=j + cell.colspan - 1)
                y.alignment = Alignment(horizontal="center")
            if cell.text and (cell.colspan > 1 or cell.text == "EUR"):
                for k in range(j, j + cell.colspan):
                    ws.cell(x, k).border = Border(bottom=THIN)
            j += cell.colspan
        x += 1
    x += 1
    for r in t.rows:
        r.x, x = x, x + 1
    for r in t.rows:
        if r.kind == "blank":
            continue
        cells = {}
        for j, f in enumerate(t.fields):
            y = ws.cell(r.x, j + 1)
            y.font = arial(bold=r.kind == "heading" or (r.kind in core.TOTAL_KINDS and bool(r.cells[t.fields.index("name")])))
            if f not in VALUE or r.kind == "heading":
                y.value = r.cells[j] or None
                continue
            y.value, y.number_format = content(t, r, f, col, values, text), "#,##0.00"
            if r.kind in ("asset", "grand_total"):
                y.border = Border(bottom=THIN if r.kind == "asset" else DOUBLE)
            if f in r.truth["printed"]:
                cells[f] = y.coordinate
        if r.truth:
            r.pointer = {"sheet": ws.title, "cell": f"A{r.x}", "cells": cells}
    if draft:  # the spacer after Umbuchungen holds an old draft of the Zugänge, hidden
        g = t.fields.index("umbuchung") + 2
        ws.cell(top + len(t.title_lines) + 2, g, "Zugänge lt. Entwurf").font = arial()
        for r in t.rows:
            if r.kind == "asset":
                ws.cell(r.x, g, draft[id(r)] / 100).number_format = "#,##0.00"
            elif r.truth:
                L = get_column_letter(g)
                ws.cell(r.x, g, "=" + "+".join(f"{L}{k.x}" for k in r.kids) if r.kind == "grand_total"
                        else f"=SUM({L}{r.kids[0].x}:{L}{r.kids[-1].x})").number_format = "#,##0.00"
        ws.column_dimensions[get_column_letter(g)].hidden = True
    for j, f in enumerate(t.fields):
        ws.column_dimensions[get_column_letter(j + 1)].width = {"num": 5, "name": 52, None: 2}.get(f, 16)
    ws.page_setup.orientation, ws.page_setup.fitToHeight = "landscape", 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = t.page_header
    return x


def recon(ws, t: core.Table, c: core.Client) -> None:
    """Italic reconciliation block below the table: other workpapers, notes, two comments. Not part of the key."""
    col = {f: get_column_letter(j + 1) for j, f in enumerate(t.fields) if f}
    g = next(r for r in t.rows if r.kind == "grand_total")
    G, nc = g.truth["fields"], t.fields.index("name") + 1
    gwg = sum(a.zugang for a in c.assets if a.method == "GWG")
    D0 = core.fmt_date(c.fy_start)
    lines = [("Abstimmung Vorjahreswerte", None, 0), (f"AHK {D0} lt. Anlagengitter Vorjahr (AV 2)", "ahk_open", -3),
             (f"kum. Abschreibungen {D0} lt. Anlagengitter Vorjahr (AV 2)", "kum_afa_open", 0), None,
             ("Abstimmung aktuelle Werte Anlagegitter", None, 0), ("Zugänge lt. Anlagegitter", "zugang", -gwg),
             ("Abgänge lt. Anlagegitter", "abgang", 0), ("Abschreibungen lt. Anlagegitter", "afa_gj", 0)]
    x = g.x + 3
    for line in lines:
        if line:
            label, f, delta = line
            ws.cell(x, nc, label).font = arial(italic=True)
            if f:
                ws.cell(x, t.fields.index(f) + 1, (G[f] + delta) / 100).number_format = "#,##0.00"
                d = ws.cell(x, t.fields.index("diff") + 1, f"={col[f]}{x}-{col[f]}{g.x}")
                d.number_format, d.font = "#,##0.00", arial(italic=True)
            if f == "ahk_open":
                ws.cell(x, t.fields.index(f) + 1).comment = Comment(
                    "Rundungsdifferenz aus dem Vorjahres-Gitter, nicht wesentlich", c.signer)
        x += 1
    notes = [f"Bearbeitungsstand {core.fmt_date(c.fy_end + timedelta(days=52))}: Werte lt. vorläufiger Bilanz"]
    if gwg:
        notes.append("Differenz Zugänge: GWG-Sammelposten im Anlagegitter nicht enthalten")
    for s in notes:
        ws.cell(x + 1, nc, s).font = arial(italic=True)
        x += 1
    ws.cell(len(t.title_lines) + 2, t.fields.index("bilanz") + 1).comment = Comment(f"Werte aus Bilanz-Entwurf vom {core.fmt_date(c.fy_end + timedelta(days=45))}", c.signer)


def recalc(src: Path, dst: Path) -> None:
    """LibreOffice calculates every formula and saves the workbook with cached values; fixed zip dates after."""
    files = sorted(src.glob("*.xlsx"))
    subprocess.run([SOFFICE, "--headless", f"-env:UserInstallation=file://{src}/lo", "--convert-to", "xlsx",
                    "--outdir", str(dst), *map(str, files)], check=True, capture_output=True, timeout=600)
    for f in files:
        assert (dst / f.name).exists(), f"soffice did not convert {f.name}"
        core._freeze_zip(dst / f.name)


def build(out: Path) -> None:
    d = out / FAMILY
    d.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)

        def key(n, c, fmt, variant, tables, planted=()):
            meta = core.key_meta(c, doc_id=f"{FAMILY}_{n:02d}", family=FAMILY, doc_type="Anlagenspiegel", format=fmt,
                                 variant=variant, years=sorted({r.truth["year"] for t in tables for r in t.rows if r.truth}),
                                 row_level="bilanzposition")
            core.write_key(d / f"{FAMILY}_{n:02d}.key.json", meta, tables, planted)

        def two_tabs(c, **kw):  # current year on Tabelle1, prior year shifted left on Tabelle2
            prev, cur = c.years(2)
            t1, t2 = table(cur), table(prev, num=False)
            wb = Workbook()
            ws1, ws2 = wb.active, wb.create_sheet("Tabelle2")
            ws1.title = "Tabelle1"
            write_tab(ws1, t1, **kw)
            if not kw.get("values"):
                recon(ws1, t1, cur)
            write_tab(ws2, t2)
            return wb, [t1, t2], core.p_multi_year([t1, t2])[1]

        def pdf(c, path):
            prev, cur = c.years(2)
            ts = [fit(table(cur)), fit(table(prev, num=False))]
            core.render_pdf(ts, path)
            return ts, core.p_multi_year(ts)[1]

        c = client(801)  # 01 perfect workbook
        wb, ts, my = two_tabs(c)
        wb.save(tmp / f"{FAMILY}_01.xlsx")
        key(1, c, "xlsx", "perfect", ts, [my])

        c = client(802, fy_end=(6, 30))  # 02 the same layout printed to PDF
        ts, my = pdf(c, d / f"{FAMILY}_02.pdf")
        key(2, c, "pdf", "perfect", ts, [my])

        c = client(803, year=2024)  # 03 scan of such a print
        ts, my = pdf(c, tmp / "print03.pdf")
        core.scan(tmp / "print03.pdf", d / f"{FAMILY}_03.pdf", dpi=200, seed=803)
        key(3, c, "scan", "perfect", ts, [my])

        c = client(804, fy_end=(9, 30))  # 04 clean print, glyphs as outlines (no text layer)
        ts, my = pdf(c, tmp / "print04.pdf")
        core.outline(tmp / "print04.pdf", d / f"{FAMILY}_04.pdf")
        key(4, c, "outline", "perfect", ts, [my])

        c = client(805)  # 05 Tabelle1 pasted from a PDF: values only, movement columns as German text
        wb, ts, my = two_tabs(c, values=True, text=MOVES)
        wb.save(tmp / f"{FAMILY}_05.xlsx")
        key(5, c, "xlsx", "problem", ts, [my, {
            "type": "numbers_as_text", "where": [{"sheet": "Tabelle1", "fields": list(MOVES)}],
            "expected_behaviour": "Tabelle1 holds pasted values; the movement columns are German-format text "
                                  "(1.234,56); parse them as numbers"}])

        c = client(806, fy_end=(6, 30))  # 06 a pasted AfA-start subtotal that is stale; the grand total adds it up
        prev, cur = c.years(2)
        t1, t2 = table(cur), table(prev, num=False)
        sub = max((r for r in t1.rows if r.kind == "position_subtotal"), key=lambda r: r.truth["fields"]["kum_afa_open"])
        grand = next(r for r in t1.rows if r.kind == "grand_total")
        delta, where = random.Random(806).choice((-1_234_500, 980_000, 2_500_000)), []
        for r in (sub, grand):
            true = r.truth["fields"]["kum_afa_open"]
            t1.set_amount(r, "kum_afa_open", true + delta)
            r.flags["*"], r.paste = "non_tying", {"kum_afa_open"} if r is sub else set()
            where.append({"row": r, "field": "kum_afa_open", "printed": true + delta, "sum_of_rows": true})
        wb = Workbook()
        ws1, ws2 = wb.active, wb.create_sheet("Tabelle2")
        ws1.title = "Tabelle1"
        write_tab(ws1, t1)
        recon(ws1, t1, cur)
        write_tab(ws2, t2)
        wb.save(tmp / f"{FAMILY}_06.xlsx")
        key(6, c, "xlsx", "problem", [t1, t2], [core.p_multi_year([t1, t2])[1], {
            "type": "non_tying_subtotal", "where": where,
            "expected_behaviour": "the group subtotal of the AfA start column is a pasted value, not a SUM; it and "
                                  f"the grand total differ from the sum of their rows by {core.fmt(delta)}; "
                                  "report the printed values and flag both"}])

        c = client(807)  # 07 hidden column with an old draft of the Zugänge between Umbuchungen and the AHK end
        prev, cur = c.years(2)
        t1, t2 = table(cur), table(prev, num=False)
        rng = random.Random(807)
        draft = {id(r): max(0, r.truth["fields"]["zugang"] - rng.choice((0, 150_000, r.truth["fields"]["zugang"] // 8)))
                 for r in t1.rows if r.kind == "asset"}
        wb = Workbook()
        ws1, ws2 = wb.active, wb.create_sheet("Tabelle2")
        ws1.title = "Tabelle1"
        write_tab(ws1, t1, draft=draft)
        recon(ws1, t1, cur)
        write_tab(ws2, t2)
        wb.save(tmp / f"{FAMILY}_07.xlsx")
        key(7, c, "xlsx", "problem", [t1, t2], [core.p_multi_year([t1, t2])[1], {
            "type": "hidden_column", "where": [{"sheet": "Tabelle1", "column": get_column_letter(t1.fields.index("umbuchung") + 2),
                                                "header": "Zugänge lt. Entwurf"}],
            "expected_behaviour": "the hidden column is a stale draft of the Zugänge, not part of the schedule; "
                                  "take Zugänge from the visible Zugänge column"}])

        c = client(808, fy_end=(6, 30))  # 08 layout mutation: one tab, both years stacked, sums above, synonyms
        prev, cur = c.years(2)
        ts = [table(cur, mutation=True), table(prev, mutation=True)]
        wb = Workbook()
        ws = wb.active
        ws.title = "AV-Spiegel"
        write_tab(ws, ts[1], top=write_tab(ws, ts[0]) + 2)
        wb.save(tmp / f"{FAMILY}_08.xlsx")
        key(8, c, "xlsx", "layout-mutation", ts, [core.p_multi_year(ts)[1]])

        recalc(tmp, d)
