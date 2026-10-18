#!/usr/bin/env python3
"""Check every *.key.json under OUT against its document. One OK/FAIL line per document, exit 1 on any FAIL.

(a) row laws on asset rows; subtotals = sum of their rows (rows flagged non_tying exempt; TEUR within rounding)
(b) perfect text PDFs: every printed amount is in the pdfplumber text of its page
(c) scans: no text layer on any page
(d) xlsx: the value at every pointer equals the key
(e) every invented name token has 0 hits in the client folders (prints token + file name only, never file text)

    python3 -I check_corpus.py OUT
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pdfplumber

sys.path.insert(0, str(Path(__file__).resolve().parent))
import core  # noqa: E402

SOURCES = [Path.home() / "Desktop/Client-Files-Confidential/anonyaa", Path.home() / "Desktop/Client-Files-Confidential/anony"]
PDFTOTEXT = core.PDFTOPPM.replace("pdftoppm", "pdftotext")
# (target, {term: sign}, terms that count as 0 when the layout does not print them)
LAWS = [
    ("ahk_close", {"ahk_open": 1, "zugang": 1, "abgang": -1, "umbuchung": 1}, {"umbuchung"}),
    ("kum_afa_close", {"kum_afa_open": 1, "afa_gj": 1, "afa_abgang": -1, "zuschreibung": -1}, {"zuschreibung"}),
    ("bw_close", {"ahk_close": 1, "kum_afa_close": -1}, set()),
    ("bw_prev", {"ahk_open": 1, "kum_afa_open": -1}, set()),
    ("bw_close", {"bw_prev": 1, "zugang": 1, "umbuchung": 1, "abgang_bw": -1, "afa_gj": -1, "zuschreibung": 1},
     {"umbuchung", "zuschreibung"}),
]
_texts: dict[str, str] | None = None


def source_texts() -> dict[str, str]:
    global _texts
    if _texts is None:
        _texts = {}
        for f in sorted(p for d in SOURCES if d.is_dir() for p in d.rglob("*") if p.is_file()):
            if f.suffix.lower() == ".pdf":
                _texts[f.name] = subprocess.run([PDFTOTEXT, "-q", str(f), "-"], capture_output=True, text=True).stdout
            elif f.suffix.lower() == ".xlsx":
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                _texts[f.name] = " ".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True)
                                          for v in r if isinstance(v, str))
                wb.close()
        if not _texts:
            raise SystemExit("no client files found for the name check: " + ", ".join(map(str, SOURCES)))
    return _texts


def arithmetic(key: dict) -> list[str]:
    errs, rows = [], key["rows"]
    tol = (lambda n: 50_000 * n) if key["meta"]["unit"] == "TEUR" else (lambda n: 0)
    assets = [r for r in rows if r["kind"] == "asset"]
    for r in assets:
        f = r["fields"]
        for target, terms, optional in LAWS:
            if target in f and all(t in f for t in terms if t not in optional):
                diff = f[target] - sum(s * f.get(t, 0) for t, s in terms.items())
                if abs(diff) > tol(len(terms) + 1):
                    errs.append(f"{r['asset_no']}/{r.get('sub_no')} {target} off by {diff}")
        errs += [f"{r['asset_no']} {k} negative" for k, v in f.items() if v < 0 and k != "umbuchung"]
    for i, r in enumerate(rows):
        if r["kind"] not in core.TOTAL_KINDS or "non_tying" in r.get("flags", {}).values():
            continue
        kids = [a for j, a in enumerate(rows) if a["kind"] == "asset" and a["year"] == r["year"]
                and all(a.get(k) == r[k] for k in ("class", "account", "position") if r.get(k))
                and (r["kind"] != "carry_over" or j < i)]
        for f, v in r["fields"].items():
            diff = v - sum(a["fields"].get(f, 0) for a in kids)
            if abs(diff) > tol(len(kids) + 1):
                errs.append(f"{r['kind']} {r.get('class') or r.get('position') or ''} {f} off by {diff}")
    return errs


def pdf_checks(key: dict, path: Path) -> list[str]:
    errs, meta = [], key["meta"]
    with pdfplumber.open(path) as pdf:
        texts = [p.extract_text() or "" for p in pdf.pages]
    if meta["format"] in ("scan", "outline"):
        return [f"page {i + 1} has a text layer" for i, t in enumerate(texts) if t.strip()]
    if meta["variant"] != "perfect":
        return []
    for r in key["rows"]:
        page = texts[r["pointer"]["page"] - 1]
        for f, s in r.get("printed", {}).items():
            if s and s not in page and s not in page.replace(" ", ""):
                errs.append(f"page {r['pointer']['page']}: '{s}' ({f} of {r.get('asset_no') or r['kind']}) not found")
    return errs


def xlsx_checks(key: dict, path: Path) -> list[str]:
    errs, scale = [], 100_000 if key["meta"]["unit"] == "TEUR" else 100
    wb = openpyxl.load_workbook(path, data_only=True)
    for r in key["rows"]:
        p = r["pointer"]
        for f, coord in p["cells"].items():
            if f in r.get("flags", {}):
                continue
            v, want, printed = wb[p["sheet"]][coord].value, r["fields"][f], r["printed"][f]
            ok = ((v in (None, "") and want == 0 and not printed)
                  or (isinstance(v, (int, float)) and abs(round(v * scale)) == abs(want))
                  or (isinstance(v, str) and v == printed))
            if not ok:
                errs.append(f"{p['sheet']}!{coord} = {v!r}, key {f} = {want}")
    return errs


def name_hits(key: dict) -> list[str]:
    tokens = {t for n in key["meta"].get("invented_names", []) for t in re.findall(r"[A-Za-zÄÖÜäöüß]{4,}", n)
              if t.lower() not in core.GENERIC}
    return [f"name hit: {t} in {name}" for t in sorted(tokens) for name, text in source_texts().items()
            if re.search(rf"\b{re.escape(t)}\b", text, re.I)]


def main(out: Path) -> int:
    out, fails, n = Path(out), 0, 0
    keys = sorted(out.rglob("*.key.json"))
    docs = {p.with_suffix("").with_suffix("") for p in keys}
    found = (p for p in out.rglob("*") if p.suffix in (".pdf", ".xlsx"))
    for d in sorted(p for p in found if not any(s.startswith("_") for s in p.relative_to(out).parts)):
        if d.with_suffix("") not in docs:
            n, fails = n + 1, fails + 1
            print(f"FAIL {d.relative_to(out)}: no answer key")
    for k in keys:
        n += 1
        try:
            key = json.loads(k.read_text(encoding="utf-8"))
            doc = k.with_name(key["meta"]["doc_id"] + (".xlsx" if key["meta"]["format"] == "xlsx" else ".pdf"))
            errs = [] if key["rows"] else ["key has no rows"]
            errs += ["document missing"] if not doc.exists() else (xlsx_checks if doc.suffix == ".xlsx" else pdf_checks)(key, doc)
            errs += arithmetic(key) + name_hits(key)
        except Exception as e:  # a broken key is a FAIL, not a crash
            errs = [f"{type(e).__name__}: {e}"]
        fails += bool(errs)
        print(("FAIL " if errs else "OK   ") + str(k.relative_to(out)).removesuffix(".key.json")
              + (": " + "; ".join(errs[:4]) + (f" (+{len(errs) - 4} more)" if len(errs) > 4 else "") if errs else ""))
    print(f"{n} documents, {n - fails} OK, {fails} FAIL")
    return 1 if fails or not n else 0


if __name__ == "__main__":
    sys.exit(main(Path(sys.argv[1])))
