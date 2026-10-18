"""Extraction contract and template writer. Runtime code never imports evaluation fixtures.

Layout specs contain identities and source references, never financial values. The writer consumes parsed records.
"""
from __future__ import annotations

import json
import math
from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path

from jsonschema import Draft202012Validator
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import coordinate_to_tuple

from carth import ingest
from carth.numbers import source_amount
from carth.ingest.xlsx import empty_amount_cell

TEMPLATE = Path(__file__).resolve().parents[1] / 'challenge/Arbeitsversion-merged-layout-20.xlsx'
GRID = dict(ahk_open='D', zugang='E', abgang='F', umbuchung='G', **{'net:bew_ahk':'H'}, ahk_close='I',
            kum_afa_open='J', afa_gj_plan='K', afa_gj='L', afa_gj_ausserplan='M', afa_abgang='N',
            zuschreibung='O', afa_umbuchung='P', **{'net:bew_wb':'Q'}, kum_afa_close='R', bw_close='S', bw_prev='T', abgang_bw='AA')
REGISTER = dict(ahk_open='W', zugang='X', abgang='Y', umbuchung='Z', ahk_close='AA', kum_afa_open='AB',
                afa_gj_plan='AC', afa_gj_ausserplan='AD', afa_abgang='AE', zuschreibung='AF', afa_umbuchung='AG',
                kum_afa_close='AH', bw_close='AI', bw_prev='AJ', afa_gj='AR', abgang_bw='AS')


def check_output_directory(out: Path):
    resolved = Path(out).resolve()
    if resolved.is_relative_to(TEMPLATE.parent.resolve()) or resolved.is_relative_to((TEMPLATE.parents[1] / 'A3-Hackathon-Corpus').resolve()):
        raise ValueError('protected output directory')


def validate(result):
    schema = json.loads(Path(__file__).with_name('extraction.schema.json').read_text())
    Draft202012Validator(schema).validate(result)
    rows = {r['id']: r for r in result['rows']}
    if len(rows) != len(result['rows']):
        raise ValueError('duplicate row identity')
    seen = set()
    for value in result['values']:
        pair = value['row'], value['field']
        if value['row'] not in rows or pair in seen:
            raise ValueError('unknown row or duplicate row/field')
        seen.add(pair)
        if value['value'] is None and not value['flags']:
            raise ValueError('blank requires a reason flag')
        if value['value'] is not None and (value['pointer'] is None or value['trace'] is None):
            raise ValueError('observed amount requires source and conversion trace')
        if rows[value['row']]['year'] is None:
            if value['value'] is not None or not rows[value['row']].get('flags'):
                raise ValueError('unknown reporting year requires flagged abstention')
        p = value['pointer']
        if p and 'bbox' in p:
            b = p['bbox']
            if not all(math.isfinite(x) for x in b) or not (b[0] < b[2] and b[1] < b[3]):
                raise ValueError('invalid PDF source box')


def extract(doc: dict, spec: dict | None, doc_id: str) -> dict:
    """Materialize a reviewed layout hypothesis using token IDs or identified empty workbook cells."""
    result = dict(version=1, document=dict(id=doc_id, source=doc['source'], sha256=doc['sha256']),
                  rows=[], values=[], flags=list(doc.get('flags', [])))
    if spec is None:
        result['flags'].append('needs_layout: no verified row/column references supplied')
        return result
    if spec.get('sha256') != doc['sha256']:
        raise ValueError('layout spec fingerprint does not match source')
    result['flags'].extend(spec.get('notes', []))
    at = {t['id']: t for t in doc['tokens']}
    cells = {(t.get('sheet'), t.get('cell')): t for t in doc['tokens']}
    used = set()
    for row in spec['rows']:
        unit = row.get('unit', spec.get('unit'))
        if unit not in ('EUR', 'TEUR'):
            raise ValueError('each layout row must identify EUR or TEUR')
        ident = {k:v for k,v in row.items() if k != 'fields'}
        result['rows'].append(ident)
        for field, ref in row['fields'].items():
            if field not in GRID:
                raise ValueError(f'unknown amount field: {field}')
            if set(ref) - {'token', 'sheet', 'cell', 'amount_column', 'blank_reason'}:
                raise ValueError('layout fields may contain source references only')
            value = dict(row=row['id'], field=field, value=None, pointer=None, flags=[], trace=None, workbook=None)
            tok = at.get(ref.get('token'))
            empty = False
            if 'sheet' in ref and 'cell' in ref:
                sheet, coord = ref['sheet'], ref['cell']
                tok = cells.get((sheet, coord))
                info = next((s for s in doc.get('sheets', []) if s['sheet'] == sheet), None)
                rr, cc = coordinate_to_tuple(coord)
                if not info or not (1 <= rr <= info['max_row'] and 1 <= cc <= info['max_col']):
                    value['flags'].append('cell outside identified source area')
                else:
                    value['pointer'] = dict(sheet=sheet, cell=coord)
                    empty = tok is None and empty_amount_cell(doc, sheet, coord)
            if tok:
                value['pointer'] = ({'sheet':tok['sheet'], 'cell':tok['cell']} if 'cell' in tok else
                                    {'page':tok['page'], 'bbox':tok['bbox']})
            reason = None
            raw, signed = None, field in ('umbuchung', 'afa_umbuchung', 'net:bew_ahk', 'net:bew_wb')
            if value['flags']:
                reason = value['flags'][0]
            elif not ref.get('amount_column'):
                reason = 'not an identified amount column'
            elif empty:
                value['value'], raw = 0, ''
            elif not tok:
                reason = ref.get('blank_reason', 'source reference unavailable')
            elif tok.get('kind') in ('formula_uncached', 'date', 'error'):
                reason = 'source value unavailable: ' + tok['kind']
            elif tok.get('orientation_uncertain'):
                reason = 'OCR orientation unresolved'
            elif tok.get('kind') == 'number':
                raw = str(tok['value'])
                cents = Decimal(raw) * (100000 if unit == 'TEUR' else 100)
                if not cents.is_finite() or cents != cents.to_integral_value():
                    reason = 'source precision is not integer cents'
                else:
                    value['value'] = int(cents) if signed else abs(int(cents))
            else:
                raw = tok['text']
                value['value'], reason = source_amount(raw, amount_column=True, unit=unit, signed=signed,
                                                       ocr=tok.get('src','').startswith('ocr'))
            if reason:
                value['flags'] = [reason]
            if value['value'] is not None:
                occurrence = json.dumps(value['pointer'], sort_keys=True)
                if occurrence in used:
                    value['value'], value['flags'] = None, ['source occurrence already assigned']
                used.add(occurrence)
            if value['value'] is not None:
                value['trace'] = dict(raw=raw, unit=unit, signed=signed, normalization='signed' if signed else 'magnitude',
                                      token=tok['id'] if tok else None, formula=tok.get('formula') if tok else None,
                                      cache_state=tok.get('cache_state') if tok else 'empty_cell')
            result['values'].append(value)
    if not result['rows'] or not result['values']:
        result['flags'].append('needs_layout: no amount fields identified')
    validate(result)
    return result


def write_outputs(result: dict, out: Path, template: Path = TEMPLATE):
    """One JSON and copied Arbeitsversion per document. Multi-year sheets stay in one workbook."""
    result = deepcopy(result)
    validate(result)
    out = Path(out)
    stem = result['document']['id']
    if Path(stem).name != stem or stem in ('.', '..'):
        raise ValueError('unsafe document ID')
    check_output_directory(out)
    if (out / f'{stem}.xlsx').resolve() == Path(result['document']['source']).resolve():
        raise ValueError('output would overwrite source workbook')
    out.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(template)
    rows = {r['id']:r for r in result['rows']}
    years = sorted({r['year'] for r in rows.values()}, key=lambda y: (y is None, y or 0))
    targets = {}
    groups = {}
    for row in rows.values():
        register = row['granularity'] == 'asset' and row['kind'] == 'asset'
        position = row['kind'] == 'position_subtotal' or (row['kind'] not in ('class_subtotal', 'account_subtotal') and row['granularity'] == 'position')
        base = 'Anlagenverzeichnis' if register else ('Gitter nach Bilanzposition' if position else 'Gitter nach Anlagenklasse')
        groups.setdefault((base, row['year']), []).append(row)
    for (base, year), records in groups.items():
        ws = wb[base] if len(years) <= 1 else wb.copy_worksheet(wb[base])
        if len(years) > 1:
            ws.title = f'{base[:26]} {year if year is not None else "?"}'
        register = base == 'Anlagenverzeichnis'
        start = 2 if register else 8
        limit = 501 if register else 27
        if start + len(records) - 1 > limit:
            ws.insert_rows(limit + 1, start + len(records) - 1 - limit)
        if register:
            ws['AT1'] = 'Unternummer (Quelle)'
            ws['AR1'] = 'AfA GJ gesamt (Quelle; Aufteilung unbekannt)'
            ws['AS1'] = 'Abgang Buchwert (Quelle)'
        else:
            # Preserve a reported year without inventing its fiscal end date.
            ws['AA7'] = 'Abgang Buchwert (Quelle)'
            ws['Y7'] = 'Quellgranularität'
            ws['Z7'] = 'Quellkonto'
            ws['C2'] = None
            ws['A4'] = f'Berichtsjahr: {year if year is not None else "unbekannt"}; Geschäftsjahresende nicht abgeleitet'
        # Keep template checks, but clear imported amount slots so blanks stay blank.
        for rr in range(start, max(limit, start + len(records) - 1) + 1):
            for col in (REGISTER if register else GRID).values():
                cell = ws[f'{col}{rr}']
                if not isinstance(cell, MergedCell):
                    cell.value = None
        for rr, row in enumerate(records, start):
            targets[row['id']] = ws, rr, REGISTER if register else GRID
            if register:
                for col, key in [('A','asset_no'), ('B','name'), ('D','class'), ('F','position'), ('H','account')]:
                    ws[f'{col}{rr}'] = row.get(key)
                ws[f'AT{rr}'] = row.get('sub_no')
            else:
                ws[f'B{rr}'] = row.get('position') if base.endswith('Bilanzposition') else row.get('class') or row.get('account')
                ws[f'C{rr}'] = row.get('name')
                ws[f'V{rr}'] = row['kind']
                ws[f'Y{rr}'] = row['granularity']
                ws[f'Z{rr}'] = row.get('account')
                if base.endswith('Anlagenklasse'):
                    ws[f'W{rr}'] = row.get('position')
    # All per-value records also survive in a source ledger, including fields without a core template equivalent.
    ledger = wb.create_sheet('Quellwerte')
    ledger.append(['Zeile', 'Jahr', 'Feld', 'EUR', 'Quelle', 'Flags', 'Original', 'Konvertierung', 'Ziel',
                   'Anlagen-Nr.', 'Unternummer', 'Klasse', 'Konto', 'Bilanzposition'])
    for v in result['values']:
        row = rows[v['row']]
        ws, rr, mapping = targets[v['row']]
        if v['field'] in mapping:
            cell = ws[f"{mapping[v['field']]}{rr}"]
            if isinstance(cell, MergedCell):
                raise ValueError(f'merged template target: {ws.title}!{cell.coordinate}')
            if v['value'] is not None:
                cell.value = v['value'] / 100
            cell.comment = Comment(json.dumps({'pointer':v['pointer'], 'trace':v['trace'], 'flags':v['flags']}, ensure_ascii=False), 'Quelle')
            v['workbook'] = {'sheet':ws.title, 'cell':cell.coordinate, 'scale':'cents/100'}
        ledger.append([v['row'], row['year'], v['field'], v['value']/100 if v['value'] is not None else None,
                       json.dumps(v['pointer']), '; '.join(v['flags']), (v['trace'] or {}).get('raw'),
                       json.dumps(v['trace']), json.dumps(v['workbook']), row.get('asset_no'), row.get('sub_no'),
                       row.get('class'), row.get('account'), row.get('position')])
    if result['flags']:
        wb['Offene Punkte']['A4'] = '; '.join(result['flags'])
    validate(result)
    wp, jp = out / f'{stem}.xlsx', out / f'{stem}.json'
    wb.save(wp)
    jp.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return jp, wp


def run(input_path: Path, out: Path, specs: Path | None = None) -> int:
    check_output_directory(out)
    failures, incomplete = [], False
    ids = ingest.document_ids(input_path)
    for path, doc_id in ids.items():
        try:
            doc = ingest.ingest(path)
            sp = specs / f'{doc_id}.layout.json' if specs else None
            spec = json.loads(sp.read_text()) if sp and sp.exists() else None
            result = extract(doc, spec, doc_id)
            write_outputs(result, out)
            incomplete |= bool(result['flags']) or any(v['value'] is None for v in result['values'])
        except Exception as exc:
            failures.append({'source':str(path), 'error':f'{type(exc).__name__}: {exc}'})
    ingest.save({'failures':failures}, out / '_failures.json')
    return int(bool(failures) or incomplete or not ids)
