"""xlsx -> cell tokens. Keeps everything a reader needs to judge relevance later, decides nothing:
hidden rows / columns / sheets (f09 hides every asset row in an outline; f08_07 hides a stale draft column),
outline level, bold, number format, merged range, and the cached value of formulas (Excel / LibreOffice save it).
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter


def _hidden_cols(ws) -> set[int]:
    out = set()
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or column_index_from_string(key)
            out.update(range(lo, (dim.max or lo) + 1))
    return out


def xlsx_tokens(path: Path) -> tuple[list[dict], list[dict]]:
    """Returns (sheets, tokens). One token per non-empty cell."""
    wb = load_workbook(path, data_only=True)
    sheets, tokens = [], []
    for ws in wb.worksheets:
        hcols = _hidden_cols(ws)
        merged = {}
        for rng in ws.merged_cells.ranges:
            merged[f"{get_column_letter(rng.min_col)}{rng.min_row}"] = str(rng)
        sheets.append({"sheet": ws.title, "state": ws.sheet_state, "max_row": ws.max_row, "max_col": ws.max_column,
                       "hidden_cols": sorted(hcols)})
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if isinstance(v, (datetime, date)):
                    v = v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
                rd = ws.row_dimensions[c.row]
                tokens.append({
                    "sheet": ws.title, "cell": c.coordinate, "row": c.row, "col": c.column,
                    "text": v if isinstance(v, str) else str(v),
                    "value": v,  # number / str / ISO date: the stored value, untouched
                    "kind": "number" if isinstance(v, (int, float)) and not isinstance(v, bool) else
                            "date" if c.is_date else "text",
                    "number_format": c.number_format,
                    "hidden": bool(rd.hidden or c.column in hcols or ws.sheet_state != "visible"),
                    "outline": rd.outlineLevel or 0,
                    "bold": bool(c.font and c.font.b),
                    "merged": merged.get(c.coordinate),
                    "src": "xlsx", "conf": 1.0,
                })
    return sheets, tokens
