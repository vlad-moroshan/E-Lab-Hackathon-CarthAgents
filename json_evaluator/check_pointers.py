#!/usr/bin/env python3
"""Pointer re-read: open the source document and check that each value really sits where its pointer says.

- xlsx: the cell at pointer.cells[field] holds the value (number x unit, German-format text, or empty = 0);
        a pointer into a hidden column is a warning (hidden columns are often stale drafts)
- text PDF: the printed string is among the characters inside the pointer box (pointer.boxes[field] if given,
        else the row bbox) and parses to the value; without a printed string, some amount in the box must
- scan / outline PDF (no text layer): not verifiable here, counted as such

Needs PyMuPDF and openpyxl (run with the project .venv).

    .venv/bin/python "json evaluator/check_pointers.py" PRED_DIR [CORPUS_DIR]
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (AMOUNT_TOKEN, DEFAULT_CORPUS, DERIVED_FLAGS, Corpus, field_box, issue, iter_predictions, load,  # noqa: E402
                    parse_amount, unit_scale)

CHECK = "pointers"
SKIP_FLAGS = DERIVED_FLAGS  # not read from their own characters: nothing to re-read


def _norm(s: str) -> str:
    return s.replace(" ", "").replace(" ", "").replace("−", "-")


def check_pdf(pred: dict, path: Path) -> tuple[list[dict], dict]:
    import pymupdf

    out, stats = [], {"verified": 0, "failed": 0, "unverifiable": 0}
    unit = pred["meta"].get("unit", "EUR")
    doc = pymupdf.open(path)
    words = [pg.get_text("words") for pg in doc]
    if not any(words):
        n = sum(1 for r in pred["rows"] for v in (r.get("fields") or {}).values() if v is not None)
        stats["unverifiable"] = n
        return [issue(CHECK, "info", f"no text layer ({n} values not re-readable without OCR)")], stats
    for r in pred["rows"]:
        p, flags = r.get("pointer") or {}, r.get("flags") or {}
        page = p.get("page")
        if not isinstance(page, int) or not 1 <= page <= len(doc):
            out.append(issue(CHECK, "error", f"page {page} outside document (1..{len(doc)})", r))
            stats["failed"] += len(r.get("fields") or {})
            continue
        for f, v in (r.get("fields") or {}).items():
            if v is None or flags.get(f) in SKIP_FLAGS:
                continue
            box = field_box(r, f)
            if not box:
                continue
            rect = pymupdf.Rect(box) + (-1, -1, 1, 1)
            text = "".join(w[4] for w in words[page - 1] if pymupdf.Rect(w[:4]).intersects(rect))
            printed = (r.get("printed") or {}).get(f)
            if printed:
                ok_text = _norm(printed) in _norm(text)
                parsed = parse_amount(printed, unit)
                ok_val = flags.get(f) == "glued" or (parsed is not None and abs(parsed) == abs(v))
                if ok_text and ok_val:
                    stats["verified"] += 1
                else:
                    stats["failed"] += 1
                    why = "not inside the box" if not ok_text else f"parses to {parsed}, value {v}"
                    out.append(issue(CHECK, "error", f"printed {printed!r} {why}", r, f))
            elif v == 0:
                stats["verified"] += 1  # empty cell = 0
            else:
                toks = [abs(parse_amount(t, unit) or 0) for w in words[page - 1]
                        if pymupdf.Rect(w[:4]).intersects(rect) for t in AMOUNT_TOKEN.findall(w[4])]
                if abs(v) in toks:
                    stats["verified"] += 1
                else:
                    stats["failed"] += 1
                    out.append(issue(CHECK, "error", f"no amount {v / 100:,.2f} inside the box", r, f))
    return out, stats


def check_xlsx(pred: dict, path: Path) -> tuple[list[dict], dict]:
    import openpyxl
    from openpyxl.utils import coordinate_to_tuple, get_column_letter

    out, stats = [], {"verified": 0, "failed": 0, "unverifiable": 0}
    unit = pred["meta"].get("unit", "EUR")
    scale = unit_scale(unit)
    wb = openpyxl.load_workbook(path, data_only=True)
    for r in pred["rows"]:
        p, flags = r.get("pointer") or {}, r.get("flags") or {}
        sheet = p.get("sheet")
        if sheet not in wb.sheetnames:
            out.append(issue(CHECK, "error", f"sheet {sheet!r} not in workbook", r))
            stats["failed"] += len(r.get("fields") or {})
            continue
        ws = wb[sheet]
        cells = p.get("cells") or {}
        for f, v in (r.get("fields") or {}).items():
            if v is None or flags.get(f) in SKIP_FLAGS:
                continue
            coord = cells.get(f)
            if not coord:
                if v == 0:
                    stats["verified"] += 1
                else:
                    stats["failed"] += 1
                    out.append(issue(CHECK, "error", "value has no cell", r, f))
                continue
            col = get_column_letter(coordinate_to_tuple(coord)[1])
            if ws.column_dimensions[col].hidden:
                out.append(issue(CHECK, "warn", f"{coord} is in hidden column {col}", r, f))
            cv = ws[coord].value
            if cv is None or cv == "":
                ok = v == 0
            elif isinstance(cv, (int, float)) and not isinstance(cv, bool):
                ok = abs(round(cv * scale)) == abs(v)
            elif isinstance(cv, str):
                parsed = parse_amount(cv, unit)
                ok = parsed is not None and abs(parsed) == abs(v)
            else:
                ok = False
            stats["verified" if ok else "failed"] += 1
            if not ok:
                out.append(issue(CHECK, "error", f"{sheet}!{coord} holds {cv!r}, value is {v / 100:,.2f}", r, f))
    return out, stats


def check(pred: dict, source: Path | None) -> tuple[list[dict], dict]:
    if source is None:
        return [issue(CHECK, "info", "source document not found")], {"verified": 0, "failed": 0, "unverifiable": 0}
    return (check_xlsx if source.suffix == ".xlsx" else check_pdf)(pred, source)


def main(pred_dir: str, corpus_dir: str | None = None) -> int:
    corpus, bad = Corpus(Path(corpus_dir) if corpus_dir else DEFAULT_CORPUS), 0
    for doc_id, path in iter_predictions(Path(pred_dir)):
        res, st = check(load(path), corpus.source(doc_id))
        errs = [i for i in res if i["severity"] == "error"]
        bad += bool(errs)
        print(("FAIL " if errs else "OK   ") + f"{doc_id} verified {st['verified']} failed {st['failed']} "
              f"unverifiable {st['unverifiable']}" + (": " + "; ".join(
                  f"{e.get('row', '')} {e.get('field', '')} {e['message']}" for e in errs[:2]) if errs else ""))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:3]))
